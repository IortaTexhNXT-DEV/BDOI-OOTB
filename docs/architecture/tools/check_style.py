"""Wording check for the architecture documents and the review and test reports.

    python3 docs/architecture/tools/check_style.py            # default file set (below)
    python3 docs/architecture/tools/check_style.py a.md b.pdf # only these files

Scans Markdown sources as text, .docx files through python-docx (paragraphs and tables) and .pdf files through
pdftotext, and reports every line with:
  - a stock word the customer does not want (seamless, robust, comprehensive, leverage, delve, empower, cutting-edge,
    streamline, unlock, elevate, holistic, synergy, state-of-the-art)
  - an emoji or pictograph
  - a long dash used as a connector (em dash, or an en dash with spaces around it)
  - a name that must not appear (BDOI, BIBS, BDO as a brand, Claude, Anthropic, AI as a word)
Exit code 1 when anything is found.
"""
import glob
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))

DEFAULT = [
    'docs/architecture/README.md',
    'docs/architecture/tools/src/*.md',
    'docs/architecture/*.docx',
    'docs/architecture/*.pdf',
    'docs/e2e/E2E_REPORT.md',
    'docs/e2e/PERSONA_WALK.md',
    'docs/review/CODE_REVIEW.md',
]

STOCK = ['seamless', 'robust', 'comprehensive', 'leverage', 'delve', 'empower', 'cutting-edge', 'streamline', 'unlock',
         'elevate', 'holistic', 'synergy', 'synergies', 'state-of-the-art']
CHECKS = [
    ('stock word', re.compile(r'\b(' + '|'.join(re.escape(w) for w in STOCK) + r')\w*', re.IGNORECASE)),
    ('emoji', re.compile('[\U0001F000-\U0001FAFF☀-➿⬀-⯿️]')),
    ('long dash', re.compile('—|―| – ')),
    ('name', re.compile(r'\b(BDOI|BIBS|Claude|Anthropic)\b|\bBDO\b(?! (bank|Unibank))|\bAI\b')),
]


def lines_of(path):
    if path.endswith('.md'):
        return open(path, encoding='utf-8').read().splitlines()
    if path.endswith('.docx'):
        from docx import Document
        doc = Document(path)
        out = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    out.extend(p.text for p in c.paragraphs)
        return out
    if path.endswith('.pdf'):
        return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True, check=True).stdout.splitlines()
    return []


def main():
    patterns = sys.argv[1:] or DEFAULT
    files = []
    for p in patterns:
        files.extend(sorted(glob.glob(os.path.join(REPO, p))) if not os.path.isabs(p) else sorted(glob.glob(p)))
    total = 0
    for f in files:
        for n, line in enumerate(lines_of(f), 1):
            for label, rx in CHECKS:
                for m in rx.finditer(line):
                    total += 1
                    print(f'{os.path.relpath(f, REPO)}:{n}: {label}: "{m.group(0)}" in: {line.strip()[:120]}')
    print(f'{len(files)} files checked, {total} findings')
    return 1 if total else 0


if __name__ == '__main__':
    sys.exit(main())
