"""Build the BrokerVerse persona guides (one .pptx per persona) with python-pptx.

Usage:  python3 docs/decks/tools/build_decks.py --shots <dir with <user>__<screen>.png> [--out docs/decks]

Screenshots come from docs/decks/tools/capture_screens.py (live screens, one sign-in per persona) and from the user
manual's screens in docs/manual/images ('manual:<name>'). Content lives in personas.py next to this file.
"""
import argparse, io, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from personas import PERSONAS, FLOW_STEPS, REPORTS, JOBS  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
LOGO = os.path.join(ROOT, 'brokerverse', 'public', 'bdoi', 'iorta-technxt.png')
MANUAL_IMAGES = os.path.join(ROOT, 'docs', 'manual', 'images')


def shot_path(img, shots):
    """'manual:<name>' -> docs/manual/images/<name>.jpg; anything else is a file in the --shots directory."""
    if img.startswith('manual:'):
        return os.path.join(MANUAL_IMAGES, img[len('manual:'):] + '.jpg')
    return os.path.join(shots, img)

DEEP = RGBColor(0x0B, 0x4F, 0x9C)
BRIGHT = RGBColor(0x1E, 0x88, 0xE5)
GREEN = RGBColor(0x43, 0xA0, 0x47)
TEXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x67, 0x78)
LIGHT = RGBColor(0xF5, 0xF8, 0xFC)
LINE = RGBColor(0xD5, 0xDE, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE_BLUE = RGBColor(0xE3, 0xEF, 0xFB)
PALE_GREEN = RGBColor(0xE6, 0xF4, 0xE7)
FONT = 'Calibri'

SW, SH = 13.333, 7.5
LOGO_RATIO = 604 / 178


# ------------------------------------------------------------------ primitives
def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, lw=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, runs, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False, margin=0.0, spacing=0):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, Inches(margin))
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        parts = p if isinstance(p, list) else [(p, {})]
        for t, o in parts:
            r = para.add_run(); r.text = t
            f = r.font; f.name = FONT; f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold)
            f.italic = o.get('italic', italic); f.color.rgb = o.get('color', color)
    return tb


def bullets(slide, x, y, w, h, items, size=13, color=TEXT, spacing=6, label_color=DEEP, marker=BRIGHT):
    """items: list of str or (label, text). Rendered as real bullet paragraphs."""
    from pptx.oxml.ns import qn
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = None
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, Inches(0.02))
    for i, it in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(spacing)
        pPr = para._p.get_or_add_pPr()
        pPr.set('marL', str(int(Inches(0.22)))); pPr.set('indent', str(-int(Inches(0.2))))
        bu_clr = pPr.makeelement(qn('a:buClr'), {})
        srgb = bu_clr.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % tuple(marker)})
        bu_clr.append(srgb); pPr.append(bu_clr)
        bu_font = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}); pPr.append(bu_font)
        bu = pPr.makeelement(qn('a:buChar'), {'char': '•'}); pPr.append(bu)
        if isinstance(it, tuple):
            lab, body = it
            r = para.add_run(); r.text = lab + '  '
            r.font.bold = True; r.font.color.rgb = label_color; r.font.size = Pt(size); r.font.name = FONT
            r = para.add_run(); r.text = body
            r.font.color.rgb = color; r.font.size = Pt(size); r.font.name = FONT
        else:
            r = para.add_run(); r.text = it
            r.font.color.rgb = color; r.font.size = Pt(size); r.font.name = FONT
    return tb


def picture(slide, path, x, y, w=None, h=None):
    if w is not None:
        return slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    return slide.shapes.add_picture(path, Inches(x), Inches(y), height=Inches(h))


_img_cache = {}


def jpeg(path, max_w=1600, crop=None):
    """Screens re-encoded as JPEG so the decks stay small; returns a file-like object."""
    key = (path, crop)
    if key not in _img_cache:
        im = Image.open(path).convert('RGB')
        if crop:
            im = im.crop(crop)
        if im.width > max_w:
            im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
        buf = io.BytesIO(); im.save(buf, 'JPEG', quality=85, optimize=True)
        _img_cache[key] = (buf.getvalue(), im.size)
    data, size = _img_cache[key]
    return io.BytesIO(data), size


# ------------------------------------------------------------------ deck
class Deck:
    def __init__(self, persona):
        self.p = persona
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW); self.prs.slide_height = Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def slide(self, title=None, kicker=None, chrome=True):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = WHITE
        if chrome:
            rect(s, 0, 0, SW, 0.08, fill=DEEP)
            # footer
            picture(s, LOGO, 0.5, 7.03, h=0.3)
            text(s, 0.5 + 0.3 * LOGO_RATIO + 0.2, 7.04, 7.5, 0.3, f"BrokerVerse | {self.p['label']}",
                 size=10, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
            text(s, SW - 1.5, 7.04, 1.0, 0.3, str(self.n), size=10, color=MUTED, align=PP_ALIGN.RIGHT,
                 anchor=MSO_ANCHOR.MIDDLE)
        if kicker:
            text(s, 0.5, 0.3, 12.3, 0.3, kicker.upper(), size=11, bold=True, color=BRIGHT)
        if title:
            text(s, 0.5, 0.55, 12.3, 0.65, title, size=28, bold=True, color=DEEP, anchor=MSO_ANCHOR.MIDDLE)
        return s

    def save(self, path):
        self.prs.save(path)


def card(s, x, y, w, h, fill=LIGHT, line=LINE):
    return rect(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)


def chip(s, x, y, label, fill=PALE_BLUE, color=DEEP, size=11, h=0.32, pad=0.18):
    w = max(0.6, len(label) * size * 0.0078 + pad * 2)
    rect(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x, y, w, h, label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def numbered_dot(s, x, y, n, fill=BRIGHT, d=0.36, size=12):
    rect(s, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    text(s, x, y, d, d, str(n), size=size, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ------------------------------------------------------------------ slides
def title_slide(d):
    p = d.p
    s = d.slide(chrome=False)
    rect(s, 0, 0, SW, 0.08, fill=DEEP)
    picture(s, LOGO, 0.8, 0.9, h=1.05)
    text(s, 0.8, 2.55, 7.9, 1.5, f"BrokerVerse — {p['label']} guide", size=38, bold=True, color=DEEP,
         anchor=MSO_ANCHOR.BOTTOM)
    text(s, 0.8, 4.15, 7.9, 0.5, 'Out-of-the-box functions, flows, reports and schedules', size=20, color=TEXT)
    text(s, 0.8, 4.75, 7.9, 0.4, 'BrokerVerse OOTB · insurance broking · Philippines · PHP', size=14, color=MUTED)
    x = 0.8
    for c in p['title_chips']:
        x += chip(s, x, 5.45, c) + 0.15
    # persona panel
    rect(s, 9.2, 0.9, 3.4, 5.6, fill=DEEP, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    text(s, 9.5, 1.2, 2.8, 0.3, 'THIS GUIDE IS FOR', size=11, bold=True, color=RGBColor(0xBF, 0xD9, 0xF5))
    text(s, 9.5, 1.55, 2.8, 1.0, p['label'], size=22, bold=True, color=WHITE)
    rows = [('Sign-in', p['users']), ('Role code', p['role_code']), ('Lands on', p['landing'])]
    y = 2.75
    for k, v in rows:
        text(s, 9.5, y, 2.8, 0.28, k.upper(), size=10, bold=True, color=RGBColor(0x9F, 0xC8, 0xF2))
        text(s, 9.5, y + 0.28, 2.8, 0.75, v, size=13, color=WHITE)
        y += 1.1
    text(s, 0.8, 6.85, 8, 0.3, 'Prepared by iorta TechNXT · September 2026', size=10, color=MUTED)


def role_slide(d):
    p = d.p
    s = d.slide('Role at a glance', 'Who you are in BrokerVerse')
    card(s, 0.5, 1.45, 7.3, 5.3, fill=WHITE)
    text(s, 0.8, 1.65, 6.8, 0.35, 'What you do', size=16, bold=True, color=DEEP)
    bullets(s, 0.8, 2.15, 6.8, 4.5, p['does'], size=16, spacing=12)
    boxes = [('Sign-in and landing page', p['signin']), ('You hand work to', p['hands_to']),
             ('Maker-checker partners', p['partners'])]
    y = 1.45
    hgt = [1.6, 1.75, 1.75]
    for (hdr, body), hh in zip(boxes, hgt):
        card(s, 8.1, y, 4.73, hh)
        text(s, 8.35, y + 0.15, 4.3, 0.3, hdr, size=13, bold=True, color=BRIGHT)
        text(s, 8.35, y + 0.5, 4.3, hh - 0.6, body, size=12, color=TEXT)
        y += hh + 0.1


def _need(items, w, pt, spacing=2):
    chars = max(1, int((w - 0.3) * 72 / (pt * 0.45)))
    lines = sum(1 + (len(i if isinstance(i, str) else i[0] + '  ' + i[1]) - 1) // chars for i in items)
    return lines * pt * 1.22 / 72 + len(items) * spacing / 72


def fit(items, w, h, spacing=6, sizes=(16, 15, 14, 13, 12.5, 12, 11.5, 11)):
    """Largest size at which a bullet list fits a box."""
    return next((pt for pt in sizes if _need(items, w, pt, spacing=spacing) <= h), sizes[-1])


def _fit_size(cards, w, h, sizes=(14, 13, 12, 11, 10, 9, 8)):
    """Largest font size at which every card's lines fit the height left under its header."""
    for pt in sizes:
        if all(0.62 + _need(items, w, pt) <= h for _, items in cards):
            return pt
    return sizes[-1]


def menu_slide(d):
    p = d.p
    s = d.slide('Menu map', 'What you see after sign-in')
    menus = p['menu']
    text(s, 0.5, 1.3, 12.3, 0.35, p.get('menu_note', 'Menus and screens granted to this role (menuPermissions); '
         'anything else is hidden and blocked by the route guard and the API.'), size=12, color=MUTED, italic=True)
    n = len(menus)
    nrows = 1 if n <= 4 else 2
    per_row = n if nrows == 1 else (n + 1) // 2
    gap = 0.2
    top = 1.8
    total_h = 6.8 - top
    cw = (12.33 - gap * (per_row - 1)) / per_row
    rows = [menus[i:i + per_row] for i in range(0, n, per_row)]
    pt = 8
    for cand in (15, 14, 13, 12, 11, 10, 9, 8):
        needs = [max(0.62 + _need(items, cw, cand) for _, items in r) for r in rows]
        if sum(needs) + gap * (len(rows) - 1) <= total_h:
            pt = cand
            break
    needs = [max(0.62 + _need(items, cw, pt) for _, items in r) for r in rows]
    extra = total_h - sum(needs) - gap * (len(rows) - 1)
    if nrows == 1:
        heights = [max(2.2, min(total_h, needs[0] + 0.4))]
    else:
        heights = [nd + extra / len(rows) for nd in needs]
    y = top
    for r, hh in zip(rows, heights):
        for c, (name, items) in enumerate(r):
            x = 0.5 + c * (cw + gap)
            card(s, x, y, cw, hh, fill=LIGHT)
            rect(s, x, y, cw, 0.42, fill=DEEP, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
            text(s, x + 0.15, y, cw - 0.3, 0.42, name, size=max(13, pt + 1), bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
            bullets(s, x + 0.12, y + 0.52, cw - 0.22, hh - 0.6, items, size=pt, spacing=2)
        y += hh + gap


def function_slide(d, i, total, f, shots):
    s = d.slide(f['title'], f'Key function {i} of {total}')
    x, y = 0.5, 1.4
    maxw, maxh = 7.95, 4.95
    if f.get('shot'):
        img = f['shot']
        path = shot_path(img, shots)
        data, (iw, ih) = jpeg(path, crop=f.get('crop'))
        w = maxw; h = w * ih / iw
        if h > maxh:
            h = maxh; w = h * iw / ih
        rect(s, x - 0.04, y - 0.04, w + 0.08, h + 0.08, fill=WHITE, line=LINE, lw=1)
        pic = s.shapes.add_picture(data, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        pic.line.color.rgb = LINE; pic.line.width = Pt(0.5)
    else:
        # step-by-step panel where no screen is shown
        w, h = maxw, maxh
        card(s, x, y, w, h, fill=WHITE)
        steps = f['steps']
        text(s, x + 0.35, y + 0.2, w - 0.7, 0.35, 'Step by step', size=14, bold=True, color=DEEP)
        gap = min(0.72, (h - 0.9) / len(steps))
        for k, st in enumerate(steps):
            yy = y + 0.75 + k * gap
            numbered_dot(s, x + 0.45, yy, k + 1, fill=GREEN if k == len(steps) - 1 else BRIGHT, d=0.42, size=13)
            if k < len(steps) - 1:
                rect(s, x + 0.65, yy + 0.44, 0.02, gap - 0.46, fill=LINE)
            text(s, x + 1.1, yy, w - 1.5, 0.42, st, size=15, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, y + h + 0.1, maxw, 0.3, f['caption'], size=10, italic=True, color=MUTED)
    cx = 8.75; cw = SW - 0.5 - cx
    card(s, cx, 1.36, cw, 5.43, fill=LIGHT)
    text(s, cx + 0.25, 1.52, cw - 0.5, 0.3, f.get('menu', ''), size=11, bold=True, color=BRIGHT)
    pts = f['points']
    size = next((pt for pt in (15, 14, 13, 12.5, 12, 11.5)
                 if _need(pts, cw - 0.45, pt, spacing=10) <= 4.55), 11.5)
    bullets(s, cx + 0.25, 1.95, cw - 0.45, 4.75, pts, size=size, spacing=10)


def flow_slide(d):
    p = d.p
    steps = p.get('flow_steps', FLOW_STEPS)
    mine = set(p['flow_mine'])
    s = d.slide(p.get('flow_title', 'End-to-end flow: your part of the cycle'), p.get('flow_kicker', 'Lead to renewal'))
    n = len(steps)
    first = (n + 1) // 2
    bw, bh, ag = 1.62, 1.12, 0.46
    row_w = first * bw + (first - 1) * ag
    x0 = 0.5 + (12.33 - row_w) / 2
    ys = [1.55, 3.35]
    pos = []
    for i in range(n):
        if i < first:
            pos.append((x0 + i * (bw + ag), ys[0]))
        else:
            j = i - first
            pos.append((x0 + (first - 1 - j) * (bw + ag), ys[1]))
    for i, ((name, owner), (x, y)) in enumerate(zip(steps, pos)):
        hl = (i + 1) in mine
        fill = GREEN if hl else LIGHT
        rect(s, x, y, bw, bh, fill=fill, line=None if hl else LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        text(s, x + 0.08, y + 0.08, 0.5, 0.25, str(i + 1), size=10, bold=True, color=WHITE if hl else BRIGHT)
        text(s, x + 0.08, y + 0.28, bw - 0.16, 0.45, name, size=13, bold=True, color=WHITE if hl else DEEP,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.08, y + 0.72, bw - 0.16, 0.34, owner, size=9.5, color=WHITE if hl else MUTED,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        # arrows
        if i < n - 1:
            nx, ny = pos[i + 1]
            if ny == y and nx > x:
                rect(s, x + bw + 0.08, y + bh / 2 - 0.13, ag - 0.16, 0.26, fill=BRIGHT, shape=MSO_SHAPE.RIGHT_ARROW)
            elif ny == y and nx < x:
                rect(s, nx + bw + 0.08, y + bh / 2 - 0.13, ag - 0.16, 0.26, fill=BRIGHT, shape=MSO_SHAPE.LEFT_ARROW)
            else:
                rect(s, x + bw / 2 - 0.15, y + bh + 0.12, 0.3, ys[1] - y - bh - 0.24, fill=BRIGHT,
                     shape=MSO_SHAPE.DOWN_ARROW)
    # legend
    ly = 4.72
    rect(s, 0.5, ly + 0.05, 0.28, 0.2, fill=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    text(s, 0.85, ly, 3.2, 0.3, 'Your steps', size=11, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 2.2, ly + 0.05, 0.28, 0.2, fill=LIGHT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    text(s, 2.55, ly, 4.5, 0.3, 'Done by other personas or by the system', size=11, color=TEXT,
         anchor=MSO_ANCHOR.MIDDLE)
    if p.get('flow_loop'):
        text(s, 7.0, ly, 5.83, 0.3, p['flow_loop'], size=11, italic=True, color=MUTED, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE)
    card(s, 0.5, 5.15, 12.33, 1.65, fill=PALE_GREEN, line=None)
    notes = p['flow_notes']
    half = (len(notes) + 1) // 2
    fs = min(fit(notes[:half], 5.85, 1.35, 5, (15, 14, 13, 12, 11)), fit(notes[half:], 5.85, 1.35, 5, (15, 14, 13, 12, 11)))
    bullets(s, 0.75, 5.3, 5.85, 1.45, notes[:half], size=fs, spacing=5, marker=GREEN)
    bullets(s, 6.75, 5.3, 5.9, 1.45, notes[half:], size=fs, spacing=5, marker=GREEN)


def _table(s, x, y, w, rows, widths, size=11, header_fill=DEEP, row_h=0.3):
    shape = s.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shape.table
    tot = sum(widths)
    for j, cw in enumerate(widths):
        t.columns[j].width = Inches(w * cw / tot)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            para = tf.paragraphs[0]
            r = para.add_run(); r.text = val
            r.font.name = FONT; r.font.size = Pt(size)
            if i == 0:
                r.font.bold = True; r.font.color.rgb = WHITE
                c.fill.solid(); c.fill.fore_color.rgb = header_fill
            else:
                r.font.color.rgb = TEXT
                if j == 0:
                    r.font.bold = True; r.font.color.rgb = DEEP
                c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
    return shape


def reports_slide(d):
    p = d.p
    codes = p['reports']
    s = d.slide(p.get('reports_title', 'Reports you use'), 'Reports > All Reports' if codes else 'Figures and evidence')
    x = 0.5
    if not codes:
        text(s, 0.5, 1.33, 12.33, 0.32, p.get('reports_note', ''), size=12, italic=True, color=MUTED,
             anchor=MSO_ANCHOR.MIDDLE)
        rows = [('Where', 'What it gives you')] + p['reports_alt']
        _table(s, 0.5, 1.95, 12.33, rows, [3, 9], size=13, row_h=0.5)
        return
    text(s, 0.5, 1.33, 1.6, 0.32, 'Output formats', size=11, bold=True, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    x = 2.05
    for c, f in (('CSV', PALE_BLUE), ('XLSX', PALE_BLUE), ('PDF', PALE_BLUE), ('On screen', PALE_BLUE),
                 ('Scheduled e-mail', PALE_GREEN)):
        x += chip(s, x, 1.33, c, fill=f, color=DEEP if f == PALE_BLUE else GREEN, size=10, h=0.3) + 0.1
    text(s, x + 0.2, 1.33, SW - 0.5 - x - 0.2, 0.32, p.get('reports_note', ''), size=11, italic=True, color=MUTED,
         anchor=MSO_ANCHOR.MIDDLE)
    if not codes:
        alt = p['reports_alt']
        rows = [('Where', 'What it gives you')] + alt
        _table(s, 0.5, 1.95, 12.33, rows, [3, 9], size=12, row_h=0.46)
        return
    items = [REPORTS[c] for c in codes]
    if len(items) <= 10:
        rows = [('Report', 'What it shows', 'Where')] + [(r[0], r[1], r[2]) for r in items]
        rh = min(0.5, 4.5 / len(rows))
        _table(s, 0.5, 1.95, 12.33, rows, [2.6, 6.6, 3.1], size=11.5 if len(rows) < 8 else 11, row_h=rh)
    else:
        ops = [r for c, r in zip(codes, items) if r[3] == 'operational']
        fin = [r for c, r in zip(codes, items) if r[3] == 'financial']
        for k, (hdr, group) in enumerate((('Operational', ops), ('Financial', fin))):
            if not group:
                continue
            rows = [(f'{hdr} report', 'What it shows')] + [(r[0], r[1]) for r in group]
            xx = 0.5 + k * 6.27
            rh = 0.4
            _table(s, xx, 1.95, 6.06, rows, [2.1, 3.96], size=9.5, row_h=rh)


def schedules_slide(d):
    p = d.p
    s = d.slide('Schedules and automations', 'Master > Configuration > Schedules')
    rows = [('Job', 'When it runs', 'What it means for you')]
    for code, effect in p['jobs']:
        name, when, _ = JOBS[code]
        rows.append((name, when, effect))
    rh = min(0.55, 4.5 / len(rows))
    _table(s, 0.5, 1.45, 12.33, rows, [2.4, 2.2, 7.7], size=12 if len(rows) <= 7 else 11, row_h=rh)
    text(s, 0.5, min(6.3, 1.45 + rh * len(rows) + 0.35), 12.33, 0.45, p.get('jobs_note', 'Times are server time (Philippine time in production). '
         'The IT administrator can change a timetable, switch a job off or press Run now; every run is recorded.'),
         size=11, italic=True, color=MUTED)


def notifications_slide(d):
    p = d.p
    s = d.slide('Notifications and approvals', 'Bell icon, top right')
    items = p['notifications']
    n = len(items)
    cols = 2
    per = (n + 1) // 2
    gap = 0.15
    ch = min(1.05, (5.35 - gap * (per - 1)) / per)
    cw = (12.33 - 0.3) / 2
    kind_color = {'Approval': (PALE_GREEN, GREEN), 'Task': (PALE_BLUE, DEEP), 'Info': (LIGHT, BRIGHT),
                  'Reminder': (PALE_BLUE, BRIGHT), 'Alert': (RGBColor(0xFD, 0xEC, 0xEA), RGBColor(0xC6, 0x28, 0x28)),
                  'E-mail': (PALE_GREEN, GREEN)}
    for i, (kind, title, body) in enumerate(items):
        c, r = divmod(i, per)
        x = 0.5 + c * (cw + 0.3); y = 1.45 + r * (ch + gap)
        card(s, x, y, cw, ch, fill=WHITE)
        f, col = kind_color.get(kind, (LIGHT, DEEP))
        rect(s, x + 0.18, y + 0.18, 1.1, 0.3, fill=f, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(s, x + 0.18, y + 0.18, 1.1, 0.3, kind, size=10, bold=True, color=col, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.45, y + 0.12, cw - 1.6, 0.36, title, size=12.5, bold=True, color=DEEP, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.45, y + 0.46, cw - 1.6, ch - 0.5, body, size=11, color=TEXT)


def limits_slide(d):
    p = d.p
    s = d.slide('Known limitations and tips', 'Good to know')
    card(s, 0.5, 1.45, 6.05, 5.3, fill=LIGHT)
    text(s, 0.8, 1.6, 5.5, 0.35, 'Open items that affect you', size=15, bold=True, color=DEEP)
    bullets(s, 0.8, 2.05, 5.55, 4.6, p['limits'], size=fit(p['limits'], 5.55, 4.45, 8, (15, 14, 13, 12, 11.5, 11)), spacing=8, marker=MUTED)
    card(s, 6.78, 1.45, 6.05, 5.3, fill=PALE_GREEN, line=None)
    text(s, 7.08, 1.6, 5.5, 0.35, 'Tips', size=15, bold=True, color=GREEN)
    bullets(s, 7.08, 2.05, 5.55, 4.6, p['tips'], size=fit(p['tips'], 5.55, 4.45, 8, (15, 14, 13, 12, 11.5, 11)), spacing=8, marker=GREEN)


def closing_slide(d):
    p = d.p
    s = d.slide(chrome=False)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = DEEP
    rect(s, 0.8, 0.9, 4.2, 1.5, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    picture(s, LOGO, 1.05, 1.15, h=1.0)
    text(s, 0.8, 2.9, 11, 0.8, 'Thank you', size=40, bold=True, color=WHITE)
    text(s, 0.8, 3.7, 11, 0.5, f"BrokerVerse — {p['label']} guide", size=18, color=RGBColor(0xBF, 0xD9, 0xF5))
    rect(s, 0.8, 4.55, 11.7, 2.05, fill=RGBColor(0x0E, 0x5C, 0xB2), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(s, 1.1, 4.7, 5, 0.35, 'SUPPORT', size=11, bold=True, color=RGBColor(0x9F, 0xC8, 0xF2))
    rows = [('Service desk', '<service desk to be confirmed>'), ('E-mail', '<support mailbox to be confirmed>'),
            ('Hours', '<business hours, Manila time, to be confirmed>'),
            ('Access requests', p.get('access_contact', 'Your User Access Administrator (Master > User Management)'))]
    y = 5.05
    for k, v in rows:
        text(s, 1.1, y, 2.4, 0.33, k, size=13, bold=True, color=WHITE)
        text(s, 3.5, y, 8.8, 0.33, v, size=13, color=WHITE)
        y += 0.36
    text(s, 0.8, 6.9, 11.7, 0.3, 'iorta TechNXT · BrokerVerse insurance-broking platform', size=10,
         color=RGBColor(0xBF, 0xD9, 0xF5))


def extra_slide(d, e, shots):
    kind = e['kind']
    if kind == 'api':
        s = d.slide(e['title'], e['kicker'])
        stats = e['stats']
        cw = (12.33 - 0.2 * (len(stats) - 1)) / len(stats)
        for i, (big, small) in enumerate(stats):
            x = 0.5 + i * (cw + 0.2)
            card(s, x, 1.45, cw, 1.35, fill=PALE_BLUE, line=None)
            text(s, x + 0.2, 1.52, cw - 0.4, 0.7, big, size=30, bold=True, color=DEEP)
            text(s, x + 0.2, 2.2, cw - 0.4, 0.5, small, size=11, color=TEXT)
        card(s, 0.5, 3.0, 6.05, 3.75, fill=LIGHT)
        text(s, 0.75, 3.12, 5.6, 0.35, e['left_title'], size=14, bold=True, color=DEEP)
        bullets(s, 0.75, 3.52, 5.6, 3.2, e['left'], size=fit(e['left'], 5.6, 3.0, 6, (14, 13, 12, 11.5, 11)), spacing=6)
        card(s, 6.78, 3.0, 6.05, 3.75, fill=LIGHT)
        text(s, 7.03, 3.12, 5.6, 0.35, e['right_title'], size=14, bold=True, color=DEEP)
        bullets(s, 7.03, 3.52, 5.6, 3.2, e['right'], size=fit(e['right'], 5.6, 3.0, 6, (14, 13, 12, 11.5, 11)), spacing=6)
    elif kind == 'twocol':
        s = d.slide(e['title'], e['kicker'])
        card(s, 0.5, 1.45, 6.05, 5.3, fill=RGBColor(0xFD, 0xEC, 0xEA) if e.get('left_warn') else LIGHT, line=None)
        text(s, 0.8, 1.6, 5.5, 0.35, e['left_title'], size=15, bold=True, color=DEEP)
        bullets(s, 0.8, 2.05, 5.55, 4.6, e['left'], size=fit(e['left'], 5.55, 4.45, 8), spacing=8)
        if e.get('shot'):
            path = shot_path(e['shot'], shots)
            data, (iw, ih) = jpeg(path)
            w = 6.05; h = w * ih / iw
            if h > 4.7:
                h = 4.7; w = h * iw / ih
            x = 6.78 + (6.05 - w) / 2
            rect(s, x - 0.04, 1.41, w + 0.08, h + 0.08, fill=WHITE, line=LINE, lw=1)
            s.shapes.add_picture(data, Inches(x), Inches(1.45), width=Inches(w), height=Inches(h))
            text(s, 6.78, 1.5 + h + 0.08, 6.05, 0.3, e['caption'], size=10, italic=True, color=MUTED)
        else:
            card(s, 6.78, 1.45, 6.05, 5.3, fill=PALE_GREEN, line=None)
            text(s, 7.08, 1.6, 5.5, 0.35, e['right_title'], size=15, bold=True, color=GREEN)
            bullets(s, 7.08, 2.05, 5.55, 4.6, e['right'], size=fit(e['right'], 5.55, 4.45, 8), spacing=8, marker=GREEN)


def build(p, shots, out):
    d = Deck(p)
    title_slide(d)
    role_slide(d)
    menu_slide(d)
    fs = p['functions']
    for i, f in enumerate(fs, 1):
        function_slide(d, i, len(fs), f, shots)
    for e in p.get('extras', []):
        extra_slide(d, e, shots)
    flow_slide(d)
    reports_slide(d)
    schedules_slide(d)
    notifications_slide(d)
    limits_slide(d)
    closing_slide(d)
    path = os.path.join(out, p['file'])
    d.save(path)
    return path, d.n


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--shots', required=True)
    ap.add_argument('--out', default=os.path.join(ROOT, 'docs', 'decks'))
    ap.add_argument('only', nargs='*')
    a = ap.parse_args()
    for p in PERSONAS:
        if a.only and p['user'] not in a.only:
            continue
        path, n = build(p, a.shots, a.out)
        print(f'{os.path.basename(path)}: {n} slides')
