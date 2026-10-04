"""Scan the text of the manual and the role decks for wording the house style does not allow.

    python3 docs/manual/tools/style_scan.py [file or folder ...]

Default targets: docs/manual/tools/manual_source.md, docs/manual/BrokerVerse_User_Manual.docx and .pdf,
docs/decks/tools/personas.py and every docs/decks/*.pptx. Text is read from .md/.py/.txt directly, from .docx with
python-docx, from .pdf with pdftotext and from .pptx with python-pptx.

It reports, with the file and a short extract:
  - stock words (seamless, robust, comprehensive, leverage, delve, empower, cutting-edge, streamline, unlock,
    elevate, holistic, synergy). "Comprehensive" as the motor policy type and "unlock" of a user, account or period
    are screen terms and are not reported;
  - emojis and pictographs;
  - long dashes: an em dash anywhere, an en dash with spaces around it (a dash used to join two sentences);
  - names that must not appear (Claude, Anthropic, AI as a word, BIBS, BDOI) and BDO outside a bank reference;
  - lines of the Markdown source with more than two bold phrases.
Exit code 1 when anything is reported.
"""
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
DEFAULTS = [
    'docs/manual/tools/manual_source.md', 'docs/manual/BrokerVerse_User_Manual.docx',
    'docs/manual/BrokerVerse_User_Manual.pdf', 'docs/decks/tools/personas.py', 'docs/decks',
]

STOCK = re.compile(r'\b(seamless\w*|robust\w*|comprehensive\w*|leverag\w*|delv(?:e|es|ed|ing)|empower\w*|cutting[- ]edge|'
                   r'streamlin\w*|unlock\w*|elevat\w*|holistic\w*|synerg\w*)\b', re.I)
ALLOWED = [
    re.compile(r'\bComprehensive\b'),                                         # motor policy type / plan tier on screen
    re.compile(r'\bunlock\w*\b(?:\W+\w+){0,4}?\W+(user|users|account|accounts|period|periods)\b', re.I),
    re.compile(r'^Unlock$|\bUnlock\b(?= \(|,| the | a |\s*\|)'),              # the Unlock action on the user list
]
EMOJI = re.compile('[\U0001F000-\U0001FAFF\U00002600-\U000027BF\U0001F900-\U0001F9FF⭐⬆⬇⤴⤵'
                   '⌚⌛⏩-⏳▪▫▶◀◻-◾️]')
EM_DASH = re.compile('—')
EN_DASH_JOIN = re.compile(r'\s–\s')
NAMES = re.compile(r'\b(Claude|Anthropic|BIBS|BDOI|BDO Insure)\b|\bAI\b')
BDO = re.compile(r'\bBDO\b')
BDO_OK = re.compile(r'BDO[- ](SAMPLE|001)|ACC-BDO|BDO \*|Banco de Oro|BDO, BPI|BDO and BPI|BDO, BPI and Metrobank|\| BDO \||BDO online', re.I)
BOLD = re.compile(r'\*\*[^*]+\*\*')


def text_of(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.md', '.py', '.txt'):
        return open(path, encoding='utf-8').read()
    if ext == '.docx':
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(' | '.join(c.text for c in row.cells))
        return '\n'.join(parts)
    if ext == '.pdf':
        return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True).stdout
    if ext == '.pptx':
        from pptx import Presentation
        parts = []
        for i, s in enumerate(Presentation(path).slides, 1):
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, 'has_table', False) and sh.has_table:
                    for row in sh.table.rows:
                        parts.append(' | '.join(c.text for c in row.cells))
            if s.has_notes_slide:
                parts.append(s.notes_slide.notes_text_frame.text)
        return '\n'.join(parts)
    return ''


def targets(args):
    out = []
    for a in args or DEFAULTS:
        p = a if os.path.isabs(a) else os.path.join(REPO, a)
        if os.path.isdir(p):
            out += sorted(os.path.join(p, f) for f in os.listdir(p) if f.endswith(('.pptx', '.docx', '.pdf', '.md')))
        elif os.path.exists(p):
            out.append(p)
    return out


def allowed(line, m):
    """True when the stock word sits inside one of the allowed screen terms."""
    return any(am.start() <= m.start() < am.end() for a in ALLOWED for am in a.finditer(line))


def scan(path):
    found = []
    text = text_of(path)
    for n, line in enumerate(text.splitlines(), 1):
        for m in STOCK.finditer(line):
            if not allowed(line, m):
                found.append((n, 'stock word', m.group(0), line))
        for rx, kind in ((EMOJI, 'emoji'), (EM_DASH, 'em dash'), (EN_DASH_JOIN, 'en dash joining text'), (NAMES, 'name')):
            for m in rx.finditer(line):
                found.append((n, kind, m.group(0), line))
        for m in BDO.finditer(line):
            if not BDO_OK.search(line[max(0, m.start() - 14):m.end() + 40]):
                found.append((n, 'BDO outside a bank reference', m.group(0), line))
        if path.endswith('.md') and len(BOLD.findall(line)) > 2:
            found.append((n, 'more than two bold phrases', '', line))
    return found


def main():
    total = 0
    for path in targets(sys.argv[1:]):
        found = scan(path)
        total += len(found)
        rel = os.path.relpath(path, REPO)
        print(f'{rel}: {len(found)} finding(s)')
        for n, kind, what, line in found[:200]:
            snippet = line.strip()
            print(f'  line {n}: {kind}: "{what}" in: {snippet[:140]}')
    print(f'Total findings: {total}')
    sys.exit(1 if total else 0)


if __name__ == '__main__':
    main()
