"""Build BrokerVerse_User_Manual.docx and .pdf from manual_source.md and the screenshots in ../images.

    python3 docs/manual/tools/build_manual.py            # docx (python-docx) -> LibreOffice: TOC, page numbers, PDF

Source format (manual_source.md), one block per blank-line-separated paragraph:
    # Chapter            (Heading 1, new page)         ## Section / ### Subsection
    plain text           paragraph; **bold**, *italic*, `code`
    1. step              numbered steps (numbering restarts in every list)
    - item               bullets
    | a | b |            table; first row is the header; a |---| row is skipped
    {widths: 25,75}      column widths (%) of the next table
    ![Caption](name)     screenshot images/name.jpg with a numbered caption
    > **Tip:** text      call-out box (Tip / Note = blue, Important / Caution = yellow)
    \\pagebreak           page break
Brand: BrokerVerse by iorta TechNXT. Header Blue #004ea8, CTA Blue #0072d8, Background Blue #e5f5ff, accent Yellow
#fdb913, font Nunito; iorta TechNXT logo on the cover.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                       # docs/manual
REPO = os.path.dirname(os.path.dirname(ROOT))
IMAGES = os.path.join(ROOT, 'images')
SOURCE = os.path.join(HERE, 'manual_source.md')
OUT_DOCX = os.path.join(ROOT, 'BrokerVerse_User_Manual.docx')
OUT_PDF = os.path.join(ROOT, 'BrokerVerse_User_Manual.pdf')
LOGO_IORTA = os.path.join(REPO, 'brokerverse', 'public', 'bdoi', 'iorta-technxt.png')

PRODUCT = 'BrokerVerse'
TITLE = 'BrokerVerse User Manual'
TAGLINE = 'Insurance broking platform'
VERSION = 'Version 1.0'
DATE = '29 September 2026'
FONT = 'Nunito'
BLUE = RGBColor(0x00, 0x4E, 0xA8)
CTA = RGBColor(0x00, 0x72, 0xD8)
TEXT = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x5B, 0x67, 0x76)
HEX_BLUE, HEX_CTA, HEX_BG, HEX_YELLOW = '004EA8', '0072D8', 'E5F5FF', 'FDB913'
PAGE_W, MARGIN = Cm(21.0), Cm(2.0)
CONTENT_W = PAGE_W - 2 * MARGIN


# ------------------------------------------------------------------ low-level helpers
def set_font(run_or_style, size=None, bold=None, color=None, italic=None, name=FONT):
    f = run_or_style.font
    f.name = name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style, 'element') else None
    if rpr is not None:
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            fonts.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcpr.append(shd)


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('bottom', bottom), ('start', left), ('end', right), ('left', left), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcpr.append(mar)


def table_borders(table, color='C9D6E3', size=4, inside=True, left_only=None):
    tblpr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        if left_only:
            val, sz, col = ('single', left_only[1], left_only[0]) if edge == 'left' else ('nil', 0, 'auto')
        elif edge.startswith('inside') and not inside:
            val, sz, col = 'nil', 0, 'auto'
        else:
            val, sz, col = 'single', size, color
        e.set(qn('w:val'), val); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        borders.append(e)
    tblpr.append(borders)


def fixed_layout(table, widths):
    tblpr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblpr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths) / 635))); tw.set(qn('w:type'), 'dxa')
    for old in tblpr.findall(qn('w:tblW')):
        tblpr.remove(old)
    tblpr.append(tw)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[i] / 635)))
    for row in table.rows:
        for i, c in enumerate(row.cells):
            c.width = widths[i]


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trpr.append(e)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); e.set(qn('w:val'), 'true'); trpr.append(e)


def para_border(p, edge='bottom', color=HEX_YELLOW, size=12, space=4):
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(size)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color)
    pbdr.append(e); ppr.append(pbdr)


def keep_next(p):
    p.paragraph_format.keep_with_next = True


def add_field(par, instr, placeholder=''):
    r = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
    r = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {instr} '; r._r.append(it)
    r = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'separate'); r._r.append(fc)
    r = par.add_run(placeholder)
    r2 = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'end'); r2._r.append(fc)
    return r


INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)')


def add_inline(par, text, size=None, color=None, bold=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**'):
            r = par.add_run(part[2:-2]); set_font(r, size=size, bold=True, color=color)
        elif part.startswith('`'):
            r = par.add_run(part[1:-1]); set_font(r, size=(size or 10.5) - 1, color=RGBColor(0x0B, 0x3D, 0x7A), name='Liberation Mono')
        elif part.startswith('*') and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, size=size, italic=True, color=color, bold=bold)
        else:
            r = par.add_run(part); set_font(r, size=size, color=color, bold=bold)


# ------------------------------------------------------------------ styles
def setup_styles(doc):
    st = doc.styles
    normal = st['Normal']
    set_font(normal, size=10.5, color=TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'en-PH'); rpr.append(lang)
    for name, size, color, before, after in (('Heading 1', 21, BLUE, 0, 10), ('Heading 2', 14.5, BLUE, 14, 5),
                                             ('Heading 3', 12, CTA, 10, 3)):
        s = st[name]
        set_font(s, size=size, bold=True, color=color)
        s.font.italic = False
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.page_break_before = False
    for name in ('List Bullet',):
        s = st[name]; set_font(s, size=10.5, color=TEXT); s.paragraph_format.space_after = Pt(2)
    for lvl in (1, 2, 3):
        name = f'TOC {lvl}' if f'TOC {lvl}' in [x.name for x in st] else None
    cap = st.add_style('Figure Caption', 1)
    cap.base_style = normal
    set_font(cap, size=9, italic=True, color=GREY)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    step = st.add_style('Step', 1)
    step.base_style = normal
    step.paragraph_format.left_indent = Cm(0.75)
    step.paragraph_format.first_line_indent = Cm(-0.75)
    step.paragraph_format.space_after = Pt(3)
    step.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
    cell = st.add_style('Table Text', 1)
    cell.base_style = normal
    set_font(cell, size=9, color=TEXT)
    cell.paragraph_format.space_after = Pt(0)
    cell.paragraph_format.line_spacing = 1.05


# ------------------------------------------------------------------ cover, header / footer, TOC
def cover(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(30)
    if os.path.exists(LOGO_IORTA):
        p.add_run().add_picture(LOGO_IORTA, width=Cm(6.2))
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = band.rows[0].cells[0]
    shade(c, HEX_BLUE); cell_margins(c, 500, 500, 450, 450)
    table_borders(band, color=HEX_BLUE, size=4)
    fixed_layout(band, [CONTENT_W])
    p = c.paragraphs[0]
    r = p.add_run(PRODUCT); set_font(r, size=40, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    p.paragraph_format.space_after = Pt(2)
    p = c.add_paragraph(); r = p.add_run(TAGLINE); set_font(r, size=18, color=RGBColor(0xFF, 0xFF, 0xFF))
    p.paragraph_format.space_after = Pt(18)
    p = c.add_paragraph(); r = p.add_run('User Manual'); set_font(r, size=28, bold=True, color=RGBColor(0xFD, 0xB9, 0x13))
    p = c.add_paragraph(); r = p.add_run('Out-of-the-box insurance broking for the Philippines — end-to-end guide for every user role')
    set_font(r, size=12, color=RGBColor(0xE5, 0xF5, 0xFF))
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(26)
    rows = [('Product', 'BrokerVerse (out-of-the-box)'), ('Document', TITLE), ('Version', VERSION), ('Date', DATE),
            ('Audience', 'Business users, approvers and administrators of BrokerVerse'),
            ('Classification', 'Internal — BrokerVerse customers and partners')]
    t = doc.add_table(rows=len(rows), cols=2)
    fixed_layout(t, [Cm(4), CONTENT_W - Cm(4)])
    table_borders(t, color='C9D6E3', size=4)
    for i, (k, v) in enumerate(rows):
        a, b = t.rows[i].cells
        shade(a, HEX_BG)
        for cc in (a, b):
            cell_margins(cc, 70, 70, 120, 120)
        pa = a.paragraphs[0]; ra = pa.add_run(k); set_font(ra, size=10, bold=True, color=BLUE)
        pb = b.paragraphs[0]; rb = pb.add_run(v); set_font(rb, size=10, color=TEXT)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run('Powered by iorta TechNXT'); set_font(r, size=11, color=GREY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def header_footer(doc):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = ''
    tabs = hp.paragraph_format.tab_stops
    for pos in (Cm(8.255), Cm(16.51), Emu(2971800), Emu(5943600)):
        tabs.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
    tabs.add_tab_stop(CONTENT_W, WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(PRODUCT + ' — ' + TAGLINE); set_font(r, size=8.5, bold=True, color=BLUE)
    r = hp.add_run('\tUser Manual · ' + VERSION); set_font(r, size=8.5, color=GREY)
    para_border(hp, 'bottom', HEX_YELLOW, 8, 4)
    fp = sec.footer.paragraphs[0]
    for pos in (Cm(8.255), Cm(16.51), Emu(2971800), Emu(5943600)):
        fp.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
    fp.paragraph_format.tab_stops.add_tab_stop(CONTENT_W, WD_TAB_ALIGNMENT.RIGHT)
    para_border(fp, 'top', 'C9D6E3', 4, 4)
    r = fp.add_run(DATE + '  ·  Powered by iorta TechNXT'); set_font(r, size=8, color=GREY)
    r = fp.add_run('\tPage '); set_font(r, size=8.5, color=GREY)
    set_font(add_field(fp, 'PAGE', '1'), size=8.5, bold=True, color=BLUE)
    r = fp.add_run(' of '); set_font(r, size=8.5, color=GREY)
    set_font(add_field(fp, 'NUMPAGES', '1'), size=8.5, color=GREY)
    for p in fp.runs:
        set_font(p, size=8.5)


def toc(doc):
    h = doc.add_paragraph()
    r = h.add_run('Contents'); set_font(r, size=21, bold=True, color=BLUE)
    h.paragraph_format.space_after = Pt(10)
    para_border(h, 'bottom', HEX_YELLOW, 12, 6)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', 'Right-click and choose Update Field to build the table of contents.')


# ------------------------------------------------------------------ blocks
class Builder:
    def __init__(self, doc):
        self.doc = doc
        self.fig = 0
        self.widths = None
        self.first_h1 = True
        self.missing = []
        self.ch = 0
        self.sec = 0
        self.number_h2 = True

    def heading(self, level, text):
        if level == 1:
            self.ch += 1; self.sec = 0
            self.number_h2 = not text.startswith('Appendices')
            text = f'{self.ch}  {text}'
        elif level == 2 and self.number_h2:
            self.sec += 1
            text = f'{self.ch}.{self.sec}  {text}'
        p = self.doc.add_paragraph(style=f'Heading {level}')
        add_inline(p, text)
        if level == 1:
            p.paragraph_format.page_break_before = True
            para_border(p, 'bottom', HEX_YELLOW, 12, 6)

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        add_inline(p, text)

    def steps(self, items):
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph(style='Step')
            r = p.add_run(f'{i}.\t'); set_font(r, bold=True, color=CTA)
            add_inline(p, it)
            if i < len(items):
                keep_next(p) if len(items) <= 6 else None

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style='List Bullet')
            add_inline(p, it)

    def table(self, rows):
        ncols = max(len(r) for r in rows)
        rows = [r + [''] * (ncols - len(r)) for r in rows]
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        pct = self.widths or [100 / ncols] * ncols
        total = sum(pct)
        widths = [Emu(int(CONTENT_W * w / total)) for w in pct]
        fixed_layout(t, widths)
        table_borders(t)
        for i, row in enumerate(rows):
            tr = t.rows[i]
            no_split(tr)
            for j, val in enumerate(row):
                c = tr.cells[j]
                cell_margins(c, 50, 50, 90, 90)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if i == 0 else WD_CELL_VERTICAL_ALIGNMENT.TOP
                p = c.paragraphs[0]; p.style = self.doc.styles['Table Text']
                parts = val.split('<br>')
                for k, part in enumerate(parts):
                    if k:
                        p = c.add_paragraph(style='Table Text')
                    if i == 0:
                        add_inline(p, part.strip(), size=9, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                    else:
                        add_inline(p, part.strip(), size=9)
                if i == 0:
                    shade(c, HEX_BLUE)
                elif i % 2 == 0:
                    shade(c, 'F3F9FF')
            if i == 0:
                repeat_header(tr)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        self.widths = None

    def image(self, caption, name):
        path = os.path.join(IMAGES, name if name.endswith('.jpg') else name + '.jpg')
        if not os.path.exists(path):
            self.missing.append(name)
            self.paragraph(f'[Screenshot not available: {name}]')
            return
        from PIL import Image
        w, h = Image.open(path).size
        max_w, max_h = CONTENT_W, Cm(17.5)
        width = max_w
        if w and h and max_w * h / w > max_h:
            width = Emu(int(max_h * w / h))
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
        keep_next(p)
        run = p.add_run(); run.add_picture(path, width=width)
        # thin frame around the screenshot
        inline = run._r.find('.//' + qn('wp:inline'))
        pic = inline.find('.//' + qn('pic:spPr')) if inline is not None else None
        if pic is not None:
            ln = OxmlElement('a:ln'); ln.set('w', '6350')
            sf = OxmlElement('a:solidFill'); clr = OxmlElement('a:srgbClr'); clr.set('val', 'C9D6E3'); sf.append(clr); ln.append(sf)
            pic.append(ln)
        self.fig += 1
        cp = self.doc.add_paragraph(style='Figure Caption')
        r = cp.add_run(f'Figure {self.fig}. '); set_font(r, size=9, bold=True, italic=True, color=BLUE)
        add_inline(cp, caption, size=9, color=GREY)

    def callout(self, text):
        m = re.match(r'\*\*(\w[\w ]*):\*\*\s*(.*)', text)
        kind = m.group(1).lower() if m else 'note'
        warm = kind in ('important', 'caution', 'warning', 'known issue')
        t = self.doc.add_table(rows=1, cols=1)
        fixed_layout(t, [CONTENT_W])
        table_borders(t, left_only=(HEX_YELLOW if warm else HEX_CTA, 24))
        c = t.rows[0].cells[0]
        shade(c, 'FFF6DD' if warm else HEX_BG); cell_margins(c, 90, 90, 160, 140)
        parts = text.split('<br>')
        for k, part in enumerate(parts):
            p = c.paragraphs[0] if k == 0 else c.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, part.strip(), size=9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_and_build(b, text):
    blocks = re.split(r'\n\s*\n', text.strip())
    for block in blocks:
        lines = [l.rstrip() for l in block.split('\n') if l.strip() and not l.strip().startswith('<!--')]
        if not lines:
            continue
        first = lines[0].strip()
        if first.startswith('{widths:'):
            b.widths = [float(x) for x in first[8:-1].split(',')]
            lines = lines[1:]
            if not lines:
                continue
            first = lines[0].strip()
        if first == '\\pagebreak':
            b.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); continue
        m = re.match(r'^(#{1,3})\s+(.*)', first)
        if m:
            b.heading(len(m.group(1)), m.group(2)); lines = lines[1:]
            if not lines:
                continue
            first = lines[0].strip()
        if first.startswith('|'):
            rows = []
            for l in lines:
                cells = [c.strip() for c in l.strip().strip('|').split('|')]
                if all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    continue
                rows.append(cells)
            b.table(rows)
        elif re.match(r'^!\[', first):
            for l in lines:
                mm = re.match(r'!\[(.*)\]\((.*)\)', l.strip())
                if mm:
                    b.image(mm.group(1), mm.group(2))
        elif first.startswith('>'):
            b.callout('<br>'.join(l.strip()[1:].strip() for l in lines))
        elif re.match(r'^\d+\.\s', first):
            items = []
            for l in lines:
                if re.match(r'^\s*\d+\.\s', l):
                    items.append(re.sub(r'^\s*\d+\.\s+', '', l))
                else:
                    items[-1] += ' ' + l.strip()
            b.steps(items)
        elif first.startswith('- '):
            items = []
            for l in lines:
                if l.strip().startswith('- '):
                    items.append(l.strip()[2:])
                else:
                    items[-1] += ' ' + l.strip()
            b.bullets(items)
        else:
            b.paragraph(' '.join(l.strip() for l in lines))


def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = MARGIN
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
    setup_styles(doc)
    cp = doc.core_properties
    cp.title = TITLE
    cp.subject = 'BrokerVerse user manual — ' + TAGLINE; cp.author = 'iorta TechNXT'
    cp.keywords = 'BrokerVerse, iorta TechNXT, insurance broking, user manual'; cp.category = 'User manual'
    cover(doc)
    header_footer(doc)
    toc(doc)
    b = Builder(doc)
    parse_and_build(b, open(SOURCE, encoding='utf-8').read())
    doc.save(path)
    return b


# ------------------------------------------------------------------ LibreOffice: update the TOC and export
UNO_SCRIPT = r'''
import sys, time, uno
from com.sun.star.beans import PropertyValue
def P(n, v):
    p = PropertyValue(); p.Name = n; p.Value = v; return p
src, out_docx, out_pdf, pipe = sys.argv[1:5]
local = uno.getComponentContext()
resolver = local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', local)
for i in range(60):
    try:
        ctx = resolver.resolve(f'uno:pipe,name={pipe};urp;StarOffice.ComponentContext'); break
    except Exception:
        time.sleep(1)
desktop = ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', ctx)
doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(src), '_blank', 0, (P('Hidden', True),))
for _ in range(2):
    idx = doc.getDocumentIndexes()
    for i in range(idx.getCount()):
        idx.getByIndex(i).update()
    doc.getTextFields().refresh()
    doc.refresh()
doc.storeToURL(uno.systemPathToFileUrl(out_docx), (P('FilterName', 'MS Word 2007 XML'),))
fd = (P('Quality', 72), P('ReduceImageResolution', True), P('MaxImageResolution', 150),
      P('ExportBookmarks', True), P('UseTaggedPDF', True))
doc.storeToURL(uno.systemPathToFileUrl(out_pdf), (P('FilterName', 'writer_pdf_Export'), P('FilterData', uno.Any('[]com.sun.star.beans.PropertyValue', fd))))
doc.close(True)
try:
    desktop.terminate()
except Exception:
    pass
'''


def libreoffice(src, out_docx, out_pdf):
    work = tempfile.mkdtemp(prefix='bv-manual-lo-')
    pipe = f'bvmanual{os.getpid()}'
    profile = 'file://' + os.path.join(work, 'profile')
    office = subprocess.Popen(['soffice', f'-env:UserInstallation={profile}', '--headless', '--invisible', '--nologo',
                               '--norestore', f'--accept=pipe,name={pipe};urp;StarOffice.ComponentContext'],
                              stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        script = os.path.join(work, 'export.py')
        open(script, 'w').write(UNO_SCRIPT)
        subprocess.run([sys.executable, script, src, out_docx, out_pdf, pipe], check=True, timeout=900)
    finally:
        try:
            office.wait(timeout=30)
        except Exception:
            office.kill()
        shutil.rmtree(work, ignore_errors=True)


def main():
    tmp = tempfile.mkdtemp(prefix='bv-manual-')
    raw = os.path.join(tmp, 'manual_raw.docx')
    b = build_docx(raw)
    print(f'figures: {b.fig}; missing screenshots: {b.missing or "none"}')
    libreoffice(raw, OUT_DOCX, OUT_PDF)
    shutil.rmtree(tmp, ignore_errors=True)
    for f in (OUT_DOCX, OUT_PDF):
        print(f'{f}: {os.path.getsize(f) / 1e6:.1f} MB')


if __name__ == '__main__':
    main()
