"""Builds sales/iNXT_BrokerVerse_Price_Book.xlsx (formula driven) from the recommended price points of
out/BrokerVerse_Commercials_and_Pricing.xlsx.

    python3 build_price_book.py

The pricing workbook is recalculated with LibreOffice (headless) into a temporary folder and its values are
read by label, so the price book always carries the same price points as the pricing workbook.
"""
import os
import shutil
import subprocess
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
SRC = os.path.join(PKG, 'out', 'BrokerVerse_Commercials_and_Pricing.xlsx')
OUT = os.path.join(PKG, 'sales', 'iNXT_BrokerVerse_Price_Book.xlsx')


# ------------------------------------------------------------------ read the pricing workbook
def recalculated(path):
    tmp = tempfile.mkdtemp()
    prof = tempfile.mkdtemp()
    subprocess.run(['soffice', f'-env:UserInstallation=file://{prof}', '--headless', '--convert-to', 'xlsx',
                    '--outdir', tmp, path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    out = os.path.join(tmp, os.path.basename(path))
    wb = load_workbook(out, data_only=True)
    shutil.rmtree(tmp, ignore_errors=True)
    shutil.rmtree(prof, ignore_errors=True)
    return wb


src = recalculated(SRC)
inp = src['Inputs']


def by_label(ws, label, col=2):
    for row in ws.iter_rows():
        if row[0].value == label:
            return row[col - 1].value
    raise KeyError(label)


G = {
    'FX': by_label(inp, 'Exchange rate'),
    'FX_Date': by_label(inp, 'Exchange rate date'),
    'VAT': by_label(inp, 'VAT rate'),
    'AMC_Rate': by_label(inp, 'AMC rate (of perpetual licence fee)'),
    'Esc': by_label(inp, 'Annual escalation (AMC, subscription, support)'),
    'AMC_Start': by_label(inp, 'AMC starts in contract year'),
    'Onboard_Pct': by_label(inp, 'Subscription onboarding fee as % of implementation fee'),
    'Infra_Esc': by_label(inp, 'Infrastructure annual change'),
    'Days_per_LOB': by_label(inp, 'Extra implementation man-days per line of business above the included count'),
    'Round_To': by_label(inp, 'Round implementation and service fees to'),
    'Validity': by_label(inp, 'Quotation validity'),
    'Blended': by_label(inp, 'Blended day rate (weighted by mix)'),
    'Rate_BA': by_label(inp, 'Business analyst'),
    'Rate_Dev': by_label(inp, 'Developer'),
    'Rate_QA': by_label(inp, 'QA engineer'),
    'Rate_PM': by_label(inp, 'Project manager'),
    'Rate_Trainer': by_label(inp, 'Trainer day rate (end-user or train-the-trainer)'),
    'Rate_Onsite': by_label(inp, 'On-site consultant day rate (Metro Manila)'),
}
TIERS = []
for row in inp.iter_rows(values_only=True):
    if row[0] in ('Small', 'Medium', 'Large', 'Enterprise') and isinstance(row[1], (int, float)) and len(TIERS) < 4:
        TIERS.append(dict(name=row[0], frm=row[1], to=row[2], perp=row[3], sub=row[4], minb=row[5], base=row[6],
                          lobinc=row[7], train=row[8], ref=row[9], reflob=row[10], sup=row[11]))
assert [t['name'] for t in TIERS] == ['Small', 'Medium', 'Large', 'Enterprise'], TIERS

HOST = {}
for row in src['Infrastructure'].iter_rows(min_row=5, values_only=True):
    if row[0] in ('AWS', 'Azure', 'Local partner', 'Customer') and row[1] in ('Small', 'Medium', 'Large', 'Enterprise'):
        HOST[(row[0], row[1])] = row[5]
OPT = {}
for row in src['Optional Services'].iter_rows(min_row=5, values_only=True):
    if row[0] and isinstance(row[2], (int, float)):
        OPT[row[0]] = row[2]
O = {
    'Env_Setup': OPT['Additional environment (training, second UAT), set-up'],
    'Env_Month': OPT['Additional environment hosted by iorta TechNXT'],
    'Mig_Source': OPT['Data migration beyond the standard templates'],
    'Mig_Day': OPT['Data migration, additional effort'],
    'Int_Std': OPT['Additional integration, standard'],
    'Int_Cx': OPT['Additional integration, complex'],
    'LOB_After': OPT['Additional line of business after go-live'],
    'Report_Item': OPT['New report or document template'],
}
SRCNOTE = 'Read from BrokerVerse_Commercials_and_Pricing.xlsx (recommended price points) when this book was built.'

# ------------------------------------------------------------------ styles
NAVY = '0F4761'
F_HEAD = PatternFill('solid', fgColor=NAVY)
F_SUB = PatternFill('solid', fgColor='D9E2F3')
F_IN = PatternFill('solid', fgColor='FFF2CC')
F_TOT = PatternFill('solid', fgColor='E2EFDA')
F_INT = PatternFill('solid', fgColor='FCE4D6')
F_RED = PatternFill('solid', fgColor='F8CBAD')
WH = Font(name='Calibri', bold=True, color='FFFFFF')
B = Font(name='Calibri', bold=True)
BLUE = Font(name='Calibri', color='1F3FBF')
TITLE = Font(name='Calibri', bold=True, size=14, color=NAVY)
ITAL = Font(name='Calibri', italic=True, color='595959')
RED = Font(name='Calibri', bold=True, color='C00000')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
PHP = '#,##0'
PHP2 = '#,##0.00'
PCT = '0%'
PCT1 = '0.0%'

wb = Workbook()
wb.remove(wb.active)


def sheet(name, widths, title, subtitle=None, footer='iNXT BrokerVerse OOTB price book, iorta TechNXT Corp. INTERNAL. Page &P of &N'):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws['A1'] = title
    ws['A1'].font = TITLE
    if subtitle:
        ws['A2'] = subtitle
        ws['A2'].font = ITAL
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = footer
    return ws


def head(ws, r, labels, c0=1, height=None):
    for i, t in enumerate(labels):
        c = ws.cell(r, c0 + i, t)
        c.font = WH
        c.fill = F_HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BOX
    if height:
        ws.row_dimensions[r].height = height


def sec(ws, r, text, span=8):
    c = ws.cell(r, 1, text)
    c.font = Font(name='Calibri', bold=True, color=NAVY, size=12)
    for i in range(1, span + 1):
        ws.cell(r, i).fill = F_SUB


def put(ws, r, c, v, fmt=None, inp=False, bold=False, fill=None, wrap=False, font=None):
    cell = ws.cell(r, c, v)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    if inp:
        cell.fill = F_IN
        cell.font = BLUE
    if bold:
        cell.font = B
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if wrap:
        cell.alignment = WRAP
    return cell


def text_row(ws, r, text, span, height=None, font=None):
    c = ws.cell(r, 1, text)
    c.alignment = WRAP
    if font:
        c.font = font
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    if height:
        ws.row_dimensions[r].height = height


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def dv_list(ws, formula, ref):
    dv = DataValidation(type='list', formula1=formula, allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ref)


# ------------------------------------------------------------------ Read Me
rm = sheet('Read Me', [24, 120], 'iNXT BrokerVerse OOTB: price book',
           'iorta TechNXT Corp. Prepared 03 October 2026. INTERNAL: contains discount limits and floor prices. Do not send this workbook to a client.')
rows = [
    ('Purpose', 'Pre-set commercials for presentations and first quotations of iNXT BrokerVerse OOTB (the out-of-the-box, as-is version) '
     'to Philippine non-life insurance brokers. The presenter picks the size, licence model, hosting, term and options and the client sees the price on the spot.'),
    ('Source of prices', 'Every price point comes from the recommended price points of the pricing workbook BrokerVerse_Commercials_and_Pricing.xlsx '
     '(Inputs, Infrastructure and Optional Services sheets), copied to the Price Basis sheet when this book was built. '
     'If the pricing workbook changes, rebuild this book with docs/package/tools/build_price_book.py. Do not type new prices here.'),
    ('Sheets', 'Quick Quote: interactive quotation. Pre-set Packages: four ready packages per tier at list price and at the approved floor. '
     'Negotiation: approval matrix, floors and give/get levers (internal). Validity and Terms: validity, taxes and payment terms to read out with any quotation. '
     'Price Basis: the price points and drivers.'),
    ('How to use Quick Quote', '1. Choose "Tier reference" and pick a tier, or choose "Custom" and type the user count and the number of lines of business.\n'
     '2. Pick the licence model, hosting, contract term and optional services (yellow cells).\n'
     '3. Leave the discount cells at 0% in front of the client. Any discount needs the approval shown in the result (Negotiation sheet).\n'
     '4. Read the results and the payment schedule. The comparison block shows the other licence model with the same choices.\n'
     '5. Recalculate (F9) if your spreadsheet tool does not recalculate on open.'),
    ('What to show the client', 'Quick Quote at list price, the client columns of Pre-set Packages (list price only) and Validity and Terms. '
     'Never show the floor columns or the Negotiation sheet. The client-facing rate card is iNXT_BrokerVerse_Commercial_Proposal_Rate_Card.pdf.'),
    ('Prices and taxes', 'All prices are in PHP and exclude 12% VAT. USD figures are for reference only at the exchange rate on Price Basis. '
     'Withholding tax deducted by the client does not reduce the invoice price; the client issues BIR Form 2307.'),
    ('Year convention', 'Years are contract years counted from go-live, as in the pricing workbook. Implementation and the perpetual licence fall in Year 1. '
     'Subscription, hosting and 24x7 support are priced for 12 months in each year; AMC starts in the AMC start year (default Year 2, after a 12-month warranty). '
     'Recurring prices increase 5% at each anniversary; hosting follows the infrastructure change rate (default 0%).'),
    ('Colour key', 'Yellow cell with blue text: input. Green: result. Light orange: internal only (discounts, floors). Navy header: column titles.'),
]
head(rm, 4, ['Topic', 'Notes'])
for i, (k, v) in enumerate(rows):
    r = 5 + i
    put(rm, r, 1, k, bold=True, wrap=True)
    put(rm, r, 2, v, wrap=True)
    rm.row_dimensions[r].height = max(30, 15 * (len(v) // 120 + 1 + v.count('\n')))

# ------------------------------------------------------------------ Price Basis
pb = sheet('Price Basis', [44, 14, 14, 14, 14, 14, 14, 14, 14, 14, 16, 30],
           'Price basis: recommended price points', 'From the pricing workbook. PHP, exclusive of VAT. Change only by rebuilding from the pricing workbook.')
sec(pb, 4, 'General drivers', 12)
head(pb, 5, ['Driver', 'Value', 'Unit', 'Source'])
pb.merge_cells('D5:L5')
gen = [('FX', 'Exchange rate (reference only)', G['FX'], 'PHP per USD', '0.00'),
       ('FX_Date', 'Exchange rate date', G['FX_Date'], 'date', None),
       ('VAT', 'VAT rate', G['VAT'], '%', PCT),
       ('AMC_Rate', 'AMC rate (of perpetual licence fee)', G['AMC_Rate'], '%', PCT),
       ('Esc', 'Annual escalation (AMC, subscription, 24x7 support)', G['Esc'], '%', PCT),
       ('AMC_Start', 'AMC starts in contract year', G['AMC_Start'], 'year', '0'),
       ('Onboard_Pct', 'Subscription onboarding fee as % of implementation fee', G['Onboard_Pct'], '%', PCT),
       ('Infra_Esc', 'Infrastructure annual change', G['Infra_Esc'], '%', PCT),
       ('Days_per_LOB', 'Extra implementation man-days per line of business above the included count', G['Days_per_LOB'], 'man-days', '0'),
       ('Round_To', 'Rounding of implementation fees', G['Round_To'], 'PHP', PHP),
       ('Validity', 'Quotation validity', G['Validity'], 'days', '0'),
       ('Blended', 'Blended day rate', G['Blended'], 'PHP per man-day', PHP),
       ('Rate_BA', 'Business analyst day rate', G['Rate_BA'], 'PHP per man-day', PHP),
       ('Rate_Dev', 'Developer day rate', G['Rate_Dev'], 'PHP per man-day', PHP),
       ('Rate_QA', 'QA engineer day rate', G['Rate_QA'], 'PHP per man-day', PHP),
       ('Rate_PM', 'Project manager day rate', G['Rate_PM'], 'PHP per man-day', PHP),
       ('Rate_Trainer', 'Trainer day rate (beyond included days)', G['Rate_Trainer'], 'PHP per trainer day', PHP),
       ('Rate_Onsite', 'On-site consultant day (Metro Manila)', G['Rate_Onsite'], 'PHP per day', PHP)]
r = 5
for n, lab, v, unit, fmt in gen:
    r += 1
    put(pb, r, 1, lab)
    put(pb, r, 2, v, fmt, inp=True)
    put(pb, r, 3, unit)
    put(pb, r, 4, SRCNOTE if n in ('FX', 'AMC_Rate') else '')
    pb.merge_cells(start_row=r, start_column=4, end_row=r, end_column=12)
    name(n, f"'Price Basis'!$B${r}")

r += 2
sec(pb, r, 'Size tiers, graduated user slabs and tier parameters', 12)
r += 1
head(pb, r, ['Tier', 'From users', 'To users', 'Perpetual licence per user, slab (PHP)', 'Subscription per user per month, slab (PHP)',
             'Minimum billable users (subscription)', 'Implementation base man-days', 'Lines of business included',
             'Training days included', 'Reference users', '24x7 Severity 1 support per year (PHP)', 'Implementation duration (Implementation Approach)'], height=60)
T0 = r + 1
dur = {'Small': '8 weeks', 'Medium': '12 weeks', 'Large': '16 to 20 weeks', 'Enterprise': '20 weeks or more, planned at mobilisation'}
for t in TIERS:
    r += 1
    put(pb, r, 1, t['name'], bold=True)
    for j, k in enumerate(['frm', 'to', 'perp', 'sub', 'minb', 'base', 'lobinc', 'train', 'ref', 'sup'], 2):
        put(pb, r, j, t[k], PHP if k in ('perp', 'sub', 'sup') else '0', inp=True)
    put(pb, r, 12, dur[t['name']], wrap=True)
T1 = r
for n, col in [('TierName', 'A'), ('TierFrom', 'B'), ('TierTo', 'C'), ('PerpRate', 'D'), ('SubRate', 'E'), ('MinBill', 'F'),
               ('BaseDays', 'G'), ('LOBIncl', 'H'), ('TrainDays', 'I'), ('RefUsers', 'J'), ('Sup24', 'K'), ('Duration', 'L')]:
    name(n, f"'Price Basis'!${col}${T0}:${col}${T1}")
r += 1
text_row(pb, r, 'Graduated slabs: the first 25 users are charged at the Small rate, users 26 to 100 at the Medium rate, users 101 to 300 at the Large rate and the rest at the Enterprise rate. '
         'The tier, set by the total user count, decides the implementation effort, the minimum billable users for subscription, the training days included and the hosting size.', 12, 30)

r += 2
sec(pb, r, 'Hosting per month (production plus one UAT environment, PHP)', 12)
r += 1
HOSTS = ['Self-hosted', 'AWS', 'Azure', 'Local partner']
head(pb, r, ['Tier'] + HOSTS)
name('Host_Hdr', f"'Price Basis'!$B${r}:$E${r}")
H0 = r + 1
for t in TIERS:
    r += 1
    put(pb, r, 1, t['name'], bold=True)
    for j, h in enumerate(HOSTS):
        put(pb, r, 2 + j, HOST[('Customer' if h == 'Self-hosted' else h, t['name'])], PHP, inp=True)
name('Host_Tbl', f"'Price Basis'!$B${H0}:$E${r}")
r += 1
text_row(pb, r, 'Self-hosted: the client hosts in its own cloud account or data centre; no hosting fee, deployment guide supplied, set-up support at day rates. '
         'AWS and Azure are priced on their Singapore regions; the local partner option keeps data in the Philippines. Cloud prices are estimates to be confirmed with the provider before a binding offer.', 12, 30)

r += 2
sec(pb, r, 'Optional services (PHP)', 12)
r += 1
head(pb, r, ['Item', 'Price', 'Unit'])
for n, lab, unit in [('Env_Setup', 'Additional environment, set-up', 'one-time'),
                     ('Env_Month', 'Additional environment hosted by iorta TechNXT', 'per month'),
                     ('Mig_Source', 'Data migration beyond the standard templates', 'per legacy source'),
                     ('Mig_Day', 'Data migration, additional effort', 'per man-day'),
                     ('Int_Std', 'Additional integration, standard', 'per integration'),
                     ('Int_Cx', 'Additional integration, complex', 'per integration'),
                     ('LOB_After', 'Additional line of business after go-live', 'per line'),
                     ('Report_Item', 'New report or document template (typical)', 'per item')]:
    r += 1
    put(pb, r, 1, lab)
    put(pb, r, 2, O[n], PHP, inp=True)
    put(pb, r, 3, unit)
    name(n, f"'Price Basis'!$B${r}")
r += 1
put(pb, r, 1, 'Training beyond the included days')
put(pb, r, 2, '=Rate_Trainer', PHP)
put(pb, r, 3, 'per trainer day')
r += 1
put(pb, r, 1, 'On-site consultant day (Metro Manila)')
put(pb, r, 2, '=Rate_Onsite', PHP)
put(pb, r, 3, 'per day')
r += 2
text_row(pb, r, SRCNOTE + ' Pricing workbook state: recalculated with LibreOffice before reading.', 12, 18, ITAL)
pb.freeze_panes = 'B4'

# ------------------------------------------------------------------ Negotiation parameters (sheet built later, names fixed now)
NEG = {}


def slab(n_ref, col):
    return '+'.join(f"MAX(0,MIN({n_ref},'Price Basis'!$C${t})-'Price Basis'!$B${t}+1)*'Price Basis'!${col}${t}" for t in range(T0, T1 + 1))


# ------------------------------------------------------------------ calculation engine
ENGINE = [
    # key, label, fmt
    ('idx', 'Tier index', '0'),
    ('tier', 'Tier', None),
    ('bill', 'Billable users (subscription minimum applies)', '0'),
    ('md', 'Implementation man-days', '0'),
    ('impl_list', 'Implementation or onboarding fee, list', PHP),
    ('impl_net', 'Implementation or onboarding fee, quoted', PHP),
    ('lic_list', 'Perpetual licence fee, list (one-time)', PHP),
    ('lic_net', 'Perpetual licence fee, quoted (one-time)', PHP),
    ('sub_list', 'Subscription per month, Year 1, list', PHP),
    ('sub_net', 'Subscription per month, Year 1, quoted', PHP),
    ('eff', 'Effective price per user (licence one-time, or subscription per month)', PHP),
    ('host_list', 'Hosting per month, list', PHP),
    ('host_net', 'Hosting per month, quoted', PHP),
    ('env_m', 'Additional hosted environments per month', PHP),
    ('month', 'Monthly recurring at go-live (subscription, hosting, environments)', PHP),
    ('once_list', 'One-time optional services, list', PHP),
    ('once_net', 'One-time optional services, quoted', PHP),
    ('sup_list', '24x7 Severity 1 support per year, Year 1, list', PHP),
    ('train', 'Training days included', '0'),
] + [(f'amc{y}', f'AMC Year {y}', PHP) for y in range(1, 6)] \
  + [(f'sub{y}', f'Subscription fees Year {y}', PHP) for y in range(1, 6)] \
  + [(f'host{y}', f'Hosting and environments Year {y}', PHP) for y in range(1, 6)] \
  + [(f'sup{y}', f'24x7 support Year {y}', PHP) for y in range(1, 6)] \
  + [(f'tot{y}', f'Total Year {y}, excl. VAT', PHP) for y in range(1, 6)] \
  + [('termtot', 'Total for the contract term, excl. VAT', PHP),
     ('tot5y', '5-year total cost, excl. VAT', PHP),
     ('tot5vat', '5-year total cost, incl. VAT', PHP),
     ('tot5usd', '5-year total cost, USD reference', PHP)]
ENGINE_KEYS = [e[0] for e in ENGINE]


def engine(ws, col, r0, I, labels=True):
    """Write the engine in column `col` from row r0. I maps input keys to absolute cell references."""
    R = {k: f'{col}{r0 + i}' for i, k in enumerate(ENGINE_KEYS)}
    m = I['model']
    F = {
        'idx': f"MATCH({I['users']},TierFrom,1)",
        'tier': f"INDEX(TierName,{R['idx']})",
        'bill': f"IF({m}=\"Subscription\",MAX({I['users']},INDEX(MinBill,{R['idx']})),{I['users']})",
        'md': f"INDEX(BaseDays,{R['idx']})+MAX(0,{I['lobs']}-INDEX(LOBIncl,{R['idx']}))*Days_per_LOB",
        'impl_list': f"ROUND({R['md']}*Blended/Round_To,0)*Round_To*IF({m}=\"Subscription\",Onboard_Pct,1)",
        'impl_net': f"{R['impl_list']}*(1-{I['dsvc']})",
        'lic_list': f"IF({m}=\"Perpetual\",{slab(I['users'], 'D')},0)",
        'lic_net': f"{R['lic_list']}*(1-{I['dsw']})",
        'sub_list': f"IF({m}=\"Subscription\",{slab(R['bill'], 'E')},0)",
        'sub_net': f"{R['sub_list']}*(1-{I['dsw']})",
        'eff': f"IF({m}=\"Perpetual\",{R['lic_net']}/{I['users']},{R['sub_net']}/{R['bill']})",
        'host_list': f"INDEX(Host_Tbl,{R['idx']},MATCH({I['host']},Host_Hdr,0))",
        'host_net': f"{R['host_list']}*(1-{I['dhost']})",
        'env_m': f"IF({I['host']}=\"Self-hosted\",0,{I['env']}*Env_Month)",
        'month': f"{R['sub_net']}+{R['host_net']}+{R['env_m']}",
        'once_list': f"{I['trn']}*Rate_Trainer+{I['mig']}*Mig_Source+{I['ints']}*Int_Std+{I['intc']}*Int_Cx+{I['env']}*Env_Setup+{I['onsite']}*Rate_Onsite",
        'once_net': f"{R['once_list']}*(1-{I['dsvc']})",
        'sup_list': f"IF({I['sup']}=\"Yes\",INDEX(Sup24,{R['idx']}),0)",
        'train': f"INDEX(TrainDays,{R['idx']})",
    }
    for y in range(1, 6):
        F[f'amc{y}'] = f"IF(AND({m}=\"Perpetual\",{y}>=AMC_Start),{R['lic_net']}*AMC_Rate*(1+Esc)^({y}-AMC_Start),0)"
        F[f'sub{y}'] = f"{R['sub_net']}*12*(1+Esc)^({y}-1)"
        F[f'host{y}'] = f"{R['host_net']}*12*(1+Infra_Esc)^({y}-1)+{R['env_m']}*12*(1+Esc)^({y}-1)"
        F[f'sup{y}'] = f"{R['sup_list']}*(1-{I['dsvc']})*(1+Esc)^({y}-1)"
        one = f"{R['impl_net']}+{R['lic_net']}+{R['once_net']}+" if y == 1 else ''
        F[f'tot{y}'] = f"{one}{R[f'amc{y}']}+{R[f'sub{y}']}+{R[f'host{y}']}+{R[f'sup{y}']}"
    F['termtot'] = '+'.join(f"IF({I['term']}>={y},{R[f'tot{y}']},0)" for y in range(1, 6))
    F['tot5y'] = '+'.join(R[f'tot{y}'] for y in range(1, 6))
    F['tot5vat'] = f"{R['tot5y']}*(1+VAT)"
    F['tot5usd'] = f"{R['tot5y']}/FX"
    cidx = ws[f'{col}1'].column
    for i, (k, lab, fmt) in enumerate(ENGINE):
        if labels:
            put(ws, r0 + i, 1, lab)
        c = put(ws, r0 + i, cidx, '=' + F[k], fmt)
        if k in ('tot5y', 'termtot', 'tot1'):
            c.fill = F_TOT
            c.font = B
    return R


# ------------------------------------------------------------------ Quick Quote
qq = sheet('Quick Quote', [52, 22, 18, 18, 18, 18, 18, 18], 'Quick Quote',
           'Pick the options in the yellow cells. PHP, exclusive of 12% VAT unless stated. Valid 90 days from the quotation date.',
           footer='iNXT BrokerVerse OOTB quotation, iorta TechNXT Corp. Page &P of &N')
sec(qq, 4, 'Choices', 8)
head(qq, 5, ['Choice', 'Value', 'Notes'])
qq.merge_cells('C5:H5')
QI = {}
choices = [
    ('client', 'Client name', '[Client name]', None, 'Shown on the quotation.'),
    ('qdate', 'Quotation date', '=DATE(2026,10,3)', 'dd mmm yyyy', 'Valid until the date shown in the results.'),
    ('method', 'Sizing method', 'Tier reference', None, '"Tier reference" uses the reference users and lines of business of the tier. "Custom" uses the two cells below.'),
    ('tierpick', 'Tier (when sizing by tier reference)', 'Medium', None, 'Small 1 to 25 users, Medium 26 to 100, Large 101 to 300, Enterprise above 300.'),
    ('cusers', 'Named users (when Custom)', 40, '0', 'A user is a named person with a sign-in.'),
    ('clobs', 'Lines of business to configure (when Custom)', 6, '0', 'Motor, fire, marine, casualty, engineering, bonds and so on.'),
    ('model', 'Licence model', 'Subscription', None, 'Perpetual (one-time licence plus AMC) or Subscription (monthly).'),
    ('host', 'Hosting', 'AWS', None, 'Self-hosted, AWS, Azure or Local partner.'),
    ('term', 'Contract term (years)', 3, '0', '1, 3 or 5 years. The minimum subscription term is 12 months.'),
    ('sup', '24x7 Severity 1 support', 'No', None, 'Standard support is 08:00 to 18:00 PHT, Monday to Friday.'),
    ('trn', 'Extra training days', 0, '0', 'Beyond the days included for the tier.'),
    ('mig', 'Legacy data sources to migrate', 0, '0', 'Beyond the standard upload templates.'),
    ('ints', 'Additional standard integrations', 0, '0', 'One documented API or file exchange, one direction.'),
    ('intc', 'Additional complex integrations', 0, '0', 'Two-way, no API, or batch with reconciliation.'),
    ('env', 'Additional environments (training or second UAT)', 0, '0', 'Set-up fee, plus a monthly fee when iorta TechNXT hosts.'),
    ('onsite', 'On-site consultant days (Metro Manila)', 0, '0', 'Remote delivery is the default.'),
    ('dsw', 'Discount on licence or subscription', 0, PCT1, 'INTERNAL. Leave at 0% in front of the client. Approval per the Negotiation sheet.'),
    ('dsvc', 'Discount on implementation, services and 24x7 support', 0, PCT1, 'INTERNAL. As above.'),
    ('dhost', 'Discount on hosting', 0, PCT1, 'INTERNAL. As above.'),
]
r = 5
for k, lab, v, fmt, note in choices:
    r += 1
    put(qq, r, 1, lab)
    c = put(qq, r, 2, v, fmt, inp=True)
    put(qq, r, 3, note, wrap=False, fill=F_INT if note.startswith('INTERNAL') else None)
    qq.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    QI[k] = f'$B${r}'
dv_list(qq, '"Tier reference,Custom"', QI['method'].replace('$', ''))
dv_list(qq, '"Small,Medium,Large,Enterprise"', QI['tierpick'].replace('$', ''))
dv_list(qq, '"Perpetual,Subscription"', QI['model'].replace('$', ''))
dv_list(qq, '"Self-hosted,AWS,Azure,Local partner"', QI['host'].replace('$', ''))
dv_list(qq, '"1,3,5"', QI['term'].replace('$', ''))
dv_list(qq, '"Yes,No"', QI['sup'].replace('$', ''))
for k in ('cusers', 'clobs', 'trn', 'mig', 'ints', 'intc', 'env', 'onsite'):
    dv = DataValidation(type='whole', operator='between', formula1='0' if k != 'cusers' else '1', formula2='100000')
    qq.add_data_validation(dv)
    dv.add(QI[k].replace('$', ''))
for k in ('dsw', 'dsvc', 'dhost'):
    dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='0.5')
    qq.add_data_validation(dv)
    dv.add(QI[k].replace('$', ''))
r += 1
put(qq, r, 1, 'Named users used', bold=True)
put(qq, r, 2, f"=IF({QI['method']}=\"Custom\",{QI['cusers']},INDEX(RefUsers,MATCH({QI['tierpick']},TierName,0)))", '0', fill=F_TOT)
QI['users'] = f'$B${r}'
r += 1
put(qq, r, 1, 'Lines of business used', bold=True)
put(qq, r, 2, f"=IF({QI['method']}=\"Custom\",{QI['clobs']},INDEX(LOBIncl,MATCH({QI['tierpick']},TierName,0)))", '0', fill=F_TOT)
QI['lobs'] = f'$B${r}'

# calculation block (lower on the sheet): selected, selected at list, perpetual and subscription comparisons
CALC0 = 100
sec(qq, CALC0 - 2, 'Calculation detail (do not edit)', 8)
head(qq, CALC0 - 1, ['Line', 'Selected, quoted', 'Selected, at list', 'Perpetual, quoted', 'Subscription, quoted'])
QQ_PERP = '"Perpetual"'
I_sel = dict(QI)
I_list = dict(QI, dsw='0', dsvc='0', dhost='0')
I_perp = dict(QI, model=QQ_PERP)
I_sub = dict(QI, model='"Subscription"')
RS = engine(qq, 'B', CALC0, I_sel)
RL = engine(qq, 'C', CALC0, I_list, labels=False)
RP = engine(qq, 'D', CALC0, I_perp, labels=False)
RB = engine(qq, 'E', CALC0, I_sub, labels=False)

r += 2
sec(qq, r, 'Results', 8)
r += 1
head(qq, r, ['Item', 'PHP excl. VAT', 'PHP incl. VAT', 'USD reference', 'Notes'])
qq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
res = [
    ('Tier', f"={RS['tier']}", None, False, f"=\"Implementation about \"&INDEX(Duration,{RS['idx']})&\"; \"&{RS['train']}&\" training days included\""),
    ('One-time implementation fee (subscription: onboarding fee)', f"={RS['impl_net']}", PHP, True, f"={RS['md']}&\" man-days at the blended rate; \"&{QI['lobs']}&\" lines of business\""),
    ('Perpetual licence fee (one-time, on go-live)', f"={RS['lic_net']}", PHP, True, f"=IF({QI['model']}=\"Perpetual\",TEXT({RS['eff']},\"#,##0\")&\" per user on average, graduated slabs\",\"Not applicable to subscription\")"),
    ('Subscription per month, Year 1', f"={RS['sub_net']}", PHP, True, f"=IF({QI['model']}=\"Subscription\",{RS['bill']}&\" billable users, \"&TEXT({RS['eff']},\"#,##0\")&\" per user per month on average\",\"Not applicable to perpetual\")"),
    ('AMC per year, first AMC year', f"=CHOOSE(MIN(AMC_Start,5),{RS['amc1']},{RS['amc2']},{RS['amc3']},{RS['amc4']},{RS['amc5']})", PHP, True, f"=IF({QI['model']}=\"Perpetual\",\"22% of the licence fee from Year \"&AMC_Start&\", then +5% a year\",\"Not applicable: support is in the subscription\")"),
    ('Hosting per month', f"={RS['host_net']}+{RS['env_m']}", PHP, True, f"={QI['host']}&IF({QI['env']}>0,\", including \"&{QI['env']}&\" additional environment(s)\",\"\")"),
    ('One-time optional services', f"={RS['once_net']}", PHP, True, 'Training, migration, integrations, environment set-up, on-site days chosen above.'),
    ('24x7 Severity 1 support per year, Year 1', f"={RS['sup1']}", PHP, True, 'Optional; +5% a year.'),
    ('Year 1 total', f"={RS['tot1']}", PHP, True, 'Everything payable in the first contract year.'),
    ('Total for the contract term', f"={RS['termtot']}", PHP, True, f"=\"Contract term of \"&{QI['term']}&\" year(s)\""),
    ('5-year total cost', f"={RS['tot5y']}", PHP, True, 'Assumes renewal to Year 5 at the escalated prices.'),
    ('Discount included in the 5-year total', f"={RL['tot5y']}-{RS['tot5y']}", PHP, True, 'INTERNAL. 0 at list price.'),
]
RES = {}
for lab, f, fmt, money, note in res:
    r += 1
    put(qq, r, 1, lab, bold=lab in ('Year 1 total', '5-year total cost'))
    c = put(qq, r, 2, f, fmt, fill=F_TOT if money else None)
    if lab in ('Year 1 total', '5-year total cost', 'Total for the contract term'):
        c.font = B
    if money:
        put(qq, r, 3, f'=B{r}*(1+VAT)', PHP)
        put(qq, r, 4, f'=B{r}/FX', PHP)
    put(qq, r, 5, note, fill=F_INT if str(note).startswith('INTERNAL') else None)
    qq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    RES[lab] = r
r += 1
put(qq, r, 1, 'Approval needed for the discounts entered', bold=True)
lvl_sw = f"IF({QI['dsw']}<=0,0,IF({QI['dsw']}<=Disc_AM,1,IF({QI['dsw']}<=Disc_SH,2,IF({QI['dsw']}<=Disc_CEO,3,4))))"
lvl_svc = f"IF({QI['dsvc']}<=0,0,IF({QI['dsvc']}<=Disc_AM,1,IF({QI['dsvc']}<=Disc_SH,2,IF({QI['dsvc']}<=Disc_CEO,3,4))))"
lvl_h = f"IF({QI['dhost']}<=0,0,IF({QI['dhost']}<=Host_AM,1,IF({QI['dhost']}<=Host_SH,2,IF({QI['dhost']}<=Host_CEO,3,4))))"
put(qq, r, 2, f'=CHOOSE(MAX({lvl_sw},{lvl_svc},{lvl_h})+1,"None: list price","Account manager","Sales head","CEO","Not allowed: below the floor")', fill=F_INT, font=RED)
qq.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
put(qq, r, 5, 'INTERNAL. Limits on the Negotiation sheet.', fill=F_INT)
qq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
r += 1
put(qq, r, 1, 'Valid until', bold=True)
put(qq, r, 2, f"={QI['qdate']}+Validity", 'dd mmm yyyy', fill=F_TOT)

r += 2
sec(qq, r, 'Year by year (PHP, excl. VAT)', 8)
r += 1
head(qq, r, ['Item', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total'])
lines = [('Implementation or onboarding, licence and one-time services', lambda y: f"={RS['impl_net']}+{RS['lic_net']}+{RS['once_net']}" if y == 1 else '=0'),
         ('AMC', lambda y: f"={RS[f'amc{y}']}"),
         ('Subscription', lambda y: f"={RS[f'sub{y}']}"),
         ('Hosting and additional environments', lambda y: f"={RS[f'host{y}']}"),
         ('24x7 Severity 1 support', lambda y: f"={RS[f'sup{y}']}"),
         ('Total excl. VAT', lambda y: f"={RS[f'tot{y}']}"),
         ('VAT 12%', None),
         ('Total incl. VAT', None)]
y0 = r + 1
for k, (lab, fn) in enumerate(lines):
    r += 1
    put(qq, r, 1, lab, bold=lab.startswith('Total'))
    for y in range(1, 6):
        col = L(1 + y)
        if lab == 'VAT 12%':
            f = f'={col}{r-1}*VAT'
        elif lab == 'Total incl. VAT':
            f = f'={col}{r-2}+{col}{r-1}'
        else:
            f = fn(y)
        put(qq, r, 1 + y, f, PHP, fill=F_TOT if lab.startswith('Total') else None)
    put(qq, r, 7, f'=SUM(B{r}:F{r})', PHP, bold=True, fill=F_TOT)

r += 2
sec(qq, r, 'Payment schedule (PHP)', 8)
r += 1
head(qq, r, ['Milestone or period', 'Amount excl. VAT', 'VAT', 'Amount incl. VAT', 'When'])
qq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
pays = [
    ('Implementation or onboarding: 40% on signing', f"={RS['impl_net']}*0.4", 'On signing of the Order Form'),
    ('Implementation or onboarding: 40% on UAT sign-off', f"={RS['impl_net']}*0.4", 'On written UAT sign-off'),
    ('Implementation or onboarding: 20% on go-live', f"={RS['impl_net']}*0.2", 'On go-live'),
    ('Perpetual licence fee: 100% on go-live', f"={RS['lic_net']}", 'On go-live (perpetual model only)'),
    ('One-time optional services', f"={RS['once_net']}", 'On delivery, or as quoted for fixed-price items'),
    ('Subscription, per month (Year 1 rate)', f"={RS['sub_net']}", 'Monthly in advance from go-live; +5% at each anniversary'),
    ('Hosting and environments, per month', f"={RS['host_net']}+{RS['env_m']}", 'Monthly in advance from environment handover'),
    ('AMC, per year (first AMC year)', f"=CHOOSE(MIN(AMC_Start,5),{RS['amc1']},{RS['amc2']},{RS['amc3']},{RS['amc4']},{RS['amc5']})", f"=\"Yearly in advance from Year \"&AMC_Start&\"; +5% a year\""),
    ('24x7 Severity 1 support, per year (Year 1)', f"={RS['sup1']}", 'Yearly in advance; +5% a year'),
]
for lab, f, when in pays:
    r += 1
    put(qq, r, 1, lab)
    put(qq, r, 2, f, PHP)
    put(qq, r, 3, f'=B{r}*VAT', PHP)
    put(qq, r, 4, f'=B{r}+C{r}', PHP, fill=F_TOT)
    put(qq, r, 5, when)
    qq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
r += 1
text_row(qq, r, 'Invoices are payable within 30 days. Clients that are withholding agents deduct creditable withholding tax as their status and the payment type require and issue BIR Form 2307; '
         'withholding does not reduce the invoice price.', 8, 30, ITAL)

r += 2
sec(qq, r, 'Comparison: the same choices under each licence model (PHP, excl. VAT)', 8)
r += 1
head(qq, r, ['Item', 'Perpetual', 'Subscription', 'Difference (subscription less perpetual)'])
for lab, key in [('One-time fees in Year 1 (implementation, licence, services)', None), ('Year 1 total', 'tot1'), ('5-year total cost', 'tot5y')]:
    r += 1
    put(qq, r, 1, lab, bold=key == 'tot5y')
    if key:
        put(qq, r, 2, f"={RP[key]}", PHP, fill=F_TOT)
        put(qq, r, 3, f"={RB[key]}", PHP, fill=F_TOT)
    else:
        put(qq, r, 2, f"={RP['impl_net']}+{RP['lic_net']}+{RP['once_net']}", PHP)
        put(qq, r, 3, f"={RB['impl_net']}+{RB['lic_net']}+{RB['once_net']}", PHP)
    put(qq, r, 4, f'=C{r}-B{r}', PHP)
r += 1
put(qq, r, 1, 'Year in which subscription costs more in total than perpetual')
cum_p = lambda y: '+'.join(RP[f'tot{k}'] for k in range(1, y + 1))  # noqa: E731
cum_s = lambda y: '+'.join(RB[f'tot{k}'] for k in range(1, y + 1))  # noqa: E731
put(qq, r, 2, '=' + ''.join(f'IF(({cum_s(y)})>({cum_p(y)}),"Year {y}",' for y in range(1, 6)) + '"After Year 5"' + ')' * 5)
assert r < CALC0 - 3, r
qq.freeze_panes = 'A4'
qq.print_area = f'A1:H{CALC0 - 3}'

# ------------------------------------------------------------------ Negotiation
ng = sheet('Negotiation', [34, 16, 16, 16, 16, 18, 60], 'Negotiation: approval matrix, floors and levers',
           'INTERNAL. Never shown or sent to a client. Discounts are off the list price of each component.')
sec(ng, 4, 'Approval matrix: maximum discount off list by approver', 7)
head(ng, 5, ['Component', 'Account manager', 'Sales head', 'CEO', 'Floor: never below (discount)', 'Floor as % of list', 'Rules'], height=32)
matrix = [
    ('Perpetual licence fee', ('Disc_AM', 0.05), ('Disc_SH', 0.10), ('Disc_CEO', 0.15), ('Floor_SW', '=D6'), 'AMC follows the quoted (net) licence fee. AMC rate 22% and 5% escalation are not discounted (see levers).'),
    ('Subscription per user per month', '=B6', '=C6', '=D6', '=E6', 'Discount fixed for the initial term; renewals at the then list price less the same discount unless agreed otherwise.'),
    ('Implementation or onboarding fee', '=B6', '=C6', '=D6', ('Floor_Svc', '=D6'), 'Reducing scope lowers the fee at the blended rate and is not a discount.'),
    ('One-time optional services and 24x7 support', '=B6', '=C6', '=D6', '=E8', 'Same limits as implementation.'),
    ('Hosting per month', ('Host_AM', 0.0), ('Host_SH', 0.05), ('Host_CEO', 0.10), ('Floor_Host', '=D10'), 'Hosting carries a 20% margin over estimated cloud cost; 10% off keeps about 8% for FX and operations.'),
    ('Change request and day rates', 0.0, 0.05, 0.10, 0.10, 'Day rates are fixed for 12 months from signing. No discount on CRs below 5 man-days.'),
    ('AMC rate (22%) and escalation (5%)', 0.0, 0.0, 0.0, 0.0, 'Not discounted. Only the multi-year AMC prepayment lever may hold the escalation, with CEO approval.'),
]
for i, (lab, *vals, rule) in enumerate(matrix):
    r = 6 + i
    put(ng, r, 1, lab, bold=True)
    for j, v in enumerate(vals):
        if isinstance(v, tuple):
            n, val = v
            put(ng, r, 2 + j, val, PCT, inp=True)
            name(n, f'Negotiation!${L(2 + j)}${r}')
        else:
            put(ng, r, 2 + j, v, PCT, inp=not str(v).startswith('='))
    put(ng, r, 6, f'=1-E{r}', PCT, fill=F_TOT)
    put(ng, r, 7, rule, wrap=True)
    ng.row_dimensions[r].height = 30
r = 6 + len(matrix)
text_row(ng, r, 'The floor equals the CEO limit: nobody may approve a price below it. A deal that needs more must be restructured with the levers below '
         '(scope, phasing, term, payment) or declined. Total discount from all levers and discretion together may not pass the CEO limit.', 7, 30, RED)

r += 2
sec(ng, r, 'Floor prices per tier (reference users and lines of business)', 7)
r += 1
head(ng, r, ['Tier', 'Perpetual licence per user, list (first slab of tier)', 'At floor', 'Subscription per user per month, list', 'At floor',
             'Implementation fee, list / at floor', 'Hosting per month on AWS, list / at floor'], height=45)
for i, t in enumerate(TIERS):
    r += 1
    tr = T0 + i
    put(ng, r, 1, t['name'], bold=True)
    put(ng, r, 2, f"='Price Basis'!D{tr}", PHP)
    put(ng, r, 3, f'=B{r}*(1-Floor_SW)', PHP, fill=F_INT)
    put(ng, r, 4, f"='Price Basis'!E{tr}", PHP)
    put(ng, r, 5, f'=D{r}*(1-Floor_SW)', PHP, fill=F_INT)
    impl = f"ROUND('Price Basis'!G{tr}*Blended/Round_To,0)*Round_To"
    put(ng, r, 6, f'=TEXT({impl},"#,##0")&" / "&TEXT({impl}*(1-Floor_Svc),"#,##0")', fill=F_INT)
    put(ng, r, 7, f"=TEXT('Price Basis'!C{H0 + i},\"#,##0\")&\" / \"&TEXT('Price Basis'!C{H0 + i}*(1-Floor_Host),\"#,##0\")", fill=F_INT)
r += 1
text_row(ng, r, 'Per-user rates are the slab rates; the graduated slabs still apply at the floor (each slab rate less the floor discount).', 7, 18, ITAL)

r += 2
sec(ng, r, 'Give and get levers', 7)
r += 1
head(ng, r, ['Lever', 'What iorta TechNXT gets', 'Allowed price effect', 'Maximum', 'Approver', 'Counts toward the discount limit', 'Conditions'], height=32)
levers = [
    ('Longer term', '3-year or 5-year committed subscription or AMC term, no termination for convenience in the term',
     'Off subscription fee', '3 years: 3%; 5 years: 5%', 'Account manager (3%); sales head (5%)', 'Yes',
     'Early termination fee equal to the remaining committed fees of the term, or 50% of them if the owner prefers (decision).'),
    ('Upfront payment', 'Subscription paid yearly in advance, or implementation 100% on signing',
     'Off the amount prepaid', 'Subscription yearly in advance: 4%; implementation 100% on signing: 3%', 'Account manager', 'Yes',
     'Discount applies only to the prepaid period; lost if a payment is late more than 30 days.'),
    ('Reference customer', 'Named as a client; two reference calls a year with prospects; logo use',
     'Off licence or subscription', '2%', 'Account manager', 'Yes', 'Signed reference clause in the Order Form; given after go-live.'),
    ('Case study', 'Joint published case study within 6 months of go-live, with a quote from the sponsor',
     'Off licence or subscription, or a credit', '2%, or PHP 100,000 credit against CRs', 'Sales head', 'Yes',
     'Credit used within 12 months; not paid out in cash.'),
    ('Multi-year AMC prepayment', 'AMC for 3 years paid in advance at the start of the AMC term',
     'AMC for the prepaid years held at the first AMC year rate (no 5% escalation)', 'About 5% of the prepaid AMC', 'CEO', 'No (AMC lever)',
     'Non-refundable except for iorta TechNXT breach; escalation resumes after the prepaid years.'),
    ('Reduced scope', 'Fewer lines of business at go-live, broker loads its own data, remote training only',
     'Implementation fee falls by the man-days removed at the blended rate', 'Up to the man-days removed (6 man-days per line of business)', 'Account manager', 'No (not a discount)',
     'Removed scope can be added later at the optional services price. Base man-days of the tier are not reduced.'),
    ('Phased users', 'Contract for the full user count, delivered in tranches',
     'Licence or subscription billed per tranche; slab prices locked for 24 months', 'Cash timing only', 'Sales head', 'No',
     'Subscription minimum billable users of the tier still apply. Perpetual tranches invoiced on activation.'),
    ('Signing date', 'Order Form signed by a stated date (for example quarter end)',
     'Off licence or subscription', '2%', 'Sales head', 'Yes', 'Only with a dated written offer; lapses after the date.'),
    ('Onboarding reduction (subscription)', '5-year subscription term', 'Onboarding fee reduced to 90% of the implementation fee', '10% off onboarding', 'Sales head', 'Yes',
     'The pricing workbook allows the onboarding percentage as a concession; the reduction is recovered in the term.'),
]
for lv in levers:
    r += 1
    for j, v in enumerate(lv):
        put(ng, r, 1 + j, v, wrap=True, bold=j == 0)
    ng.row_dimensions[r].height = 48
r += 2
sec(ng, r, 'Rules for every discount', 7)
for line in ['1. Start at list price. Offer a lever (a get for a give) before any discretionary discount.',
             '2. Each discount is shown as a separate line on the quotation and the Order Form, against the list price, with the lever named.',
             '3. Record the approval (approver, date, reason) in the deal file before the quotation leaves iorta TechNXT. E-mail approval is enough.',
             '4. Discounts apply to the initial term only unless the Order Form says otherwise. Added users are priced at list less the same discount in the initial term.',
             '5. Hosting is priced from cloud cost; never discount it to win the software deal.',
             '6. No side letters, free periods or unpriced extras. A free item is entered at list price with a 100% discount and counts toward the limit.',
             '7. Quotations are valid 90 days. A re-issued quotation after expiry is repriced at the current price book.']:
    r += 1
    text_row(ng, r, line, 7, 18)
ng.freeze_panes = 'A4'

# ------------------------------------------------------------------ Pre-set Packages
pp = sheet('Pre-set Packages', [46] + [15] * 8, 'Pre-set packages: four per tier, at list and at the approved floor',
           'Reference users and lines of business of each tier. PHP, exclusive of VAT. List columns may be shown to the client; floor columns are INTERNAL.')
PACKS = [
    ('Essentials', 'Subscription', 'Self-hosted', 'No', 0, 0, 3, 'Subscription; the client hosts; standard support; 3-year term.'),
    ('Standard', 'Subscription', 'AWS', 'No', 0, 0, 3, 'Subscription hosted by iorta TechNXT on AWS (Singapore); 3-year term.'),
    ('Ownership', 'Perpetual', 'AWS', 'No', 0, 0, 5, 'Perpetual licence with AMC, hosted on AWS (Singapore); 5-year view.'),
    ('Ownership Plus', 'Perpetual', 'Local partner', 'Yes', 2, 1, 5, 'Perpetual with AMC, hosted in the Philippines, 24x7 Severity 1 support, 2 extra training days, 1 legacy source migrated.'),
]
SUMMARY_R0 = 4
sec(pp, SUMMARY_R0, 'Summary', 9)
head(pp, SUMMARY_R0 + 1, ['Tier and package', 'One-time fees, list', 'One-time fees, floor', 'Monthly fees at go-live, list', 'Monthly fees, floor',
                          'Year 1 total, list', 'Year 1 total, floor', '5-year total, list', '5-year total, floor'], height=45)
for c in (3, 5, 7, 9):
    pp.cell(SUMMARY_R0 + 1, c).fill = PatternFill('solid', fgColor='843C0C')
sum_r = SUMMARY_R0 + 2
BLOCK0 = SUMMARY_R0 + 2 + 16 + 2
INPUT_ROWS = [('users', 'Named users', '0'), ('lobs', 'Lines of business', '0'), ('model', 'Licence model', None), ('host', 'Hosting', None),
              ('sup', '24x7 Severity 1 support', None), ('trn', 'Extra training days', '0'), ('mig', 'Legacy data sources migrated', '0'),
              ('ints', 'Additional standard integrations', '0'), ('intc', 'Additional complex integrations', '0'), ('env', 'Additional environments', '0'),
              ('onsite', 'On-site days', '0'), ('term', 'Contract term (years)', '0'), ('dsw', 'Discount on licence or subscription', PCT),
              ('dsvc', 'Discount on implementation and services', PCT), ('dhost', 'Discount on hosting', PCT)]
r = BLOCK0
PACK_CELLS = []
for ti, t in enumerate(TIERS):
    tr = T0 + ti
    sec(pp, r, f"{t['name']} tier: {t['ref']} users, {t['lobinc']} lines of business", 9)
    r += 1
    hdr = ['Line']
    for p in PACKS:
        hdr += [f'{p[0]}, list', f'{p[0]}, floor']
    head(pp, r, hdr, height=32)
    for j in range(8):
        if j % 2 == 1:
            pp.cell(r, 2 + j).fill = PatternFill('solid', fgColor='843C0C')
    r += 1
    put(pp, r, 1, 'What it is', bold=True)
    for pi, p in enumerate(PACKS):
        put(pp, r, 2 + pi * 2, p[7], wrap=True)
        pp.merge_cells(start_row=r, start_column=2 + pi * 2, end_row=r, end_column=3 + pi * 2)
    pp.row_dimensions[r].height = 62
    in0 = r + 1
    for k, (key, lab, fmt) in enumerate(INPUT_ROWS):
        put(pp, in0 + k, 1, lab)
    for pi, p in enumerate(PACKS):
        for side in (0, 1):
            col = 2 + pi * 2 + side
            vals = {'users': f"='Price Basis'!J{tr}", 'lobs': f"='Price Basis'!H{tr}", 'model': p[1], 'host': p[2], 'sup': p[3],
                    'trn': p[4], 'mig': p[5], 'ints': 0, 'intc': 0, 'env': 0, 'onsite': 0, 'term': p[6],
                    'dsw': '=Floor_SW' if side else 0, 'dsvc': '=Floor_Svc' if side else 0, 'dhost': '=Floor_Host' if side else 0}
            for k, (key, lab, fmt) in enumerate(INPUT_ROWS):
                put(pp, in0 + k, col, vals[key], fmt, fill=F_INT if (side and key.startswith('d')) else None)
    r = in0 + len(INPUT_ROWS)
    for pi, p in enumerate(PACKS):
        for side in (0, 1):
            col = L(2 + pi * 2 + side)
            I = {key: f'{col}${in0 + k}' for k, (key, lab, fmt) in enumerate(INPUT_ROWS)}
            I = {k: v.replace(col, '$' + col) for k, v in I.items()}
            R = engine(pp, col, r, I, labels=(pi == 0 and side == 0))
            PACK_CELLS.append((t['name'], p[0], p[1], p[2], side, R))
    r += len(ENGINE) + 2

# summary rows
for ti, t in enumerate(TIERS):
    for pi, p in enumerate(PACKS):
        Rl = [x for x in PACK_CELLS if x[0] == t['name'] and x[1] == p[0] and x[4] == 0][0][5]
        Rf = [x for x in PACK_CELLS if x[0] == t['name'] and x[1] == p[0] and x[4] == 1][0][5]
        put(pp, sum_r, 1, f"{t['name']}: {p[0]} ({p[1].lower()}, {p[2]})", bold=pi == 0)
        put(pp, sum_r, 2, f"={Rl['impl_net']}+{Rl['lic_net']}+{Rl['once_net']}", PHP)
        put(pp, sum_r, 3, f"={Rf['impl_net']}+{Rf['lic_net']}+{Rf['once_net']}", PHP, fill=F_INT)
        put(pp, sum_r, 4, f"={Rl['month']}", PHP)
        put(pp, sum_r, 5, f"={Rf['month']}", PHP, fill=F_INT)
        put(pp, sum_r, 6, f"={Rl['tot1']}", PHP, fill=F_TOT)
        put(pp, sum_r, 7, f"={Rf['tot1']}", PHP, fill=F_INT)
        put(pp, sum_r, 8, f"={Rl['tot5y']}", PHP, fill=F_TOT)
        put(pp, sum_r, 9, f"={Rf['tot5y']}", PHP, fill=F_INT)
        sum_r += 1
text_row(pp, sum_r, 'One-time fees: implementation or onboarding, perpetual licence and one-time services. Monthly fees: subscription, hosting and additional environments. '
         'AMC (perpetual) and 24x7 support are yearly and are in the Year 1 and 5-year totals. Floor = CEO limit on every component.', 9, 30, ITAL)
pp.freeze_panes = 'B4'

# ------------------------------------------------------------------ Validity and Terms
vt = sheet('Validity and Terms', [40, 100], 'Validity and terms', 'Read out or attach with every quotation. Applies to the Quick Quote and the pre-set packages.',
           footer='iNXT BrokerVerse OOTB quotation terms, iorta TechNXT Corp. Page &P of &N')
put(vt, 4, 1, 'Quotation date', bold=True)
put(vt, 4, 2, "='Quick Quote'!" + QI['qdate'], 'dd mmmm yyyy', fill=F_TOT)
put(vt, 5, 1, 'Valid until', bold=True)
put(vt, 5, 2, "='Quick Quote'!" + QI['qdate'] + '+Validity', 'dd mmmm yyyy', fill=F_TOT)
terms = [
    ('Validity', '=Validity&" days from the quotation date, for an Order Form signed in that period. Day rates are fixed for 12 months from signing."'),
    ('Currency', '="Philippine peso (PHP). USD figures are for reference only at PHP "&TEXT(FX,"0.00")&" per USD ("&FX_Date&")."'),
    ('VAT', 'All prices exclude 12% VAT, which is added on each invoice.'),
    ('Withholding tax', 'A client that is a withholding agent deducts creditable withholding tax as its status and the payment type require and issues BIR Form 2307 for each deduction. '
     'Withholding does not reduce the invoice price. The treatment of the perpetual licence fee (sale of software or royalty) is confirmed by both parties\' tax advisers before signing.'),
    ('Implementation or onboarding fee', '40% on signing, 40% on UAT sign-off, 20% on go-live.'),
    ('Perpetual licence fee', '100% on go-live.'),
    ('AMC', '22% of the licence fee a year, yearly in advance from the AMC start year (Year 2, after a 12-month warranty from go-live). Increases 5% at each anniversary.'),
    ('Subscription', 'Monthly in advance from go-live. Increases 5% at each contract anniversary. Minimum term 12 months; minimum billable users by tier.'),
    ('Hosting', 'Monthly in advance from the date the environment is handed over. Cloud prices are estimates confirmed with the provider before a binding offer.'),
    ('Change requests', 'Small: on delivery. Medium: 50% on approval, 50% on UAT sign-off. Large: 30% on approval, 50% on UAT sign-off, 20% on deployment.'),
    ('Optional services', 'Monthly in arrears at actual days, or as quoted for fixed-price items. Travel outside Metro Manila at cost.'),
    ('Payment', 'Invoices are payable within 30 days. Subscription and hosting may be suspended after 60 days of non-payment, with written notice.'),
    ('Scope', 'iNXT BrokerVerse OOTB as delivered. Changes to the product are change requests at the published day rates.'),
    ('Contract', 'Master Services Agreement with the Order Form and the schedules for the chosen model (licence or subscription, implementation SOW, support and SLA, hosting, data processing).'),
]
head(vt, 7, ['Term', 'Detail'])
for i, (k, v) in enumerate(terms):
    rr = 8 + i
    put(vt, rr, 1, k, bold=True)
    put(vt, rr, 2, v, wrap=True)
    vt.row_dimensions[rr].height = max(18, 15 * (len(v) // 95 + 1))

order = ['Read Me', 'Quick Quote', 'Pre-set Packages', 'Negotiation', 'Validity and Terms', 'Price Basis']
wb._sheets = [wb[n] for n in order]
for n in ('Negotiation',):
    wb[n].sheet_properties.tabColor = 'C00000'
wb.active = 1
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('wrote', OUT)
