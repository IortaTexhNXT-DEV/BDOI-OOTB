"""Builds 01_Sales/iNXT_BrokerVerse_ROI_Calculator.xlsx (formula driven).

    python3 build_roi.py

Package costs are the list prices of the pre-set packages in 02_Commercials/iNXT_BrokerVerse_Price_Book.xlsx, read after a
LibreOffice (headless) recalculation. Run build_price_book.py first.
"""
import os
import shutil
import subprocess
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
BOOK = os.path.join(PKG, '02_Commercials', 'iNXT_BrokerVerse_Price_Book.xlsx')
OUT = os.path.join(PKG, '01_Sales', 'iNXT_BrokerVerse_ROI_Calculator.xlsx')


def recalculated(path):
    tmp, prof = tempfile.mkdtemp(), tempfile.mkdtemp()
    subprocess.run(['soffice', f'-env:UserInstallation=file://{prof}', '--headless', '--convert-to', 'xlsx', '--outdir', tmp, path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    wb = load_workbook(os.path.join(tmp, os.path.basename(path)), data_only=True)
    shutil.rmtree(tmp, ignore_errors=True)
    shutil.rmtree(prof, ignore_errors=True)
    return wb


# ------------------------------------------------------------------ package costs from the price book (list price)
pp = recalculated(BOOK)['Pre-set Packages']
PACKS = []
tier = None
hdr = None
for row in pp.iter_rows(min_row=23):
    a = row[0].value
    if isinstance(a, str) and ' tier: ' in a:
        tier = a.split(' tier:')[0]
    elif a == 'Line':
        hdr = [c.value for c in row]
    elif isinstance(a, str) and a.startswith('Total Year ') and tier:
        y = int(a.split()[2][0])
        for j in (1, 3, 5, 7):
            pname = hdr[j].replace(', list', '')
            key = (tier, pname)
            rec = next((p for p in PACKS if p['key'] == key), None)
            if rec is None:
                rec = {'key': key, 'years': [0] * 5}
                PACKS.append(rec)
            rec['years'][y - 1] = row[j].value
for row in pp.iter_rows(min_row=6, max_row=21, values_only=True):
    t, rest = row[0].split(': ', 1)
    pname, desc = rest.split(' (', 1)
    rec = next(p for p in PACKS if p['key'] == (t, pname))
    rec['label'] = f'{t}: {pname} ({desc}'
assert len(PACKS) == 16 and all(len(p['years']) == 5 and all(isinstance(v, (int, float)) for v in p['years']) for p in PACKS)

# ------------------------------------------------------------------ styles
NAVY = '0F4761'
F_HEAD = PatternFill('solid', fgColor=NAVY)
F_SUB = PatternFill('solid', fgColor='D9E2F3')
F_IN = PatternFill('solid', fgColor='FFF2CC')
F_TOT = PatternFill('solid', fgColor='E2EFDA')
F_EST = PatternFill('solid', fgColor='FCE4D6')
WH = Font(name='Calibri', bold=True, color='FFFFFF')
B = Font(name='Calibri', bold=True)
BLUE = Font(name='Calibri', color='1F3FBF')
TITLE = Font(name='Calibri', bold=True, size=14, color=NAVY)
ITAL = Font(name='Calibri', italic=True, color='595959')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
PHP = '#,##0'
NUM = '#,##0'
DEC = '#,##0.0'
PCT = '0%'
PCT1 = '0.0%'

wb = Workbook()
wb.remove(wb.active)


def sheet(name, widths, title, subtitle=None):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws['A1'] = title
    ws['A1'].font = TITLE
    if subtitle:
        ws['A2'] = subtitle
        ws['A2'].font = ITAL
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = 'iNXT BrokerVerse OOTB ROI calculator, iorta TechNXT Corp. Page &P of &N'
    return ws


def head(ws, r, labels, c0=1, height=None):
    for i, t in enumerate(labels):
        c = ws.cell(r, c0 + i, t)
        c.font = WH
        c.fill = F_HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BOX
    if height:
        ws.row_dimensions[r].height = height


def sec(ws, r, text, span=8):
    c = ws.cell(r, 1, text)
    c.font = Font(name='Calibri', bold=True, color=NAVY, size=12)
    for i in range(1, span + 1):
        ws.cell(r, i).fill = F_SUB


def put(ws, r, c, v, fmt=None, inp=False, bold=False, fill=None, wrap=False):
    cell = ws.cell(r, c, v)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    if inp:
        cell.fill = F_IN
        cell.font = BLUE
    if bold:
        cell.font = B
    if fill:
        cell.fill = fill
    if wrap:
        cell.alignment = WRAP
    return cell


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


# ------------------------------------------------------------------ Read Me
rm = sheet('Read Me', [24, 120], 'iNXT BrokerVerse OOTB: ROI calculator',
           'iorta TechNXT Corp. Prepared 03 October 2026. PHP, exclusive of VAT. Example figures only: replace them with the broker\'s own.')
notes = [
    ('Purpose', 'Estimates the yearly benefit of iNXT BrokerVerse OOTB for a broker (staff hours saved, rework avoided, remittance and commission '
     'cash effects, renewal retention) and compares it with the cost of a chosen package to give the payback period, the 5-year net benefit and the return.'),
    ('How to use', '1. On Inputs, replace every yellow cell with the broker\'s figures. Each input has a label, a unit and its basis.\n'
     '2. Pick the package (pre-set packages at list price from the price book), or choose "Custom" and type the yearly costs of a quotation.\n'
     '3. Read Results. Recalculate (F9) if your spreadsheet tool does not recalculate on open.'),
    ('Status of the defaults', 'The default inputs describe a fictitious medium broker (60 users). Time per task, error rates, improvement rates and money figures '
     'are assumptions for illustration, not measured results of BrokerVerse. Cells with a light orange basis are assumptions that the broker must confirm, '
     'ideally from a short time study of the current process.'),
    ('Method', 'Hours saved = volume x current time per item x expected reduction. Cost saved = hours saved x hourly staff cost x the share of hours that turns '
     'into cash (avoided hiring, overtime or temporary staff). Rework, remittance charges, commission cash timing, commission leakage and renewal retention '
     'are added as separate lines. Year 1 counts only part of the benefit (go-live during the year and learning); volumes grow each year.'),
    ('Package costs', 'Package costs are the list prices of the pre-set packages in iNXT_BrokerVerse_Price_Book.xlsx for the tier reference user counts, copied as values when '
     'this workbook was built, by contract year from go-live (implementation and licence in Year 1). For any other configuration, use Custom with the Quick Quote figures.'),
    ('Colour key', 'Yellow cell with blue text: input. Light orange: assumption to confirm. Green: result.'),
]
head(rm, 4, ['Topic', 'Notes'])
for i, (k, v) in enumerate(notes):
    put(rm, 5 + i, 1, k, bold=True, wrap=True)
    put(rm, 5 + i, 2, v, wrap=True)
    rm.row_dimensions[5 + i].height = max(30, 15 * (len(v) // 120 + 1 + v.count('\n')))

# ------------------------------------------------------------------ Packages
pk = sheet('Packages', [58, 15, 15, 15, 15, 15, 16], 'Package costs by contract year (PHP, excl. VAT, list price)',
           'From the Pre-set Packages sheet of the price book. Custom: type the yearly costs of a specific quotation (Quick Quote, year by year table).')
head(pk, 4, ['Package', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total'])
r = 4
for p in PACKS:
    r += 1
    put(pk, r, 1, p['label'])
    for y in range(5):
        put(pk, r, 2 + y, round(p['years'][y], 2), PHP)
    put(pk, r, 7, f'=SUM(B{r}:F{r})', PHP, fill=F_TOT)
r += 1
put(pk, r, 1, 'Custom', bold=True)
for y in range(5):
    put(pk, r, 2 + y, [4970000, 2986800, 3098940, 3216687, 3340321][y], PHP, inp=True)
put(pk, r, 7, f'=SUM(B{r}:F{r})', PHP, fill=F_TOT)
PK0, PK1 = 5, r
name('PackNames', f'Packages!$A${PK0}:$A${PK1}')
name('PackCosts', f'Packages!$B${PK0}:$F${PK1}')
pk.freeze_panes = 'B5'

# ------------------------------------------------------------------ Inputs
ws = sheet('Inputs', [62, 16, 18, 70], 'Inputs: all assumptions are editable',
           'Yellow cells are inputs. Replace the example figures with the broker\'s own. Light orange basis = assumption to confirm.')
r = 3
IN = {}


def section(title):
    global r
    r += 1
    sec(ws, r, title, 4)
    r += 1
    head(ws, r, ['Input', 'Value', 'Unit', 'Basis'])


def inp(key, label, value, unit, fmt, basis, est=True):
    global r
    r += 1
    put(ws, r, 1, label)
    put(ws, r, 2, value, fmt, inp=True)
    put(ws, r, 3, unit)
    put(ws, r, 4, basis, wrap=True, fill=F_EST if est else None)
    name(key, f'Inputs!$B${r}')
    IN[key] = r


section('Package and broker profile')
inp('Pack', 'Package to compare against', 'Medium: Standard (subscription, AWS)', 'list', None,
    'Pick a pre-set package or "Custom" (Packages sheet).', est=False)
inp('Comm_Income', 'Commission income per year', 120000000, 'PHP', PHP, 'Example: medium broker estimate used in the pricing workbook (IC 2024 rank about 10 to 30).')
inp('Growth', 'Business volume growth per year', 0.05, '% a year', PCT, 'Assumption. Applied to volumes and money benefits from Year 2.')

section('Volumes per year')
VOL = [('V_Policy', 'New policies issued', 12000), ('V_Renew', 'Renewals processed', 8000), ('V_Endt', 'Endorsements', 3000),
       ('V_Claim', 'Claims notified and followed up', 1500), ('V_OR', 'Official receipts and collections posted', 20000)]
for k, lab, v in VOL:
    inp(k, lab, v, 'items a year', NUM, 'Example. Take from the current system or registers for the last 12 months.')
inp('V_Remit', 'Insurer remittance batches per month', 40, 'batches a month', NUM, 'Example: number of insurers x remittances per insurer per month.')
inp('V_Bank', 'Bank accounts reconciled each month', 6, 'accounts', NUM, 'Example.')

section('Staff cost')
inp('Staff_Cost', 'Fully loaded monthly cost of an operations or accounting staff member', 40000, 'PHP a month', PHP,
    'Assumption: salary, 13th month pay, statutory contributions and benefits. Use the broker\'s payroll average.')
inp('Hours_Month', 'Working hours per staff member per month', 176, 'hours', NUM, 'Assumption: 22 days x 8 hours.')
r += 1
put(ws, r, 1, 'Hourly staff cost', bold=True)
put(ws, r, 2, '=Staff_Cost/Hours_Month', PHP, fill=F_TOT)
put(ws, r, 3, 'PHP an hour')
name('Hour_Cost', f'Inputs!$B${r}')
inp('Cash_Share', 'Share of hours saved that becomes a cost saving', 0.6, '%', PCT,
    'Assumption: avoided hiring, overtime and temporary staff. The rest is capacity for growth and service, not counted as cash.')

section('Time on manual tasks today and expected reduction')
head(ws, r, ['Task', 'Minutes per item today', 'Expected reduction', 'Basis'])
TASKS = [
    ('T_Policy', 'Quotation, placement and policy issuance (per new policy)', 'V_Policy', 1, 60, 0.40,
     'Quotation, placement slip, policy and billing in one flow; product templates and motor tariff compute premium and taxes.'),
    ('T_Renew', 'Renewal processing (per renewal)', 'V_Renew', 1, 45, 0.50, 'Renewal lists and notices from expiring policies instead of manual tracking.'),
    ('T_Endt', 'Endorsement (per endorsement)', 'V_Endt', 1, 30, 0.35, 'Endorsement on the policy record with automatic premium and tax lines.'),
    ('T_Claim', 'Claim notification and follow-up (per claim)', 'V_Claim', 1, 90, 0.25, 'Claim record, documents and status in one place.'),
    ('T_OR', 'Official receipt and collection posting (per receipt)', 'V_OR', 1, 10, 0.50, 'Receipt against open items with automatic journal from posting rules.'),
    ('T_Remit', 'Remittance preparation and approval (per batch)', 'V_Remit', 12, 180, 0.60, 'Remittance built from collected premiums by insurer, with approval and ageing.'),
    ('T_Bank', 'Bank reconciliation (per account per month)', 'V_Bank', 12, 960, 0.60, 'Bank statement import and matching instead of manual ticking.'),
    ('T_Close', 'Month-end close and BIR working papers (per month)', None, 12, 4800, 0.50, 'Month-End Close run and BIR working papers (2307, VAT Summary, SAWT, QAP, SLSP) from posted data.'),
    ('T_Comm', 'Commission computation and statements (per month)', None, 12, 2400, 0.60, 'Commission from rate matrix and referrer sharing; statements from the system.'),
]
for k, lab, vol, mult, mins, red, basis in TASKS:
    r += 1
    put(ws, r, 1, lab)
    put(ws, r, 2, mins, NUM, inp=True)
    put(ws, r, 3, red, PCT, inp=True)
    put(ws, r, 4, 'Assumption. ' + basis, wrap=True, fill=F_EST)
    IN[k] = r
r += 1
put(ws, r, 4, 'Month-based tasks: minutes per month (4,800 minutes = 80 hours). Reductions are planning assumptions, to be confirmed by a time study.', wrap=True)

section('Errors and rework')
inp('Err_Rate', 'Transactions with an error that needs rework today (policies, endorsements, receipts)', 0.04, '% of transactions', PCT1, 'Assumption: wrong premium, tax, commission, client or insurer detail.')
inp('Err_Hours', 'Hours to find and correct one error', 1.5, 'hours', DEC, 'Assumption: includes reversal, re-issue and communication.')
inp('Err_Red', 'Expected reduction in errors', 0.5, '%', PCT, 'Assumption: validations, computed taxes, maker-checker approvals.')

section('Remittance and commission cash')
inp('Late_Charges', 'Late remittance charges, penalties or write-offs paid to insurers today', 300000, 'PHP a year', PHP, 'Example. Use the broker\'s records; 0 if none.')
inp('Late_Red', 'Expected reduction in late remittance charges', 0.7, '%', PCT, 'Assumption: remittance ageing on the insurer\'s terms and approval workflow.')
inp('Comm_Delay', 'Average delay in collecting commission from insurers today', 45, 'days', NUM, 'Example: reconciliation backlog with insurers.')
inp('Comm_Delay_Red', 'Expected reduction of that delay', 20, 'days', NUM, 'Assumption: insurer statement reconciliation in the system.')
inp('Cost_Money', 'Cost of money (borrowing rate or deposit rate forgone)', 0.06, '% a year', PCT1, 'Assumption.')
inp('Leak_Rate', 'Commission never collected or written off today', 0.005, '% of commission income', '0.0%', 'Assumption: commission not matched to insurer statements.')
inp('Leak_Red', 'Expected recovery of that leakage', 0.5, '%', PCT, 'Assumption.')

section('Renewal retention')
inp('Renew_Share', 'Share of commission income that comes from renewable policies', 0.7, '%', PCT, 'Assumption: non-life book renews yearly.')
inp('Ret_Now', 'Renewal retention rate today', 0.8, '%', PCT, 'Example. Renewed policies / expiring policies over the last 12 months.')
inp('Ret_Gain', 'Expected improvement in retention', 0.02, 'percentage points', PCT1, 'Assumption: renewal lists, reminders and follow-up on time. Net of new churn.')

section('Other costs avoided')
inp('Legacy', 'Current system costs that stop (licences, maintenance, spreadsheets support)', 0, 'PHP a year', PHP, 'Enter only costs that will actually stop after go-live.', est=False)

section('Realisation and evaluation')
inp('Real_Y1', 'Share of the full benefit realised in Year 1', 0.5, '%', PCT, 'Assumption: go-live during Year 1, then learning. Payback is measured from contract year 1.')
inp('Real_Y2', 'Share realised in Year 2', 0.9, '%', PCT, 'Assumption.')
inp('Real_Y3', 'Share realised from Year 3', 1.0, '%', PCT, 'Assumption.')
inp('Disc_Rate', 'Discount rate for net present value', 0.10, '% a year', PCT, 'Assumption: broker cost of capital.')

dv = DataValidation(type='list', formula1='PackNames', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"B{IN['Pack']}")
ws.freeze_panes = 'A4'

# ------------------------------------------------------------------ Results
rs = sheet('Results', [58, 17, 17, 17, 17, 17, 18], 'Results', 'PHP, exclusive of VAT. Formula driven from Inputs and Packages.')
r = 4
sec(rs, r, 'Hours saved by task (full year)', 7)
r += 1
head(rs, r, ['Task', 'Items a year', 'Hours a year today', 'Reduction', 'Hours saved a year', 'Cost of hours saved (PHP)', 'Cash saving (PHP)'], height=32)
h0 = r + 1
for k, lab, vol, mult, mins, red, basis in TASKS:
    r += 1
    ir = IN[k]
    put(rs, r, 1, f'=Inputs!A{ir}')
    put(rs, r, 2, f'={vol}*{mult}' if vol else f'={mult}', NUM)
    put(rs, r, 3, f'=B{r}*Inputs!B{ir}/60', NUM)
    put(rs, r, 4, f'=Inputs!C{ir}', PCT)
    put(rs, r, 5, f'=C{r}*D{r}', NUM)
    put(rs, r, 6, f'=E{r}*Hour_Cost', PHP)
    put(rs, r, 7, f'=F{r}*Cash_Share', PHP)
h1 = r
r += 1
put(rs, r, 1, 'Total', bold=True)
for c in (3, 5, 6, 7):
    put(rs, r, c, f'=SUM({L(c)}{h0}:{L(c)}{h1})', PHP, fill=F_TOT, bold=True)
TOT_H = r
r += 1
put(rs, r, 1, 'Full-time equivalents released (hours saved / yearly hours of one person)')
put(rs, r, 5, f'=E{TOT_H}/(Hours_Month*12)', DEC, fill=F_TOT)

r += 2
sec(rs, r, 'Yearly benefit at full run rate', 7)
r += 1
head(rs, r, ['Benefit', 'PHP a year', 'How it is computed'])
rs.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
BEN = [
    ('Staff time saved (cash share)', f'=G{TOT_H}', 'Hours saved x hourly cost x share that becomes a cost saving.'),
    ('Rework avoided', '=(V_Policy+V_Renew+V_Endt+V_OR)*Err_Rate*Err_Red*Err_Hours*Hour_Cost', 'Transactions x error rate x reduction x hours per error x hourly cost.'),
    ('Late remittance charges avoided', '=Late_Charges*Late_Red', 'Charges today x expected reduction.'),
    ('Commission collected earlier (cost of money)', '=Comm_Income*Comm_Delay_Red/365*Cost_Money', 'Commission income x days gained / 365 x cost of money.'),
    ('Commission leakage recovered', '=Comm_Income*Leak_Rate*Leak_Red', 'Commission income x leakage rate x recovery.'),
    ('Renewal retention gain', '=Comm_Income*Renew_Share*Ret_Gain', 'Renewable commission x improvement in retention (percentage points).'),
    ('Current system costs avoided', '=Legacy', 'As entered.'),
]
b0 = r + 1
for lab, f, how in BEN:
    r += 1
    put(rs, r, 1, lab)
    put(rs, r, 2, f, PHP)
    put(rs, r, 3, how)
    rs.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
b1 = r
r += 1
put(rs, r, 1, 'Total yearly benefit at full run rate', bold=True)
put(rs, r, 2, f'=SUM(B{b0}:B{b1})', PHP, fill=F_TOT, bold=True)
FULL = r
r += 1
put(rs, r, 1, 'Benefit as a share of commission income')
put(rs, r, 2, f'=B{FULL}/Comm_Income', PCT1)

r += 2
sec(rs, r, 'Benefit against the cost of the selected package', 7)
r += 1
put(rs, r, 1, 'Selected package', bold=True)
put(rs, r, 2, '=Pack', fill=F_TOT)
rs.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
r += 1
head(rs, r, ['Line', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total'])
y_hdr = r
rows = {}
for key, lab in [('real', 'Share of full benefit realised'), ('grow', 'Volume growth factor'), ('ben', 'Benefit'), ('cost', 'Package cost'),
                 ('net', 'Net benefit'), ('cum', 'Cumulative net benefit'), ('help', 'Payback helper (months)')]:
    r += 1
    rows[key] = r
    put(rs, r, 1, lab, bold=key in ('net', 'cum'))
for y in range(1, 6):
    c = L(1 + y)
    put(rs, rows['real'], 1 + y, '=Real_Y1' if y == 1 else ('=Real_Y2' if y == 2 else '=Real_Y3'), PCT)
    put(rs, rows['grow'], 1 + y, f'=(1+Growth)^({y}-1)', '0.000')
    put(rs, rows['ben'], 1 + y, f'=$B${FULL}*{c}{rows["real"]}*{c}{rows["grow"]}', PHP)
    put(rs, rows['cost'], 1 + y, f'=INDEX(PackCosts,MATCH(Pack,PackNames,0),{y})', PHP)
    put(rs, rows['net'], 1 + y, f'={c}{rows["ben"]}-{c}{rows["cost"]}', PHP, fill=F_TOT)
    prev = f'{L(y)}{rows["cum"]}' if y > 1 else '0'
    put(rs, rows['cum'], 1 + y, f'={prev}+{c}{rows["net"]}', PHP, fill=F_TOT)
    put(rs, rows['help'], 1 + y, f'=IF(AND({prev}<0,{c}{rows["cum"]}>=0),({y}-1)*12+12*(-{prev})/{c}{rows["net"]},IF(AND({y}=1,{c}{rows["cum"]}>=0),12*{c}{rows["cost"]}/{c}{rows["ben"]},""))', DEC)
for key in ('ben', 'cost', 'net'):
    put(rs, rows[key], 7, f'=SUM(B{rows[key]}:F{rows[key]})', PHP, fill=F_TOT, bold=True)
rs.row_dimensions[rows['help']].hidden = False
r += 2
sec(rs, r, 'Summary', 7)
summ = [
    ('Hours saved a year at full run rate', f'=E{TOT_H}', NUM),
    ('Full-time equivalents released', f'=E{TOT_H}/(Hours_Month*12)', DEC),
    ('Cost saved a year at full run rate (staff time and rework)', f'=B{b0}+B{b0+1}', PHP),
    ('Renewal retention gain a year', f'=B{b0+5}', PHP),
    ('Total yearly benefit at full run rate', f'=B{FULL}', PHP),
    ('5-year benefit', f'=G{rows["ben"]}', PHP),
    ('5-year package cost', f'=G{rows["cost"]}', PHP),
    ('5-year net benefit', f'=G{rows["net"]}', PHP),
    ('Return on investment over 5 years (net benefit / cost)', f'=IF(G{rows["cost"]}=0,"",G{rows["net"]}/G{rows["cost"]})', PCT),
    ('Net present value over 5 years', f'=NPV(Disc_Rate,B{rows["net"]}:F{rows["net"]})', PHP),
    ('Payback period (months from the start of contract year 1)', f'=IF(COUNT(B{rows["help"]}:F{rows["help"]})=0,"Not within 5 years",MIN(B{rows["help"]}:F{rows["help"]}))', DEC),
]
for lab, f, fmt in summ:
    r += 1
    put(rs, r, 1, lab, bold=True)
    put(rs, r, 2, f, fmt, fill=F_TOT, bold=True)
r += 2
rs.cell(r, 1, 'The results are estimates from the inputs. They are not a commitment by iorta TechNXT Corp. to any benefit; the broker validates the inputs.').font = ITAL
rs.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
rs.freeze_panes = 'A4'

wb._sheets = [wb[n] for n in ['Read Me', 'Inputs', 'Results', 'Packages']]
wb.active = 1
wb.save(OUT)
print('wrote', OUT)
