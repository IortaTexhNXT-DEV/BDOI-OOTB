"""Builds BrokerVerse_Commercials_and_Pricing.xlsx (formula driven)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

OUT = '/home/user/BDOI-OOTB/docs/package/out/BrokerVerse_Commercials_and_Pricing.xlsx'
ACC = '03 Oct 2026'

NAVY = '0F4761'
F_HEAD = PatternFill('solid', fgColor=NAVY)
F_SUB = PatternFill('solid', fgColor='D9E2F3')
F_IN = PatternFill('solid', fgColor='FFF2CC')
F_TOT = PatternFill('solid', fgColor='E2EFDA')
F_EST = PatternFill('solid', fgColor='FCE4D6')
WH = Font(name='Calibri', bold=True, color='FFFFFF')
B = Font(name='Calibri', bold=True)
BLUE = Font(name='Calibri', color='1F3FBF')
TITLE = Font(name='Calibri', bold=True, size=14, color=NAVY)
ITAL = Font(name='Calibri', italic=True, color='595959')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
PHP = '#,##0'
USD = '#,##0'
PCT = '0%'
PCT1 = '0.0%'
DEC = '#,##0.0'

wb = Workbook()


def sheet(name, widths, title, subtitle=None):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws['A1'] = title
    ws['A1'].font = TITLE
    if subtitle:
        ws['A2'] = subtitle
        ws['A2'].font = ITAL
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = 'BrokerVerse OOTB commercials, iorta TechNXT. Page &P of &N'
    return ws


def head(ws, r, labels, c0=1):
    for i, t in enumerate(labels):
        c = ws.cell(r, c0 + i, t)
        c.font = WH
        c.fill = F_HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BOX


def sec(ws, r, text, span=8):
    c = ws.cell(r, 1, text)
    c.font = Font(name='Calibri', bold=True, color=NAVY, size=12)
    for i in range(1, span + 1):
        ws.cell(r, i).fill = F_SUB


def put(ws, r, c, v, fmt=None, inp=False, bold=False, fill=None, wrap=False):
    cell = ws.cell(r, c, v)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    if inp:
        cell.fill = F_IN
        cell.font = BLUE
    if bold:
        cell.font = B
    if fill:
        cell.fill = fill
    if wrap:
        cell.alignment = WRAP
    return cell


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


wb.remove(wb.active)

# ------------------------------------------------------------------ Read Me
rm = sheet('Read Me', [26, 120], 'BrokerVerse OOTB: commercials and pricing workbook',
           'iorta TechNXT. Prepared 03 October 2026. Prices in PHP, exclusive of 12% VAT. For discussion with prospective broker customers.')
rows = [
    ('Purpose', 'Sets the commercial model for BrokerVerse OOTB (the out-of-the-box, as-is version) for Philippine non-life insurance brokers: '
     'Model A perpetual licence with Annual Maintenance Contract (AMC), Model B monthly subscription (SaaS), optional infrastructure, '
     'change request (CR) rates and optional services, with three worked quotes and the market evidence behind the price points.'),
    ('How to use', '1. Change drivers only on the Inputs sheet (yellow cells, blue text). Every other sheet is formula driven.\n'
     '2. To price a prospect, enter its user count and number of lines of business in Inputs (Calculator section). '
     'The calculators on Perpetual Model and Subscription Model update, with the break-even year.\n'
     '3. To change the three worked examples, edit the Sample quote scenarios table in Inputs.\n'
     '4. Infrastructure is optional and separate. Pick AWS, Azure or Local partner per scenario; the customer may host itself at no fee.\n'
     '5. Recalculate (F9) if your spreadsheet tool does not recalculate on open.'),
    ('Size tiers and user slabs', 'Small 1 to 25 users, Medium 26 to 100, Large 101 to 300, Enterprise above 300. '
     'User prices are graduated: the first 25 users are charged at the Small slab rate, users 26 to 100 at the Medium rate, and so on. '
     'This avoids a price drop when a broker crosses a tier boundary. The tier of the customer (by total users) sets the implementation effort, '
     'the minimum billable users for subscription, the included training days and the infrastructure size.'),
    ('Users', 'A user is a named person with a BrokerVerse sign-in (one role per user, as set up in Master > Generals > User Management > User). '
     'Inactive users are not counted. Read-only or customer-portal access is outside this price list.'),
    ('Model A: perpetual', 'One-time implementation fee by size and lines of business; one-time perpetual licence fee by user slab, payable on go-live; '
     'AMC at 22% of the licence fee from the AMC start year (default Year 2, after a 12-month warranty), increasing 5% a year. '
     'AMC covers corrections, product updates of the OOTB version, regulatory form updates released for all customers and standard support.'),
    ('Model B: subscription', 'Monthly price per user by slab, billed monthly in advance, increasing 5% at each contract anniversary. '
     'Minimum billable users apply per tier. The onboarding fee equals the implementation fee times the onboarding percentage in Inputs (default 100%). '
     'Subscription includes the licence right for the term, updates and standard support. Hosting is not included and is quoted separately.'),
    ('Implementation effort', 'Implementation = (base man-days for the tier + extra man-days for each line of business above the number included) x blended day rate, '
     'rounded to PHP 10,000. Each line of business needs product and cover set-up in the product configurator, rating and tax lines '
     '(documentary stamp tax, VAT or premium tax, local government tax, fire service tax where relevant), policy and endorsement document templates, '
     'insurer commission rates and UAT test cases. That is why the count of lines of business changes the effort.'),
    ('Taxes', 'All prices exclude 12% VAT, which is added on the invoice (Philippine VAT on services and on the sale or licensing of software). '
     'Customers who are withholding agents deduct creditable withholding tax as their status and the payment type require, '
     'commonly 2% for services and 1% to 2% on software or subscription fees under RR 2-98 as amended, and issue BIR Form 2307 to iorta TechNXT. '
     'Withholding does not reduce the invoice price. The rates are to be confirmed by the customer\'s and iorta TechNXT\'s tax advisers.'),
    ('Exchange rate', 'USD columns are for reference only and use the PHP per USD rate in Inputs (BSP reference rate 62.75 on 25 September 2026). '
     'The contract currency is PHP.'),
    ('Estimates and assumptions', 'Rows and cells shaded light orange, and figures marked "Estimate" or "Assumption", are not published facts. '
     'They are the basis stated beside them and must be confirmed before a binding quotation. '
     'Competitor prices in Market Benchmark are mostly third-party estimates because the vendors do not publish list prices.'),
    ('Validity', 'Prices are valid for 90 days from the quotation date and for contracts signed in that period. Day rates are fixed for 12 months from signing.'),
    ('Colour key', 'Yellow cell with blue text: input. Green: total. Light orange: estimate or assumption. Navy header: column titles.'),
    ('Sheets', 'Inputs, Size Tiers, Perpetual Model, Subscription Model, Infrastructure, Change Requests, Optional Services, Sample Quotes, Market Benchmark, Sources.'),
]
r = 4
head(rm, r, ['Topic', 'Notes'])
for k, v in rows:
    r += 1
    put(rm, r, 1, k, bold=True, wrap=True)
    put(rm, r, 2, v, wrap=True)
    rm.row_dimensions[r].height = max(30, 15 * (len(v) // 125 + 1 + v.count('\n')))
rm.freeze_panes = 'A5'

# ------------------------------------------------------------------ Inputs
ws = sheet('Inputs', [46, 13, 12, 14, 14, 12, 13, 12, 11, 12, 12, 14, 16, 44],
           'Inputs: all editable drivers', 'Change yellow cells only. Light orange notes mark estimates or assumptions.')
sec(ws, 4, 'General drivers', 14)
head(ws, 5, ['Driver', 'Value', 'Unit', 'Basis'])
ws.merge_cells('D5:N5')
gen = [
    ('FX', 'Exchange rate', 62.75, 'PHP per USD', '0.00', 'BSP reference rate, 25 September 2026 (source S14). Reference only.'),
    ('FX_Date', 'Exchange rate date', '25 Sep 2026', 'date', None, 'Update with the rate on the quotation date.'),
    ('VAT', 'VAT rate', 0.12, '%', PCT, 'Philippine VAT, added to all prices (prices exclusive of VAT).'),
    ('AMC_Rate', 'AMC rate (of perpetual licence fee)', 0.22, '%', PCT, 'Business owner instruction.'),
    ('Esc', 'Annual escalation (AMC, subscription, support)', 0.05, '%', PCT, 'Business owner instruction. Applied at each anniversary.'),
    ('AMC_Start', 'AMC starts in contract year', 2, 'year', '0', 'Assumption: Year 1 after go-live is covered by a 12-month warranty. Set to 1 to bill AMC from go-live.'),
    ('Onboard_Pct', 'Subscription onboarding fee as % of implementation fee', 1.0, '%', PCT, 'Assumption: same set-up work in both models.'),
    ('Infra_Margin', 'Infrastructure margin on cloud cost', 0.20, '%', PCT, 'Assumption: covers FX movement, monitoring, patching and backup checks.'),
    ('Infra_Esc', 'Infrastructure annual change', 0.0, '%', PCT, 'Assumption: cloud list prices are flat; review yearly against FX.'),
    ('Days_per_LOB', 'Extra implementation man-days per line of business above the included count', 6, 'man-days', '0', 'Estimate: product and cover set-up, rating and tax lines, 2 to 3 document templates, insurer commission rates, UAT cases.'),
    ('Round_To', 'Round implementation and service fees to', 10000, 'PHP', PHP, 'Presentation.'),
    ('IT_Low', 'IT budget benchmark, low (share of revenue)', 0.03, '%', PCT1, 'Gartner insurance IT spend 3.1% of revenue (2010); later benchmarks 4% to 7% (sources S10, S11).'),
    ('IT_High', 'IT budget benchmark, high (share of revenue)', 0.07, '%', PCT1, 'As above.'),
    ('Validity', 'Quotation validity', 90, 'days', '0', 'Commercial policy.'),
]
r = 5
for n, lab, v, unit, fmt, basis in gen:
    r += 1
    put(ws, r, 1, lab)
    put(ws, r, 2, v, fmt, inp=True)
    put(ws, r, 3, unit)
    put(ws, r, 4, basis)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
    if basis.startswith(('Assumption', 'Estimate')):
        ws.cell(r, 4).fill = F_EST
    name(n, f"Inputs!$B${r}")

r += 2
sec(ws, r, 'Calculator (used by the calculators on Perpetual Model and Subscription Model)', 14)
r += 1
head(ws, r, ['Driver', 'Value', 'Unit', 'Basis'])
for n, lab, v, unit in [('Calc_Users', 'Number of named users', 40, 'users'), ('Calc_LOBs', 'Number of lines of business to configure', 6, 'lines')]:
    r += 1
    put(ws, r, 1, lab)
    put(ws, r, 2, v, '0', inp=True)
    put(ws, r, 3, unit)
    put(ws, r, 4, 'Enter the prospect figures.')
    name(n, f"Inputs!$B${r}")

r += 2
sec(ws, r, 'Size tiers, user slabs and tier parameters', 14)
r += 1
head(ws, r, ['Tier', 'From users', 'To users', 'Perpetual licence per user (PHP, slab)', 'Subscription per user per month (PHP, slab)',
             'Minimum billable users (subscription)', 'Implementation base man-days', 'Lines of business included',
             'Training days included', 'Reference users (5-year view)', 'Reference lines of business',
             '24x7 Severity 1 support, annual (PHP)', 'Indicative annual commission income (PHP, estimate)', 'Basis'])
ws.row_dimensions[r].height = 60
TIER_HDR = r
tiers = [
    ('Small', 1, 25, 90000, 3200, 10, 70, 5, 4, 15, 5, 240000, 25000000,
     'Commission income estimate: IC 2024 ranks 30 and below (rank 30 earned PHP 55.2 million). Price slab: see Market Benchmark.'),
    ('Medium', 26, 100, 78000, 2800, 26, 130, 8, 6, 60, 8, 480000, 120000000,
     'IC 2024 ranks about 10 to 30 (rank 20 earned PHP 98.0 million).'),
    ('Large', 101, 300, 66000, 2400, 101, 230, 12, 10, 180, 12, 900000, 500000000,
     'IC 2024 ranks 4 to 10 (Lockton PHP 712 million, PhilPacific PHP 598 million).'),
    ('Enterprise', 301, 100000, 54000, 2000, 301, 360, 16, 15, 400, 16, 1500000, 1500000000,
     'IC 2024 top 3 earned PHP 2.0 to 2.4 billion each (BDO, Marsh, Aon).'),
]
T0 = r + 1
for t in tiers:
    r += 1
    put(ws, r, 1, t[0], bold=True)
    for j, v in enumerate(t[1:13], 2):
        fmt = PHP if j in (4, 5, 12, 13) else '0'
        put(ws, r, j, v, fmt, inp=True)
    put(ws, r, 14, t[13], wrap=True, fill=F_EST)
    ws.row_dimensions[r].height = 30
T1 = r
TR = lambda col: f"Inputs!${col}${T0}:${col}${T1}"  # noqa: E731
r += 1
put(ws, r, 1, 'Implementation base man-days, lines of business and training days are estimates for the OOTB version (configuration, data load with the standard templates, UAT support, training, go-live and hypercare). Commission income is an estimate used only for the affordability check.', wrap=True, fill=F_EST)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
ws.row_dimensions[r].height = 30

r += 2
sec(ws, r, 'Day rates (PHP per man-day of 8 hours)', 14)
r += 1
head(ws, r, ['Role', 'Day rate (PHP)', 'Mix for blended rate', 'USD per day', 'Basis'])
ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=14)
roles = [('Business analyst', 18000, 0.25, 'Estimate. PH senior BA salary about PHP 70,500 a month (S17); vendor rate covers benefits, overhead, bench and margin.'),
         ('Developer', 16000, 0.40, 'Estimate. PH outsourcing market USD 30 to 50 an hour mid-level, 50 to 75 senior (S15, S16); PHP 16,000 = USD 255 a day.'),
         ('QA engineer', 12000, 0.25, 'Estimate. PH senior QA salary about PHP 65,000 a month (S18).'),
         ('Project manager', 22000, 0.10, 'Estimate. PH IT project manager salary PHP 60,000 to 123,000 a month (S19).')]
R0 = r + 1
for role, rate, mix, basis in roles:
    r += 1
    put(ws, r, 1, role)
    put(ws, r, 2, rate, PHP, inp=True)
    put(ws, r, 3, mix, PCT, inp=True)
    put(ws, r, 4, f'=B{r}/FX', USD)
    put(ws, r, 5, basis, fill=F_EST)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=14)
R1 = r
ROLE_ROWS = list(range(R0, R1 + 1))
name('Rate_BA', f'Inputs!$B${R0}')
name('Rate_Dev', f'Inputs!$B${R0+1}')
name('Rate_QA', f'Inputs!$B${R0+2}')
name('Rate_PM', f'Inputs!$B${R0+3}')
r += 1
put(ws, r, 1, 'Blended day rate (weighted by mix)', bold=True)
put(ws, r, 2, f'=ROUND(SUMPRODUCT(B{R0}:B{R1},C{R0}:C{R1})/SUM(C{R0}:C{R1}),-2)', PHP, fill=F_TOT)
put(ws, r, 3, f'=SUM(C{R0}:C{R1})', PCT)
put(ws, r, 4, f'=B{r}/FX', USD)
name('Blended', f'Inputs!$B${r}')
r += 1
put(ws, r, 1, 'Trainer day rate (end-user or train-the-trainer)')
put(ws, r, 2, 30000, PHP, inp=True)
put(ws, r, 4, f'=B{r}/FX', USD)
put(ws, r, 5, 'Assumption: one trainer, up to 15 participants, preparation included.', fill=F_EST)
ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=14)
name('Rate_Trainer', f'Inputs!$B${r}')
r += 1
put(ws, r, 1, 'On-site consultant day rate (Metro Manila)')
put(ws, r, 2, 20000, PHP, inp=True)
put(ws, r, 4, f'=B{r}/FX', USD)
put(ws, r, 5, 'Assumption: travel, lodging and meals outside Metro Manila billed at cost.', fill=F_EST)
ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=14)
name('Rate_Onsite', f'Inputs!$B${r}')

r += 2
sec(ws, r, 'Change request estimation ratios', 14)
r += 1
head(ws, r, ['Driver', 'Value', 'Unit', 'Basis'])
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
for n, lab, v, unit, fmt, basis in [
    ('CR_BA', 'Business analysis effort as % of development', 0.30, '%', PCT, 'Assumption: requirement note, impact analysis, UAT support.'),
    ('CR_QA', 'Testing effort as % of development', 0.40, '%', PCT, 'Assumption: test cases, regression of affected screens and reports.'),
    ('CR_PM', 'Project management as % of BA + development + testing', 0.10, '%', PCT, 'Assumption.'),
    ('CR_Cont', 'Contingency', 0.10, '%', PCT, 'Assumption: unknowns found during analysis.'),
    ('CR_Small_Max', 'Small CR band: total man-days up to', 5, 'man-days', '0', 'Commercial policy.'),
    ('CR_Med_Max', 'Medium CR band: total man-days up to', 20, 'man-days', '0', 'Commercial policy.'),
    ('CR_Large_Max', 'Large CR band: total man-days up to (above this, a separate project)', 60, 'man-days', '0', 'Commercial policy.'),
    ('CR_Dev_Days', 'CR calculator: development man-days estimated', 8, 'man-days', '0.0', 'Enter the developer estimate for the change.')]:
    r += 1
    put(ws, r, 1, lab)
    put(ws, r, 2, v, fmt, inp=True)
    put(ws, r, 3, unit)
    put(ws, r, 4, basis, fill=F_EST if basis.startswith('Assumption') else None)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
    name(n, f'Inputs!$B${r}')

r += 2
sec(ws, r, 'Optional services drivers', 14)
r += 1
head(ws, r, ['Driver', 'Value', 'Unit', 'Basis'])
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
for n, lab, v, unit, fmt, basis in [
    ('Env_Setup_Days', 'Additional environment: set-up man-days', 6, 'man-days', '0', 'Estimate: provisioning, deployment, configuration copy, smoke test.'),
    ('Env_Monthly_USD', 'Additional non-production environment: cloud cost per month', 150, 'USD', USD, 'Estimate: 1 small container, small single-zone database, storage (S3 for compute and database pricing references).'),
    ('Mig_Days_Source', 'Data migration beyond templates: man-days per additional legacy source', 15, 'man-days', '0', 'Estimate: mapping, extraction scripts, 2 trial loads, reconciliation.'),
    ('Int_Std_Days', 'Additional integration, standard (documented API or file)', 20, 'man-days', '0', 'Estimate: e.g. one insurer API, SMS gateway, payment gateway variant.'),
    ('Int_Cx_Days', 'Additional integration, complex (no API, two-way, or batch reconciliation)', 45, 'man-days', '0', 'Estimate.'),
    ('Report_Days', 'New report or document template (typical)', 4, 'man-days', '0', 'Estimate: falls in the Small CR band.')]:
    r += 1
    put(ws, r, 1, lab)
    put(ws, r, 2, v, fmt, inp=True)
    put(ws, r, 3, unit)
    put(ws, r, 4, basis, fill=F_EST)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
    name(n, f'Inputs!$B${r}')

r += 2
sec(ws, r, 'Infrastructure: AWS reference cost per month (USD, production plus one UAT environment)', 14)
r += 1
head(ws, r, ['Component', 'Small', 'Medium', 'Large', 'Enterprise', 'Basis'])
ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
infra = [
    ('Application containers (API, web, scheduler)', 90, 135, 180, 360, 'Estimate: 2/3/4/8 containers of 1 vCPU, 2 GB (Fargate, Singapore).'),
    ('Managed PostgreSQL 16 (instance and storage)', 85, 320, 640, 1280, 'Estimate: db.t4g.medium single-AZ / db.m7g.large Multi-AZ / m7g.xlarge Multi-AZ / m7g.2xlarge Multi-AZ (S20).'),
    ('Object storage and backups (documents, snapshots)', 20, 40, 90, 180, 'Estimate: S3 and snapshot storage, growing with policies and uploads.'),
    ('Network (load balancer, CDN, data transfer)', 45, 70, 120, 220, 'Estimate.'),
    ('Security and monitoring (WAF, logs, alarms, secrets, keys)', 35, 60, 110, 200, 'Estimate.'),
    ('UAT environment (scaled down)', 120, 200, 350, 600, 'Estimate.'),
    ('Disaster recovery copy (cross-region backups)', 0, 0, 0, 400, 'Estimate: Enterprise only; optional for others.'),
]
I0 = r + 1
for comp in infra:
    r += 1
    put(ws, r, 1, comp[0])
    for j in range(4):
        put(ws, r, 2 + j, comp[1 + j], USD, inp=True)
    put(ws, r, 6, comp[5], fill=F_EST)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
I1 = r
r += 1
put(ws, r, 1, 'Total AWS reference cost per month (USD)', bold=True)
for j in range(4):
    col = L(2 + j)
    put(ws, r, 2 + j, f'=SUM({col}{I0}:{col}{I1})', USD, fill=F_TOT)
INFRA_TOT = r
r += 1
put(ws, r, 1, 'Azure cost relative to AWS')
put(ws, r, 2, 1.05, '0.00', inp=True)
put(ws, r, 6, 'Assumption: comparable list prices in Southeast Asia (D4s v5 Linux about USD 158 a month, S21); +5% for managed PostgreSQL and WAF differences.', fill=F_EST)
ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
name('Azure_Factor', f'Inputs!$B${r}')
r += 1
put(ws, r, 1, 'Local partner cloud cost relative to AWS')
put(ws, r, 2, 0.95, '0.00', inp=True)
put(ws, r, 6, 'Assumption: Philippine data-centre or partner cloud priced in PHP; lower compute and bandwidth, database run as managed VM. Confirm with a partner quotation.', fill=F_EST)
ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
name('Local_Factor', f'Inputs!$B${r}')

r += 2
sec(ws, r, 'Sample quote scenarios (used by Sample Quotes)', 14)
r += 1
head(ws, r, ['Driver', 'Example 1', 'Example 2', 'Example 3', '', 'Notes'])
ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
scen = [
    ('Broker profile', 'Small broker', 'Medium broker', 'Large broker', None, 'Label only.'),
    ('Named users', 15, 60, 180, '0', 'Tier follows from the user count.'),
    ('Lines of business', 5, 8, 12, '0', 'Motor, fire, marine, casualty, engineering, bonds and so on.'),
    ('Infrastructure provider (AWS, Azure, Local partner, Customer)', 'AWS', 'Azure', 'Local partner', None, '"Customer" means the broker hosts itself: no infrastructure fee.'),
    ('24x7 Severity 1 support (Yes/No)', 'No', 'Yes', 'Yes', None, 'Optional.'),
    ('Extra training days', 0, 2, 4, '0', 'Beyond the days included for the tier.'),
    ('Extra legacy data sources to migrate', 0, 1, 2, '0', 'Beyond the standard templates.'),
    ('Additional standard integrations', 0, 1, 2, '0', 'One-time.'),
]
S0 = r + 1
for lab, a, b2, c3, fmt, note in scen:
    r += 1
    put(ws, r, 1, lab)
    for j, v in enumerate((a, b2, c3)):
        put(ws, r, 2 + j, v, fmt, inp=True)
    put(ws, r, 6, note)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
SC = {k: S0 + i for i, k in enumerate(['label', 'users', 'lobs', 'prov', 'sup', 'train', 'mig', 'int'])}
dv = DataValidation(type='list', formula1='"AWS,Azure,Local partner,Customer"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f'B{SC["prov"]}:D{SC["prov"]}')
dv2 = DataValidation(type='list', formula1='"Yes,No"', allow_blank=False)
ws.add_data_validation(dv2)
dv2.add(f'B{SC["sup"]}:D{SC["sup"]}')
ws.freeze_panes = 'B4'


# helpers -----------------------------------------------------------------
def slab_sum(N, pricecol):
    return '+'.join(f'MAX(0,MIN({N},Inputs!$C${t})-Inputs!$B${t}+1)*Inputs!${pricecol}${t}' for t in range(T0, T1 + 1))


def tidx(N):
    return f'MATCH({N},{TR("B")},1)'


def tget(col, N):
    return f'INDEX({TR(col)},{tidx(N)})'


def impl(N, LOB):
    return f'ROUND(({tget("G", N)}+MAX(0,{LOB}-{tget("H", N)})*Days_per_LOB)*Blended/Round_To,0)*Round_To'


def billable(N):
    return f'MAX({N},{tget("F", N)})'


YC = ['C', 'D', 'E', 'F', 'G']  # year columns


def amc(lic, y):
    return f'=IF({y}>=AMC_Start,{lic}*AMC_Rate*(1+Esc)^({y}-AMC_Start),0)'


# ------------------------------------------------------------------ Size Tiers
st = sheet('Size Tiers', [16, 12, 12, 34, 30, 22, 20, 20, 18, 18], 'Size tiers and user slabs',
           'Tier definitions come from Inputs. Profile columns are estimates from the Insurance Commission broker ranking (2024) and market practice.')
head(st, 4, ['Tier', 'From users', 'To users', 'Typical broker profile (estimate)', 'IC 2024 commission ranking position (evidence)',
             'Perpetual licence per user, slab (PHP)', 'Subscription per user per month, slab (PHP)', 'Minimum billable users',
             'Lines of business included', 'Training days included'])
st.row_dimensions[4].height = 45
prof = [('Single office or 2 to 3 branches; 10 to 40 staff; SME, motor and fire focus; owner-managed finance.', 'Rank 30 to 67: commission up to about PHP 55 million (rank 30, Moneyhero, PHP 55.2 million).'),
        ('Several branches; 40 to 150 staff; corporate and affinity accounts; in-house accounting team.', 'Rank 10 to 30: PHP 55 to 250 million (rank 20, East West, PHP 98.0 million).'),
        ('National branch network or multinational; 150 to 400 staff; large corporate and reinsurance placements.', 'Rank 4 to 10: PHP 250 to 800 million (Lockton PHP 712 million; PhilPacific PHP 598 million).'),
        ('Bank-affiliated or global broker; over 400 staff; high transaction volumes and bancassurance.', 'Top 3: PHP 2.0 to 2.4 billion (BDO PHP 2,434 million; Marsh PHP 2,075 million; Aon PHP 2,047 million).')]
for i in range(4):
    rr = 5 + i
    t = T0 + i
    put(st, rr, 1, f'=Inputs!A{t}', bold=True)
    put(st, rr, 2, f'=Inputs!B{t}', '0')
    put(st, rr, 3, f'=IF(Inputs!C{t}>=100000,"and above",Inputs!C{t})', '0')
    put(st, rr, 4, prof[i][0], wrap=True, fill=F_EST)
    put(st, rr, 5, prof[i][1], wrap=True)
    put(st, rr, 6, f'=Inputs!D{t}', PHP)
    put(st, rr, 7, f'=Inputs!E{t}', PHP)
    put(st, rr, 8, f'=Inputs!F{t}', '0')
    put(st, rr, 9, f'=Inputs!H{t}', '0')
    put(st, rr, 10, f'=Inputs!I{t}', '0')
    st.row_dimensions[rr].height = 48
r = 10
sec(st, r, 'How the graduated slabs work', 10)
st.cell(r + 1, 1, 'Each user is priced at the rate of the slab it falls in. Example: 40 users = 25 users at the Small rate + 15 users at the Medium rate. '
        'The broker tier (by total users) sets implementation effort, minimum billable users, included training and infrastructure size.').alignment = WRAP
st.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=10)
st.row_dimensions[r + 1].height = 32
r = 13
sec(st, r, 'Affordability check against IT budget norms (reference users per tier, Year 1 prices)', 10)
head(st, r + 1, ['Tier', 'Reference users', 'Commission income (PHP, estimate)', 'IT budget at low benchmark (PHP)',
                 'IT budget at high benchmark (PHP)', 'Subscription Year 1 (PHP)', 'Subscription as % of income',
                 'Perpetual 5-year average per year incl. implementation (PHP)', 'Perpetual average as % of income', 'Within IT budget?'])
st.row_dimensions[r + 1].height = 60
for i in range(4):
    rr = r + 2 + i
    t = T0 + i
    put(st, rr, 1, f'=Inputs!A{t}', bold=True)
    put(st, rr, 2, f'=Inputs!J{t}', '0')
    put(st, rr, 3, f'=Inputs!M{t}', PHP, fill=F_EST)
    put(st, rr, 4, f'=C{rr}*IT_Low', PHP)
    put(st, rr, 5, f'=C{rr}*IT_High', PHP)
    put(st, rr, 6, f"='Subscription Model'!C{8 + i * 8}", PHP)  # filled in later block layout
    put(st, rr, 7, f'=F{rr}/C{rr}', PCT1)
    put(st, rr, 8, f"='Perpetual Model'!H{9 + i * 8}/5", PHP)
    put(st, rr, 9, f'=H{rr}/C{rr}', PCT1)
    put(st, rr, 10, f'=IF(MAX(G{rr},I{rr})<=IT_High,"Yes","Review")')
st.cell(r + 7, 1, 'Note: the IT budget also pays for hardware, network, office systems and staff. BrokerVerse should take well under the full budget; '
        'the check shows the software cost against the benchmark band, not a recommendation of spend.').font = ITAL
st.merge_cells(start_row=r + 7, start_column=1, end_row=r + 7, end_column=10)
st.freeze_panes = 'B5'

# ------------------------------------------------------------------ Perpetual Model
pm = sheet('Perpetual Model', [30, 26, 15, 15, 15, 15, 15, 17, 15], 'Model A: perpetual licence with AMC',
           'Implementation and licence in Year 1; AMC at 22% of the licence fee from the AMC start year, +5% a year. PHP, exclusive of VAT unless stated.')
head(pm, 4, ['Tier / item', 'Basis', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total (PHP)', '5-year total (USD)'])
PM_TOT = {}
for i in range(4):
    t = T0 + i
    b0 = 5 + i * 8
    N, LOB = f'$C${b0}', f'$D${b0}'
    put(pm, b0, 1, f'=Inputs!A{t}&" (reference)"', bold=True, fill=F_SUB)
    put(pm, b0, 2, 'Users / lines of business', fill=F_SUB)
    put(pm, b0, 3, f'=Inputs!J{t}', '0', fill=F_SUB)
    put(pm, b0, 4, f'=Inputs!K{t}', '0', fill=F_SUB)
    for c in range(5, 10):
        pm.cell(b0, c).fill = F_SUB
    items = [('Implementation fee', 'Base + extra lines x blended rate'),
             ('Perpetual licence fee', 'Graduated user slabs'),
             ('AMC', '22% of licence, +5% a year')]
    for k, (lab, basis) in enumerate(items):
        rr = b0 + 1 + k
        put(pm, rr, 1, lab)
        put(pm, rr, 2, basis)
        for y, col in enumerate(YC, 1):
            if k == 0:
                f = f'={impl(N, LOB)}' if y == 1 else 0
            elif k == 1:
                f = f'={slab_sum(N, "D")}' if y == 1 else 0
            else:
                f = amc(f'$C${b0 + 2}', y)
            put(pm, rr, 3 + y - 1, f, PHP)
        put(pm, rr, 8, f'=SUM(C{rr}:G{rr})', PHP)
        put(pm, rr, 9, f'=H{rr}/FX', USD)
    rr = b0 + 4
    put(pm, rr, 1, 'Total excl. VAT', bold=True, fill=F_TOT)
    put(pm, rr, 2, '', fill=F_TOT)
    for c in range(3, 9):
        col = L(c)
        put(pm, rr, c, f'=SUM({col}{b0+1}:{col}{b0+3})', PHP, bold=True, fill=F_TOT)
    put(pm, rr, 9, f'=H{rr}/FX', USD, fill=F_TOT)
    PM_TOT[i] = rr
    rr = b0 + 5
    put(pm, rr, 1, 'VAT')
    put(pm, rr, 2, '=TEXT(VAT,"0%")')
    for c in range(3, 9):
        put(pm, rr, c, f'={L(c)}{b0+4}*VAT', PHP)
    put(pm, rr, 9, f'=H{rr}/FX', USD)
    rr = b0 + 6
    put(pm, rr, 1, 'Total incl. VAT', bold=True)
    put(pm, rr, 2, '')
    for c in range(3, 9):
        put(pm, rr, c, f'={L(c)}{b0+4}+{L(c)}{b0+5}', PHP, bold=True)
    put(pm, rr, 9, f'=H{rr}/FX', USD)
# calculator
c0 = 5 + 4 * 8 + 1
sec(pm, c0, 'Calculator for a chosen user count (change users and lines of business in Inputs, Calculator section)', 9)
r = c0 + 1
put(pm, r, 1, 'Named users')
put(pm, r, 3, '=Calc_Users', '0', fill=F_IN)
put(pm, r + 1, 1, 'Lines of business')
put(pm, r + 1, 3, '=Calc_LOBs', '0', fill=F_IN)
put(pm, r + 2, 1, 'Tier')
put(pm, r + 2, 3, f'=INDEX({TR("A")},{tidx(f"C{r}")})', bold=True)
CN, CL = f'$C${r}', f'$C${r+1}'
r += 4
head(pm, r, ['Slab', 'Users from - to', 'Users in slab', 'Price per user (PHP)', 'Amount (PHP)'])
s0 = r + 1
for i in range(4):
    t = T0 + i
    rr = r + 1 + i
    put(pm, rr, 1, f'=Inputs!A{t}')
    put(pm, rr, 2, f'=Inputs!B{t}&" - "&IF(Inputs!C{t}>=100000,"above",Inputs!C{t})')
    put(pm, rr, 3, f'=MAX(0,MIN({CN},Inputs!C{t})-Inputs!B{t}+1)', '0')
    put(pm, rr, 4, f'=Inputs!D{t}', PHP)
    put(pm, rr, 5, f'=C{rr}*D{rr}', PHP)
r = r + 5
put(pm, r, 1, 'Perpetual licence fee', bold=True)
put(pm, r, 5, f'=SUM(E{s0}:E{s0+3})', PHP, bold=True, fill=F_TOT)
LICC = f'$E${r}'
put(pm, r + 1, 1, 'Effective licence price per user')
put(pm, r + 1, 5, f'=E{r}/{CN}', PHP)
put(pm, r + 2, 1, 'Implementation fee')
put(pm, r + 2, 2, 'Base man-days for tier + extra lines')
put(pm, r + 2, 5, f'={impl(CN, CL)}', PHP)
IMPC = f'$E${r+2}'
put(pm, r + 3, 1, 'Implementation man-days (before rounding)')
put(pm, r + 3, 5, f'={tget("G", CN)}+MAX(0,{CL}-{tget("H", CN)})*Days_per_LOB', '0')
r += 5
head(pm, r, ['Item', 'Basis', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total (PHP)', '5-year total (USD)'])
k0 = r + 1
for k, lab in enumerate(['Implementation fee', 'Perpetual licence fee', 'AMC']):
    rr = k0 + k
    put(pm, rr, 1, lab)
    for y, col in enumerate(YC, 1):
        if k == 0:
            f = f'={IMPC}' if y == 1 else 0
        elif k == 1:
            f = f'={LICC}' if y == 1 else 0
        else:
            f = amc(LICC, y)
        put(pm, rr, 2 + y, f, PHP)
    put(pm, rr, 8, f'=SUM(C{rr}:G{rr})', PHP)
    put(pm, rr, 9, f'=H{rr}/FX', USD)
rr = k0 + 3
put(pm, rr, 1, 'Total excl. VAT', bold=True, fill=F_TOT)
for c in range(3, 9):
    put(pm, rr, c, f'=SUM({L(c)}{k0}:{L(c)}{k0+2})', PHP, bold=True, fill=F_TOT)
put(pm, rr, 9, f'=H{rr}/FX', USD, fill=F_TOT)
PM_CALC_TOT = rr
put(pm, rr + 1, 1, 'VAT')
put(pm, rr + 2, 1, 'Total incl. VAT', bold=True)
for c in range(3, 9):
    put(pm, rr + 1, c, f'={L(c)}{rr}*VAT', PHP)
    put(pm, rr + 2, c, f'={L(c)}{rr}+{L(c)}{rr+1}', PHP, bold=True)
put(pm, rr + 1, 9, f'=H{rr+1}/FX', USD)
put(pm, rr + 2, 9, f'=H{rr+2}/FX', USD)
pm.freeze_panes = 'C5'

# ------------------------------------------------------------------ Subscription Model
sm = sheet('Subscription Model', [32, 26, 15, 15, 15, 15, 15, 17, 15], 'Model B: subscription (SaaS, as-is version)',
           'Monthly per-user price by graduated slab, billed monthly in advance, +5% at each anniversary. Hosting quoted separately. PHP, exclusive of VAT unless stated.')
head(sm, 4, ['Tier / item', 'Basis', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total (PHP)', '5-year total (USD)'])
SM_TOT = {}
for i in range(4):
    t = T0 + i
    b0 = 5 + i * 8
    N = f'$C${b0}'
    put(sm, b0, 1, f'=Inputs!A{t}&" (reference)"', bold=True, fill=F_SUB)
    put(sm, b0, 2, 'Users / billable users', fill=F_SUB)
    put(sm, b0, 3, f'=Inputs!J{t}', '0', fill=F_SUB)
    put(sm, b0, 4, f'={billable(N)}', '0', fill=F_SUB)
    for c in range(5, 10):
        sm.cell(b0, c).fill = F_SUB
    # monthly fee row (info) b0+1
    put(sm, b0 + 1, 1, 'Monthly fee (Year 1 rates)')
    put(sm, b0 + 1, 2, 'Graduated slabs on billable users')
    put(sm, b0 + 1, 3, f'={slab_sum(f"$D${b0}", "E")}', PHP)
    put(sm, b0 + 1, 4, 'per user:')
    put(sm, b0 + 1, 5, f'=C{b0+1}/$D${b0}', PHP)
    # onboarding b0+2
    put(sm, b0 + 2, 1, 'Onboarding fee')
    put(sm, b0 + 2, 2, 'Implementation x onboarding %')
    for y, col in enumerate(YC, 1):
        put(sm, b0 + 2, 2 + y, f"='Perpetual Model'!C{b0+1}*Onboard_Pct" if y == 1 else 0, PHP)
    put(sm, b0 + 2, 8, f'=SUM(C{b0+2}:G{b0+2})', PHP)
    put(sm, b0 + 2, 9, f'=H{b0+2}/FX', USD)
    put(sm, b0 + 3, 1, 'Subscription fees (12 months)')
    put(sm, b0 + 3, 2, 'Monthly fee x 12, +5% a year')
    for y, col in enumerate(YC, 1):
        put(sm, b0 + 3, 2 + y, f'=$C${b0+1}*12*(1+Esc)^({y}-1)', PHP)
    put(sm, b0 + 3, 8, f'=SUM(C{b0+3}:G{b0+3})', PHP)
    put(sm, b0 + 3, 9, f'=H{b0+3}/FX', USD)
    rr = b0 + 4
    put(sm, rr, 1, 'Total excl. VAT', bold=True, fill=F_TOT)
    put(sm, rr, 2, '', fill=F_TOT)
    for c in range(3, 9):
        put(sm, rr, c, f'={L(c)}{b0+2}+{L(c)}{b0+3}', PHP, bold=True, fill=F_TOT)
    put(sm, rr, 9, f'=H{rr}/FX', USD, fill=F_TOT)
    SM_TOT[i] = rr
    put(sm, rr + 1, 1, 'VAT')
    put(sm, rr + 2, 1, 'Total incl. VAT', bold=True)
    for c in range(3, 9):
        put(sm, rr + 1, c, f'={L(c)}{rr}*VAT', PHP)
        put(sm, rr + 2, c, f'={L(c)}{rr}+{L(c)}{rr+1}', PHP, bold=True)
    put(sm, rr + 1, 9, f'=H{rr+1}/FX', USD)
    put(sm, rr + 2, 9, f'=H{rr+2}/FX', USD)
# Fix Size Tiers references: subscription Year 1 fees row = b0+3, perpetual total row = b0+4
for i in range(4):
    b0 = 5 + i * 8
    st.cell(15 + i, 6).value = f"='Subscription Model'!C{b0+3}"
    st.cell(15 + i, 8).value = f"='Perpetual Model'!H{b0+4}/5"

# break-even
be0 = 5 + 4 * 8 + 1
sec(sm, be0, 'Break-even against the perpetual model (cumulative cost including implementation or onboarding, excl. VAT)', 9)
r = be0 + 1
head(sm, r, ['Tier / item', '', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', 'Break-even', 'Saving of perpetual over 5 years'])
for i in range(4):
    t = T0 + i
    b0 = 5 + i * 8
    rr = r + 1 + i * 5
    put(sm, rr, 1, f'=Inputs!A{t}', bold=True, fill=F_SUB)
    for c in range(2, 10):
        sm.cell(rr, c).fill = F_SUB
    put(sm, rr + 1, 1, 'Perpetual, cumulative')
    put(sm, rr + 2, 1, 'Subscription, cumulative')
    put(sm, rr + 3, 1, 'Helper: year subscription reaches perpetual')
    for y, col in enumerate(YC, 1):
        put(sm, rr + 1, 2 + y, f"=SUM('Perpetual Model'!$C${PM_TOT[i]}:{col}{PM_TOT[i]})", PHP)
        put(sm, rr + 2, 2 + y, f'=SUM($C${SM_TOT[i]}:{col}{SM_TOT[i]})', PHP)
        put(sm, rr + 3, 2 + y, f'=IF({col}{rr+2}>={col}{rr+1},{y},99)', '0')
        sm.cell(rr + 3, 2 + y).font = ITAL
    put(sm, rr + 1, 8, f'=IF(MIN(C{rr+3}:G{rr+3})=99,"Beyond Year 5","Year "&MIN(C{rr+3}:G{rr+3}))', bold=True, fill=F_TOT)
    put(sm, rr + 1, 9, f'=G{rr+2}-G{rr+1}', PHP, fill=F_TOT)
# calculator
cc = r + 1 + 4 * 5 + 1
sec(sm, cc, 'Calculator for a chosen user count (Inputs, Calculator section)', 9)
r = cc + 1
put(sm, r, 1, 'Named users')
put(sm, r, 3, '=Calc_Users', '0', fill=F_IN)
put(sm, r + 1, 1, 'Tier')
put(sm, r + 1, 3, f'=INDEX({TR("A")},{tidx(f"C{r}")})', bold=True)
put(sm, r + 2, 1, 'Billable users (minimum for tier applies)')
put(sm, r + 2, 3, f'={billable(f"C{r}")}', '0')
BN = f'$C${r+2}'
r += 4
head(sm, r, ['Slab', 'Users from - to', 'Users in slab', 'Price per user per month (PHP)', 'Monthly amount (PHP)'])
s0 = r + 1
for i in range(4):
    t = T0 + i
    rr = r + 1 + i
    put(sm, rr, 1, f'=Inputs!A{t}')
    put(sm, rr, 2, f'=Inputs!B{t}&" - "&IF(Inputs!C{t}>=100000,"above",Inputs!C{t})')
    put(sm, rr, 3, f'=MAX(0,MIN({BN},Inputs!C{t})-Inputs!B{t}+1)', '0')
    put(sm, rr, 4, f'=Inputs!E{t}', PHP)
    put(sm, rr, 5, f'=C{rr}*D{rr}', PHP)
r += 5
put(sm, r, 1, 'Monthly subscription, Year 1', bold=True)
put(sm, r, 5, f'=SUM(E{s0}:E{s0+3})', PHP, bold=True, fill=F_TOT)
MON = f'$E${r}'
put(sm, r + 1, 1, 'Effective price per user per month')
put(sm, r + 1, 5, f'=E{r}/{BN}', PHP)
put(sm, r + 2, 1, 'Onboarding fee')
put(sm, r + 2, 5, f"='Perpetual Model'!{IMPC}*Onboard_Pct", PHP)
ONB = f'$E${r+2}'
r += 4
head(sm, r, ['Item', 'Basis', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total (PHP)', '5-year total (USD)'])
k0 = r + 1
put(sm, k0, 1, 'Onboarding fee')
put(sm, k0 + 1, 1, 'Subscription fees (12 months)')
for y, col in enumerate(YC, 1):
    put(sm, k0, 2 + y, f'={ONB}' if y == 1 else 0, PHP)
    put(sm, k0 + 1, 2 + y, f'={MON}*12*(1+Esc)^({y}-1)', PHP)
for rr in (k0, k0 + 1):
    put(sm, rr, 8, f'=SUM(C{rr}:G{rr})', PHP)
    put(sm, rr, 9, f'=H{rr}/FX', USD)
rr = k0 + 2
put(sm, rr, 1, 'Total excl. VAT', bold=True, fill=F_TOT)
for c in range(3, 9):
    put(sm, rr, c, f'={L(c)}{k0}+{L(c)}{k0+1}', PHP, bold=True, fill=F_TOT)
put(sm, rr, 9, f'=H{rr}/FX', USD, fill=F_TOT)
put(sm, rr + 1, 1, 'VAT')
put(sm, rr + 2, 1, 'Total incl. VAT', bold=True)
for c in range(3, 9):
    put(sm, rr + 1, c, f'={L(c)}{rr}*VAT', PHP)
    put(sm, rr + 2, c, f'={L(c)}{rr}+{L(c)}{rr+1}', PHP, bold=True)
put(sm, rr + 1, 9, f'=H{rr+1}/FX', USD)
put(sm, rr + 2, 9, f'=H{rr+2}/FX', USD)
r = rr + 4
put(sm, r, 1, 'Perpetual, cumulative (calculator)')
put(sm, r + 1, 1, 'Subscription, cumulative (calculator)')
put(sm, r + 2, 1, 'Helper')
for y, col in enumerate(YC, 1):
    put(sm, r, 2 + y, f"=SUM('Perpetual Model'!$C${PM_CALC_TOT}:{col}{PM_CALC_TOT})", PHP)
    put(sm, r + 1, 2 + y, f'=SUM($C${k0+2}:{col}{k0+2})', PHP)
    put(sm, r + 2, 2 + y, f'=IF({col}{r+1}>={col}{r},{y},99)', '0')
put(sm, r, 8, f'=IF(MIN(C{r+2}:G{r+2})=99,"Beyond Year 5","Break-even Year "&MIN(C{r+2}:G{r+2}))', bold=True, fill=F_TOT)
sm.freeze_panes = 'C5'

# ------------------------------------------------------------------ Infrastructure
inf = sheet('Infrastructure', [20, 14, 18, 18, 18, 20, 20, 16, 50], 'Infrastructure (optional and separate)',
            'Indicative monthly cost by provider and tier, with margin. Production plus one UAT environment. The customer may host BrokerVerse itself at no fee.')
head(inf, 4, ['Provider', 'Tier', 'Cloud cost per month (USD)', 'Cloud cost per month (PHP)', 'Margin (PHP)', 'Price per month (PHP)',
              'Price per year, Year 1 (PHP)', 'Price per month (USD)', 'Notes'])
inf.row_dimensions[4].height = 45
provs = [('AWS', '1', 'Singapore region (ap-southeast-1); a Manila Local Zone exists for latency-sensitive parts (S22).'),
         ('Azure', 'Azure_Factor', 'Southeast Asia region (Singapore).'),
         ('Local partner', 'Local_Factor', 'Philippine data centre or partner cloud; data stays in the Philippines.')]
r = 4
for p, fac, note in provs:
    for j in range(4):
        r += 1
        put(inf, r, 1, p, bold=(j == 0))
        put(inf, r, 2, f'=Inputs!A{T0+j}')
        put(inf, r, 3, f'=Inputs!{L(2+j)}{INFRA_TOT}*{fac}', USD, fill=F_EST)
        put(inf, r, 4, f'=C{r}*FX', PHP)
        put(inf, r, 5, f'=D{r}*Infra_Margin', PHP)
        put(inf, r, 6, f'=ROUND(D{r}+E{r},-3)', PHP, bold=True, fill=F_TOT)
        put(inf, r, 7, f'=F{r}*12', PHP)
        put(inf, r, 8, f'=F{r}/FX', USD)
        put(inf, r, 9, note if j == 0 else '', wrap=True)
for j in range(4):
    r += 1
    put(inf, r, 1, 'Customer', bold=(j == 0))
    put(inf, r, 2, f'=Inputs!A{T0+j}')
    for c in range(3, 9):
        put(inf, r, c, 0, PHP)
    put(inf, r, 9, 'Broker hosts itself (cloud account or own data centre). Deployment guide supplied; set-up support at day rates.' if j == 0 else '', wrap=True)
INF0, INF1 = 5, r
r += 2
sec(inf, r, 'What the infrastructure price covers', 9)
for line in ['Covers: cloud resources for production and one UAT environment, daily backups, monitoring and alarms, operating system and database patching, TLS certificates.',
             'Reference sizing: the BrokerVerse capacity document sizes 200 named users (80 concurrent) at 2 API containers of 1 vCPU, 2 GB and Amazon RDS for PostgreSQL 16 db.m7g.large Multi-AZ, 100 GB gp3 storage.',
             'Excludes: e-mail sending service charges above normal volumes, SMS, payment gateway fees, domain names, extra environments (see Optional Services) and disaster-recovery copies for tiers other than Enterprise.',
             'All cloud figures are estimates from public on-demand list prices, rounded up. Confirm with the provider pricing calculator and a partner quotation before a binding offer. Reserved instances or savings plans can lower cost 20% to 40% for a 1 to 3 year commitment.',
             'Billing: monthly in advance in PHP. The margin line covers FX movement between USD cloud billing and PHP invoicing, and the operations effort.']:
    r += 1
    inf.cell(r, 1, line).alignment = WRAP
    inf.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    inf.row_dimensions[r].height = 30
inf.freeze_panes = 'C5'
INF_REF = (f"Infrastructure!$F${INF0}:$F${INF1}", f"Infrastructure!$A${INF0}:$A${INF1}", f"Infrastructure!$B${INF0}:$B${INF1}")

# ------------------------------------------------------------------ Change Requests
cr = sheet('Change Requests', [34, 16, 16, 16, 16, 16, 16, 50], 'Change requests (CR)',
           'Man-day rates by role, blended rate, CR bands and the estimation method. PHP, exclusive of VAT.')
head(cr, 4, ['Role', 'Day rate (PHP)', 'Day rate (USD)', 'Hourly equivalent (PHP)', 'Mix in blended rate'])
for k, rr in enumerate(ROLE_ROWS):
    put(cr, 5 + k, 1, f'=Inputs!A{rr}')
    put(cr, 5 + k, 2, f'=Inputs!B{rr}', PHP)
    put(cr, 5 + k, 3, f'=B{5+k}/FX', USD)
    put(cr, 5 + k, 4, f'=B{5+k}/8', PHP)
    put(cr, 5 + k, 5, f'=Inputs!C{rr}', PCT)
put(cr, 9, 1, 'Blended rate', bold=True, fill=F_TOT)
put(cr, 9, 2, '=Blended', PHP, bold=True, fill=F_TOT)
put(cr, 9, 3, '=B9/FX', USD)
put(cr, 9, 4, '=B9/8', PHP)
r = 11
sec(cr, r, 'Estimation method', 8)
steps = ['1. The broker raises the CR with the business need and examples. iorta TechNXT replies with an impact note within 5 business days (Small) or 10 (Medium, Large).',
         '2. Developer man-days are estimated from the screens, APIs, reports, database changes and jobs affected.',
         '3. Business analysis = development x BA %; testing = development x QA %; project management = (BA + development + testing) x PM %; contingency on the total. Days are rounded up to the half day.',
         '4. Cost = days by role x role day rate. The band decides the approval route and the payment terms: Small, billed on delivery; Medium, 50% on approval and 50% on UAT sign-off; Large, 30% / 50% / 20% like a project.',
         '5. CRs change the OOTB version for that customer only. The changed code is covered by AMC or subscription support only when merged into the product line; otherwise support of the change is billed at day rates.']
for s in steps:
    r += 1
    cr.cell(r, 1, s).alignment = WRAP
    cr.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    cr.row_dimensions[r].height = 30
r += 2
sec(cr, r, 'CR bands', 8)
r += 1
head(cr, r, ['Band', 'Total man-days from', 'Total man-days to', 'Indicative cost from (PHP, blended)', 'Indicative cost to (PHP, blended)', 'Approval', '', 'Typical examples'])
bands = [('Small', '0.5', 'CR_Small_Max', 'Broker IT or operations head', 'New report column or filter; document template wording; new validation; new field on a screen.'),
         ('Medium', 'CR_Small_Max+0.5', 'CR_Med_Max', 'Broker sponsor', 'New report; new letter or statement; change to a commission or tax computation; new upload template.'),
         ('Large', 'CR_Med_Max+0.5', 'CR_Large_Max', 'Steering committee', 'New module screen set; new insurer or bank integration; new workflow with approvals.')]
for b in bands:
    r += 1
    put(cr, r, 1, b[0], bold=True)
    put(cr, r, 2, f'={b[1]}', DEC)
    put(cr, r, 3, f'={b[2]}', DEC)
    put(cr, r, 4, f'=B{r}*Blended', PHP)
    put(cr, r, 5, f'=C{r}*Blended', PHP)
    put(cr, r, 6, b[3], wrap=True)
    put(cr, r, 8, b[4], wrap=True)
    cr.row_dimensions[r].height = 30
r += 1
cr.cell(r, 1, 'Above the Large band: treated as a separate project with its own statement of work.').font = ITAL
r += 2
sec(cr, r, 'Worked examples and calculator', 8)
r += 1
head(cr, r, ['Line', 'Example Small', 'Example Medium', 'Example Large', 'Calculator (Inputs)', '', '', 'Notes'])
e0 = r + 1
lines = ['Development man-days', 'Business analysis man-days', 'Testing man-days', 'Project management man-days', 'Contingency man-days',
         'Total man-days', 'Band', 'Cost: business analyst', 'Cost: developer', 'Cost: QA engineer', 'Cost: project manager', 'Contingency cost (blended)',
         'CR price excl. VAT (PHP)', 'CR price (USD)']
for k, lab in enumerate(lines):
    put(cr, e0 + k, 1, lab, bold=(k in (5, 12)))
devs = [2, 7, 20, 'CR_Dev_Days']
for j, d in enumerate(devs):
    c = L(2 + j)
    R = lambda k: f'{c}{e0+k}'  # noqa: E731
    put(cr, e0, 2 + j, f'={d}' if isinstance(d, str) else d, DEC, inp=not isinstance(d, str))
    put(cr, e0 + 1, 2 + j, f'=ROUNDUP({R(0)}*CR_BA*2,0)/2', DEC)
    put(cr, e0 + 2, 2 + j, f'=ROUNDUP({R(0)}*CR_QA*2,0)/2', DEC)
    put(cr, e0 + 3, 2 + j, f'=ROUNDUP(({R(0)}+{R(1)}+{R(2)})*CR_PM*2,0)/2', DEC)
    put(cr, e0 + 4, 2 + j, f'=ROUNDUP(({R(0)}+{R(1)}+{R(2)}+{R(3)})*CR_Cont*2,0)/2', DEC)
    put(cr, e0 + 5, 2 + j, f'=SUM({c}{e0}:{c}{e0+4})', DEC, bold=True)
    put(cr, e0 + 6, 2 + j, f'=IF({R(5)}<=CR_Small_Max,"Small",IF({R(5)}<=CR_Med_Max,"Medium",IF({R(5)}<=CR_Large_Max,"Large","Project")))', bold=True)
    put(cr, e0 + 7, 2 + j, f'={R(1)}*Rate_BA', PHP)
    put(cr, e0 + 8, 2 + j, f'={R(0)}*Rate_Dev', PHP)
    put(cr, e0 + 9, 2 + j, f'={R(2)}*Rate_QA', PHP)
    put(cr, e0 + 10, 2 + j, f'={R(3)}*Rate_PM', PHP)
    put(cr, e0 + 11, 2 + j, f'={R(4)}*Blended', PHP)
    put(cr, e0 + 12, 2 + j, f'=SUM({c}{e0+7}:{c}{e0+11})', PHP, bold=True, fill=F_TOT)
    put(cr, e0 + 13, 2 + j, f'={R(12)}/FX', USD)
put(cr, e0, 8, 'Examples are illustrative; the calculator column reads CR_Dev_Days in Inputs.')
cr.freeze_panes = 'B5'

# ------------------------------------------------------------------ Optional Services
op = sheet('Optional Services', [46, 22, 18, 16, 13, 62], 'Optional services',
           'Priced on request in addition to either model. PHP, exclusive of VAT. Recurring items escalate 5% a year.')
head(op, 4, ['Item', 'Unit', 'Price (PHP)', 'Price (USD)', 'Man-days', 'Basis and notes'])
items = [
    ('Additional environment (training, second UAT), set-up', 'one-time', '=ROUND(Env_Setup_Days*Blended/Round_To,0)*Round_To', 'Env_Setup_Days', 'Set-up man-days x blended rate.'),
    ('Additional environment hosted by iorta TechNXT', 'per month', '=ROUND(Env_Monthly_USD*FX*(1+Infra_Margin),-3)', None, 'Cloud cost plus infrastructure margin. Customer-hosted: set-up fee only.'),
    ('Training beyond the included days', 'per trainer day', '=Rate_Trainer', None, 'Up to 15 participants per session; customer provides the room and PCs.'),
    ('Data migration beyond the standard templates', 'per legacy source', '=ROUND(Mig_Days_Source*Blended/Round_To,0)*Round_To', 'Mig_Days_Source', 'Standard: data loaded by the broker with the BrokerVerse upload templates (docs/templates), with guidance. This item covers extraction and transformation from a legacy system.'),
    ('Data migration, additional effort', 'per man-day', '=Blended', None, 'For volume or data-quality work beyond the estimate.'),
    ('Additional integration, standard', 'per integration', '=ROUND(Int_Std_Days*Blended/Round_To,0)*Round_To', 'Int_Std_Days', 'One documented API or file exchange, one direction.'),
    ('Additional integration, complex', 'per integration', '=ROUND(Int_Cx_Days*Blended/Round_To,0)*Round_To', 'Int_Cx_Days', 'Two-way, no API, or batch with reconciliation.'),
    ('Additional line of business after go-live', 'per line', '=ROUND(Days_per_LOB*Blended/Round_To,0)*Round_To', 'Days_per_LOB', 'Product set-up, rating and tax lines, templates, testing.'),
    ('New report or document template', 'per item (typical)', '=ROUND(Report_Days*Blended/Round_To,0)*Round_To', 'Report_Days', 'Small CR band; firm price after the impact note.'),
    ('On-site day (Metro Manila)', 'per consultant day', '=Rate_Onsite', None, 'Outside Metro Manila: travel, lodging and meals at cost.'),
]
r = 4
for lab, unit, f, days, note in items:
    r += 1
    put(op, r, 1, lab)
    put(op, r, 2, unit)
    put(op, r, 3, f, PHP, bold=True)
    put(op, r, 4, f'=C{r}/FX', USD)
    put(op, r, 5, f'={days}' if days else '', '0')
    put(op, r, 6, note, wrap=True)
    op.row_dimensions[r].height = 30
for j in range(4):
    r += 1
    put(op, r, 1, f'="24x7 Severity 1 support, "&Inputs!A{T0+j}&" tier"')
    put(op, r, 2, 'per year, Year 1')
    put(op, r, 3, f'=Inputs!L{T0+j}', PHP, bold=True)
    put(op, r, 4, f'=C{r}/FX', USD)
    put(op, r, 6, 'Standard support is business hours (Monday to Friday, 08:00 to 18:00 Philippine time, excluding holidays). This adds round-the-clock response to Severity 1 incidents (production down or a core process stopped for all users).' if j == 0 else '', wrap=True)
    op.row_dimensions[r].height = 30 if j == 0 else 15
SUP0 = r - 3
r += 2
sec(op, r, 'Included in the standard implementation (no extra charge)', 6)
for line in ['Configuration of the company, branches, users and roles, insurers, commission rates, chart of accounts, banks and the included lines of business.',
             'Data load using the standard upload templates (about 40 templates in docs/templates), with the broker supplying clean data.',
             'Training days included for the tier (train-the-trainer), UAT support, go-live support and 4 weeks of hypercare.',
             'One production and one UAT environment when iorta TechNXT hosts; deployment guidance when the broker hosts.']:
    r += 1
    op.cell(r, 1, line).alignment = WRAP
    op.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    op.row_dimensions[r].height = 18
op.freeze_panes = 'B5'

# ------------------------------------------------------------------ Sample Quotes
sq = sheet('Sample Quotes', [40, 13, 13, 13, 13, 13, 16, 14], 'Sample quotes: three brokers, 5-year total cost of ownership',
           'Scenario drivers are in Inputs (Sample quote scenarios). PHP, exclusive of VAT unless stated. Optional items shown separately and added to both models.')
r = 3
for j in range(3):
    col = L(2 + j)
    g = lambda k: f'Inputs!${col}${SC[k]}'  # noqa: E731
    r += 1
    sec(sq, r, '', 8)
    sq.cell(r, 1).value = f'="Example {j+1}: "&{g("label")}&", "&{g("users")}&" users, "&{g("lobs")}&" lines of business, tier "&INDEX({TR("A")},{tidx(g("users"))})&", hosting: "&{g("prov")}'
    r += 1
    head(sq, r, ['Item', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '5-year total (PHP)', '5-year total (USD)'])
    N, LOB = g('users'), g('lobs')
    tier = f'INDEX({TR("A")},{tidx(N)})'
    lic = slab_sum(N, 'D')
    mon = slab_sum(billable(N), 'E')
    rows_def = []
    r += 1
    sq.cell(r, 1, 'Model A: perpetual').font = B
    r += 1
    a0 = r
    put(sq, r, 1, 'Implementation fee')
    put(sq, r + 1, 1, 'Perpetual licence fee')
    put(sq, r + 2, 1, 'AMC (22%, +5% a year)')
    put(sq, r + 3, 1, 'Software subtotal, perpetual', bold=True)
    for y, yc in enumerate(['B', 'C', 'D', 'E', 'F'], 1):
        put(sq, r, 1 + y, f'={impl(N, LOB)}' if y == 1 else 0, PHP)
        put(sq, r + 1, 1 + y, f'={lic}' if y == 1 else 0, PHP)
        put(sq, r + 2, 1 + y, amc(f'$B${r+1}', y), PHP)
        put(sq, r + 3, 1 + y, f'=SUM({yc}{r}:{yc}{r+2})', PHP, bold=True)
    r += 4
    sq.cell(r, 1, 'Model B: subscription').font = B
    r += 1
    b0r = r
    put(sq, r, 1, 'Onboarding fee')
    put(sq, r + 1, 1, f'="Subscription ("&{billable(N)}&" billable users)"')
    put(sq, r + 2, 1, 'Software subtotal, subscription', bold=True)
    for y, yc in enumerate(['B', 'C', 'D', 'E', 'F'], 1):
        put(sq, r, 1 + y, f'=$B${a0}*Onboard_Pct' if y == 1 else 0, PHP)
        put(sq, r + 1, 1 + y, f'=({mon})*12*(1+Esc)^({y}-1)', PHP)
        put(sq, r + 2, 1 + y, f'={yc}{r}+{yc}{r+1}', PHP, bold=True)
    r += 3
    sq.cell(r, 1, 'Optional items (same in both models)').font = B
    r += 1
    o0 = r
    put(sq, r, 1, f'="Infrastructure: "&{g("prov")}')
    put(sq, r + 1, 1, '24x7 Severity 1 support')
    put(sq, r + 2, 1, 'Extra training, migration and integrations (one-time)')
    put(sq, r + 3, 1, 'Optional subtotal', bold=True)
    infra_m = f'SUMIFS({INF_REF[0]},{INF_REF[1]},{g("prov")},{INF_REF[2]},{tier})'
    one_time = (f'{g("train")}*Rate_Trainer+ROUND(Mig_Days_Source*Blended/Round_To,0)*Round_To*{g("mig")}'
                f'+ROUND(Int_Std_Days*Blended/Round_To,0)*Round_To*{g("int")}')
    for y, yc in enumerate(['B', 'C', 'D', 'E', 'F'], 1):
        put(sq, r, 1 + y, f'={infra_m}*12*(1+Infra_Esc)^({y}-1)', PHP)
        put(sq, r + 1, 1 + y, f'=IF({g("sup")}="Yes",{tget("L", N)}*(1+Esc)^({y}-1),0)', PHP)
        put(sq, r + 2, 1 + y, f'={one_time}' if y == 1 else 0, PHP)
        put(sq, r + 3, 1 + y, f'=SUM({yc}{r}:{yc}{r+2})', PHP, bold=True)
    r += 4
    sq.cell(r, 1, 'Total cost of ownership').font = B
    r += 1
    t0 = r
    labels = ['Perpetual with options, excl. VAT', 'Subscription with options, excl. VAT', 'Perpetual with options, incl. VAT', 'Subscription with options, incl. VAT']
    for k, lab in enumerate(labels):
        put(sq, r + k, 1, lab, bold=True, fill=F_TOT if k < 2 else None)
    for y, yc in enumerate(['B', 'C', 'D', 'E', 'F'], 1):
        put(sq, r, 1 + y, f'={yc}{a0+3}+{yc}{o0+3}', PHP, bold=True, fill=F_TOT)
        put(sq, r + 1, 1 + y, f'={yc}{b0r+2}+{yc}{o0+3}', PHP, bold=True, fill=F_TOT)
        put(sq, r + 2, 1 + y, f'={yc}{r}*(1+VAT)', PHP)
        put(sq, r + 3, 1 + y, f'={yc}{r+1}*(1+VAT)', PHP)
    for rr in list(range(a0, a0 + 4)) + list(range(b0r, b0r + 3)) + list(range(o0, o0 + 4)) + list(range(t0, t0 + 4)):
        put(sq, rr, 7, f'=SUM(B{rr}:F{rr})', PHP, bold=True)
        put(sq, rr, 8, f'=G{rr}/FX', USD)
    r += 4
    put(sq, r, 1, 'Lower 5-year cost (excl. VAT)', bold=True)
    put(sq, r, 2, f'=IF(G{t0}<G{t0+1},"Perpetual","Subscription")', bold=True)
    put(sq, r, 3, 'by (PHP)')
    put(sq, r, 4, f'=ABS(G{t0}-G{t0+1})', PHP)
    put(sq, r, 5, 'Year 1 cash, perpetual vs subscription')
    sq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    put(sq, r, 7, f'=B{t0}', PHP)
    put(sq, r, 8, f'=B{t0+1}', PHP)
    r += 1
put(sq, r + 1, 1, 'Payment terms assumed: implementation 40% on signing, 40% on UAT sign-off, 20% on go-live; perpetual licence on go-live; AMC yearly in advance; subscription and hosting monthly in advance.').font = ITAL
sq.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=8)
sq.freeze_panes = 'B4'

# ------------------------------------------------------------------ Market Benchmark
mb = sheet('Market Benchmark', [36, 18, 22, 30, 13, 13, 17, 17, 12, 54], 'Market benchmark: comparable systems and prices',
           'Published prices are rare in this market. "Third-party estimate" means a review or comparison site, not the vendor. Converted at the Inputs exchange rate.')
head(mb, 4, ['Vendor / product', 'Origin / market', 'Pricing model', 'Price as found', 'Low (USD per user per month)', 'High (USD per user per month)',
             'Low (PHP per user per month)', 'High (PHP per user per month)', 'Source ID', 'Type of evidence and notes'])
mb.row_dimensions[4].height = 45
bm = [
    ('Applied Epic', 'USA, global', 'Subscription per user, quote only', 'USD 125 to 250 per user per month; set-up USD 10,000 to 25,000', 125, 250, 'S5', 'Third-party estimate. Large-agency system; not localised for Philippine taxes or IC reports.'),
    ('Applied Sagitta', 'USA', 'Enterprise, quote only', 'USD 200 to 500 per user per month plus set-up', 200, 500, 'S9', 'Third-party estimate.'),
    ('Vertafore AMS360', 'USA', 'Subscription per user, quote only', 'Listings from USD 99; agencies report USD 150 to 300 per user per month', 150, 300, 'S6', 'Third-party estimate.'),
    ('Insly', 'Estonia / EU, global', 'Subscription per user, tiers', 'USD 89 (Starter) to 199 (Professional) per user per month; EUR 149 basic', 89, 199, 'S7', 'Third-party listings; Enterprise tier higher.'),
    ('Open GI', 'UK', 'Subscription per user', 'From USD 70 per user per month', 70, 70, 'S8', 'Third-party listing (starting price).'),
    ('Acturis', 'UK, Europe', 'Service fee by users and modules', 'USD 2,000 to 25,000+ per month per firm; implementation USD 15,000 to 80,000', None, None, 'S8', 'Third-party estimate; per firm, not per user.'),
    ('Salesforce Financial Services Cloud', 'USA, global', 'Published per user, billed annually', 'USD 325 to 350 per user per month', 325, 350, 'S12', 'Published list price. CRM platform only: broking accounting, remittance and BIR reports must be built.'),
    ('Zoho CRM', 'India, global', 'Published per user', 'USD 23 to 50 per user per month (Professional to Enterprise)', 23, 50, 'S13', 'Published list price. Generic CRM floor price; no broking functions.'),
    ('Agiliux', 'Singapore / Malaysia; live clients in the Philippines', 'Modular; pay-as-you-sell for digital', 'No public price', None, None, 'S23', 'Main regional competitor. Price on request.'),
    ('Mzapp (Mindzen)', 'India', 'Quote only', 'No public per-user price', None, None, 'S24', 'Indian broker system; one listing shows a USD 200,000 starting price (basis unclear).'),
    ('Sibro', 'India; markets to the Philippines', 'Quote only', 'No public price', None, None, 'S25', 'Regional broker system.'),
    ('BrokerEdge (Damco)', 'India / USA', 'Quote only', 'No public price', None, None, 'S26', 'Feature-based pricing.'),
    ('Xceedance', 'USA / India', 'Services and platforms, quote only', 'No public price', None, None, 'S27', 'Builds agency and policy systems and offers operations services.'),
]
r = 4
for v in bm:
    r += 1
    put(mb, r, 1, v[0], bold=True)
    put(mb, r, 2, v[1], wrap=True)
    put(mb, r, 3, v[2], wrap=True)
    put(mb, r, 4, v[3], wrap=True)
    put(mb, r, 5, v[4], USD)
    put(mb, r, 6, v[5], USD)
    put(mb, r, 7, f'=IF(E{r}="","",E{r}*FX)', PHP)
    put(mb, r, 8, f'=IF(F{r}="","",F{r}*FX)', PHP)
    put(mb, r, 9, v[6])
    put(mb, r, 10, v[7], wrap=True, fill=F_EST if 'estimate' in v[7].lower() else None)
    mb.row_dimensions[r].height = 60
r += 1
for j in range(4):
    put(mb, r + j, 1, f'="BrokerVerse OOTB, "&Inputs!A{T0+j}&" slab"', bold=True, fill=F_TOT)
    put(mb, r + j, 2, 'Philippines')
    put(mb, r + j, 3, 'Subscription per user, graduated slabs')
    put(mb, r + j, 4, f'="PHP "&TEXT(Inputs!E{T0+j},"#,##0")&" per user per month"')
    put(mb, r + j, 5, f'=Inputs!E{T0+j}/FX', USD)
    put(mb, r + j, 6, f'=Inputs!E{T0+j}/FX', USD)
    put(mb, r + j, 7, f'=Inputs!E{T0+j}', PHP)
    put(mb, r + j, 8, f'=Inputs!E{T0+j}', PHP)
    put(mb, r + j, 10, 'Recommended price (this workbook).')
r += 5
sec(mb, r, 'Positioning summary', 10)
r += 1
put(mb, r, 1, 'Global broker systems (median of low/high, USD per user per month)')
put(mb, r, 7, '=MEDIAN(E5:F7,E9:F9)', USD)
put(mb, r, 10, 'Applied Epic, Sagitta, AMS360, Open GI (third-party estimates).')
r += 1
put(mb, r, 1, 'BrokerVerse Small slab as % of that median')
put(mb, r, 7, f'=(Inputs!E{T0}/FX)/G{r-1}', PCT)
r += 1
put(mb, r, 1, 'BrokerVerse Small slab vs Zoho CRM Enterprise monthly (USD 50)')
put(mb, r, 7, f'=(Inputs!E{T0}/FX)/50', '0.0"x"')
put(mb, r, 10, 'A broking system with accounting and BIR reports sits above generic CRM and below global broker systems.')
r += 2
sec(mb, r, 'Philippine IT services rates (basis for day rates)', 10)
r += 1
head(mb, r, ['Item', 'Market', 'Unit', 'Range as found', 'Low (USD)', 'High (USD)', 'Low (PHP)', 'High (PHP)', 'Source ID', 'Notes'])
it = [('Mid-level developer', 'Philippines', 'per hour', 'USD 30 to 50', 30, 50, 'S15', 'Outsourcing market rates; day = 8 hours.'),
      ('Senior developer', 'Philippines', 'per hour', 'USD 50 to 75', 50, 75, 'S15', ''),
      ('IT consulting', 'Philippines', 'per hour', 'USD 30 to 70; premium firms 70 to 150', 30, 70, 'S16', ''),
      ('Senior business analyst salary', 'Philippines', 'per month', 'PHP 47,054 to 105,688 (average 70,519)', None, None, 'S17', 'Employee cost, not vendor rate.'),
      ('Senior QA engineer salary', 'Philippines', 'per month', 'PHP 54,000 to 100,000 (average 65,053)', None, None, 'S18', 'Employee cost.'),
      ('IT project manager salary', 'Philippines', 'per month', 'PHP 60,000 to 123,000 (average 64,826)', None, None, 'S19', 'Employee cost.'),
      ('BrokerVerse blended day rate', 'Philippines', 'per day', '=TEXT(Blended,"#,##0")&" PHP"', None, None, '', 'Recommended (Inputs).')]
for v in it:
    r += 1
    put(mb, r, 1, v[0], bold=True)
    put(mb, r, 2, v[1])
    put(mb, r, 3, v[2])
    put(mb, r, 4, v[3])
    put(mb, r, 5, v[4], USD)
    put(mb, r, 6, v[5], USD)
    put(mb, r, 7, f'=IF(E{r}="","",E{r}*FX)', PHP)
    put(mb, r, 8, f'=IF(F{r}="","",F{r}*FX)', PHP)
    put(mb, r, 9, v[6])
    put(mb, r, 10, v[7])
r += 2
sec(mb, r, 'Philippine market facts used for tiers and affordability', 10)
facts = [('Licensed insurance brokers', 'About 67 insurance brokers (IC key data, 2024 preliminary). A small, concentrated market: price for a few dozen realistic prospects.', 'S1'),
         ('Largest brokers by commission earned, 2024', 'BDO Insurance and Reinsurance Brokers PHP 2,434 million; Marsh PHP 2,075 million; Aon PHP 2,047 million; Lockton PHP 712 million; PhilPacific PHP 598 million.', 'S2'),
         ('Mid-market brokers, 2024', 'Rank 20 East West Insurance Brokerage PHP 98.0 million; rank 30 Moneyhero Insurance Brokerage PHP 55.2 million.', 'S2'),
         ('Premiums produced by brokers, 2024', 'Aon was first at PHP 23.4 billion of premiums produced.', 'S3'),
         ('Non-life insurers, 2025', 'Gross premiums written PHP 149.11 billion (excluding reinsurers); Pioneer PHP 20.70 billion (14%), Prudential Guarantee PHP 13.86 billion (9%), Malayan PHP 13.35 billion (9%). Pioneer first in net premiums written at PHP 6.90 billion.', 'S4'),
         ('What the insurer ranking suggests', 'The top 3 insurers hold about a third of gross premiums and the rest is spread over many mid-size insurers, so a broker places with many insurers: insurer set-up, remittance and reconciliation by insurer are core needs in every tier.', 'S4'),
         ('IT budget norm', 'Insurance IT spend 3.1% of revenue (Gartner, 2010); later benchmarks 4% to 7% for insurance. Used as the affordability band.', 'S10, S11'),
         ('Broker licensing', 'Minimum paid-up capital PHP 20 million for an insurance broker, PHP 50 million for insurance and reinsurance broker.', 'S28')]
for k, v, s in facts:
    r += 1
    put(mb, r, 1, k, bold=True, wrap=True)
    put(mb, r, 2, v, wrap=True)
    mb.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    put(mb, r, 9, s)
    mb.row_dimensions[r].height = 45
mb.freeze_panes = 'B5'

# ------------------------------------------------------------------ Sources
so = sheet('Sources', [7, 52, 70, 13, 52], 'Sources', 'Accessed 03 October 2026. Some Insurance Commission and BusinessWorld pages could not be opened from the research network; their figures were read from search-engine extracts of the same documents and should be checked against the originals before external use.')
head(so, 4, ['ID', 'Title / publisher', 'URL', 'Accessed', 'Used for'])
src = [
    ('S1', 'Insurance Commission: Key Data 2020-2024 (preliminary), Annex A', 'https://www.insurance.gov.ph/wp-content/uploads/2026/02/Annex-A-Key-Data-2020-2024-Preliminary-Figures.pdf', 'Number of licensed insurance brokers (about 67, 2024). Read from search extract.'),
    ('S2', 'Insurance Commission: Performance of Insurance Brokers, Commission Earned as of 31 December 2024', 'https://www.insurance.gov.ph/wp-content/uploads/2026/01/Performance-of-Insurance-Brokers-Commission-Earned-as-of-31-December-2024.pdf', 'Broker ranking by commission earned. Read from search extract.'),
    ('S3', 'Insurance Commission: Performance of Insurance Brokers, Premiums Produced as of 31 December 2024', 'https://www.insurance.gov.ph/wp-content/uploads/2026/01/Performance-of-Insurance-Brokers-Premiums-Produced-as-of-31-December-2024.pdf', 'Aon premiums produced. Read from search extract.'),
    ('S4', 'The Manila Times, "Growth on a sure footing", 31 August 2026; Insurance Commission non-life NPW and GPW rankings 2025', 'https://www.manilatimes.net/2026/08/31/supplements/growth-on-a-sure-footing/2414786 ; https://www.insurance.gov.ph/wp-content/uploads/2026/02/Performance-of-Non-Life-Insurance-Companies-NPW-as-of-31-December-2025-Based-on-Submitted-Unaudited-EQRSFS.pdf', 'Non-life GPW 2025 and top insurers. Read from search extract.'),
    ('S5', 'Software Finder / All Insurance CRM: Applied Epic pricing', 'https://softwarefinder.com/insurance-software/applied-epic ; https://allinsurancecrm.com/compare/applied-epic/', 'Applied Epic per-user estimate and set-up cost.'),
    ('S6', 'Unlocked CRM: How much does Vertafore AMS360 cost (2026)', 'https://unlockedcrm.ai/blog/how-much-does-vertafore-ams360-really-cost', 'AMS360 per-user estimate.'),
    ('S7', 'ITQlick / GetApp: Insly pricing', 'https://www.itqlick.com/insly/pricing ; https://www.getapp.com/industries-software/a/insly-com/pricing/', 'Insly tiers.'),
    ('S8', 'Capterra: Open GI and Acturis pricing; Acturis FAQ', 'https://www.capterra.com/p/79710/Transactor/ ; https://www.capterra.com/p/238790/Acturis/ ; https://www.acturis.com/faqs/how-much-does-it-cost/', 'Open GI starting price; Acturis monthly range and model.'),
    ('S9', 'WifiTalents / ZipDo: best insurance broking management software 2026', 'https://wifitalents.com/best/insurance-broking-management-software/', 'Sagitta and general price bands.'),
    ('S10', 'Gartner IT Key Metrics (2010) via Scribd', 'https://www.scribd.com/document/670004267/Gartner-It-Spent', 'Insurance IT spend 3.1% of revenue.'),
    ('S11', 'IT Budget Calculator: IT spending benchmarks by industry 2026', 'https://itbudgetcalculator.com/by-industry', 'Insurance IT spend 4% to 7% (third-party benchmark).'),
    ('S12', 'Salesforce: Financial Services Cloud pricing', 'https://www.salesforce.com/financial-services/sales-pricing/', 'Published list price.'),
    ('S13', 'G2: Zoho CRM pricing', 'https://www.g2.com/products/zoho-crm/pricing', 'Published list price.'),
    ('S14', 'Bangko Sentral ng Pilipinas reference exchange rate (via search extract of BSP bulletin and rate sites)', 'https://www.bsp.gov.ph/Lists/RERB/Attachments/2311/10Jul2026.pdf ; https://pesohub.ph/rates/exchange-rates/usd-to-php-today/', 'PHP 62.75 per USD on 25 September 2026.'),
    ('S15', 'Hireplicity: Philippines software developer rates', 'https://www.hireplicity.com/blog/philippines-developer-rates-hourly-project-based', 'Developer hourly rates.'),
    ('S16', 'Second Talent / Full Scale: offshore development rates 2026', 'https://www.secondtalent.com/resources/software-development-outsourcing-philippines/ ; https://fullscale.io/blog/comparing-offshore-software-development-rates-by-country/', 'IT consulting rates.'),
    ('S17', 'Indeed Philippines: senior business analyst salary', 'https://ph.indeed.com/career/senior-business-analyst/salaries', 'BA salary.'),
    ('S18', 'Glassdoor: senior QA engineer salary, Philippines', 'https://www.glassdoor.com/Salaries/philippines-senior-qa-engineer-salary-SRCH_IL.0,11_IN204_KO12,30.htm', 'QA salary.'),
    ('S19', 'Glassdoor: IT project manager salary, Philippines', 'https://www.glassdoor.com/Salaries/philippines-it-project-manager-salary-SRCH_IL.0,11_IN204_KO12,30.htm', 'PM salary.'),
    ('S20', 'Vantage / Economize: Amazon RDS db.m6g.large pricing', 'https://instances.vantage.sh/aws/rds/db.m6g.large', 'Database instance price reference (about USD 0.16 to 0.18 an hour).'),
    ('S21', 'CloudPrice: Azure Standard_D4s_v5 pricing', 'https://cloudprice.net/vm/Standard_D4s_v5', 'Azure VM price reference (USD 0.216 an hour, Southeast Asia).'),
    ('S22', 'AWS: Local Zones in Manila generally available (June 2023)', 'https://aws.amazon.com/about-aws/whats-new/2023/06/general-availability-aws-local-zones-manila', 'Hosting options.'),
    ('S23', 'Agiliux: FAQs and broker solutions', 'https://www.agiliux.com/resource-centre/faqs/', 'Regional competitor, pricing model.'),
    ('S24', 'Mindzen: FAQ; SoftwareSuggest Mzapp listing', 'https://mindzen.com/faq-mindzen-insurance-products/ ; https://www.softwaresuggest.com/mzapp', 'Indian competitor.'),
    ('S25', 'Sibro: best insurance broker management software', 'https://sibro.xyz/best-insurance-broker-management-software/', 'Regional competitor.'),
    ('S26', 'Damco: BrokerEdge', 'https://www.damcogroup.com/insurance/brokeredge-broker-management-software', 'Competitor.'),
    ('S27', 'Xceedance: broking operations', 'https://www.xceedance.com/empowering-broking-operations/', 'Competitor.'),
    ('S28', 'Baker McKenzie: Guide for Insurance Sales, Advisory and Distribution, Philippines', 'https://resourcehub.bakermckenzie.com/en/resources/asia-pacific-insurance/asia-pacific/philippines/topics/guide-for-insurance-sales-advisory-and-distribution', 'Broker capital requirement.'),
    ('S29', 'Respicio & Co.: withholding tax on subscription and professional fees; HitPay withholding tax guide', 'https://www.respicio.ph/commentaries/philippine-withholding-tax-rates-on-subscription-and-professional-fees ; https://hitpayapp.com/blog/withholding-tax-philippines', 'Withholding tax note (to be confirmed by tax advisers).'),
    ('S30', 'BrokerVerse documentation: Capacity and Performance; Go-live data set-up; docs/templates', 'docs/architecture/06_BrokerVerse_Capacity_and_Performance.pdf ; docs/onboarding/GO_LIVE_DATA_SETUP.md', 'Reference sizing and the standard upload templates.'),
]
r = 4
for s in src:
    r += 1
    put(so, r, 1, s[0], bold=True)
    put(so, r, 2, s[1], wrap=True)
    put(so, r, 3, s[2], wrap=True)
    put(so, r, 4, ACC)
    put(so, r, 5, s[3], wrap=True)
    so.row_dimensions[r].height = 45
so.freeze_panes = 'A5'

for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
    w.print_title_rows = None
wb.save(OUT)
print('saved', OUT)
