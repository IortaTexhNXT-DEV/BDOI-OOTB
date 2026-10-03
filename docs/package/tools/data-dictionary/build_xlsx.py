import json, os, re, subprocess, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import gen
import tables_meta as tm
import pii

OUT = '/home/user/BDOI-OOTB/docs/package/out/BrokerVerse_Data_Dictionary.xlsx'
# Database read by dump_db.py / vals.py (DD_DB, default the loaded test database "golive"). DD_REFERENCE_ONLY=1 when it
# was built with the migrations and the reference seed data only (SEED_SAMPLE_DATA=false): row counts are then those
# of the reference data.
DD_DB = os.environ.get('DD_DB', 'golive')
REFERENCE_ONLY = os.environ.get('DD_REFERENCE_ONLY') == '1'
DB_WORDS = 'reference data only' if REFERENCE_ONLY else 'loaded test database'
D = gen.D
M = gen.M
_dyn = ['countries','states','cities','currencies','banks','insurance_companies','products','policy_types',
        'vehicle_brands','vehicle_models','vehicle_variants','coverages','signatories','branches']
for _t in _dyn:
    for _c in ('attrs', 'created_by', 'updated_by', 'created_at', 'updated_at'):
        M['colmig'].setdefault(f'{_t}.{_c}', '0040_config_masters.sql')
for _c in ('name', 'applied_at'):
    M['colmig'].setdefault(f'schema_migrations.{_c}', 'db/migrate.js')


def q(sql):
    out = subprocess.run(['su', 'postgres', '-c', f'psql -d {DD_DB} -AtF "\x1f" -c "{sql}"'], capture_output=True, text=True)
    if out.returncode:
        raise Exception(out.stderr)
    return [l.split('\x1f') for l in out.stdout.split('\n') if l]


TTYPE = {t: ty for t, ty in D['tables']}
VIEWS = {t for t, ty in D['tables'] if ty == 'VIEW'}
ROWS = {t: int(n) for t, n in D['rows']}
COLS = collections.defaultdict(list)
for r in D['columns']:
    COLS[r[0]].append(r)

PK, UQ, FKS, CHK = {}, collections.defaultdict(set), [], []
for name, ct, t, cols, ft, fcols, dl, ul, df in D['cons']:
    if ct == 'p':
        PK[t] = cols.split(',')
    elif ct == 'u':
        UQ[t].add(tuple(cols.split(',')))
    elif ct == 'f':
        FKS.append((name, t, cols, ft, fcols, dl, ul))
    elif ct == 'c':
        CHK.append((name, t, cols, df))

IDXINFO = {r[0]: r for r in D['idxinfo']}
IDX = []
for t, iname, idef in D['indexes']:
    _, uniq, prim, pred, am = IDXINFO[iname]
    m = re.search(r'USING \w+ \((.*?)\)(?: WHERE .*)?$', idef)
    cols = m.group(1) if m else ''
    IDX.append((t, iname, uniq == 'true', prim == 'true', pred, am, cols, idef))
    if uniq == 'true' and prim != 'true':
        cl = [c.strip() for c in cols.split(',')]
        if all(re.fullmatch(r'[a-z_0-9]+', c) for c in cl):
            UQ[t].add(tuple(cl) + (('partial',) if pred else ()))

FK_OF = {}
for name, t, cols, ft, fcols, dl, ul in FKS:
    for i, c in enumerate(cols.split(',')):
        FK_OF[(t, c)] = f"{ft}.{fcols.split(',')[i]}"

# FK supporting index: an index whose first column is the FK column
FIRSTCOL = collections.defaultdict(set)
for t, iname, u, p, pred, am, cols, idef in IDX:
    first = cols.split(',')[0].strip()
    FIRSTCOL[t].add(first)

ACT = {'a': 'No action', 'r': 'Restrict', 'c': 'Cascade', 'n': 'Set null', 'd': 'Set default'}


def dtype(r):
    t, c, pos, dt, udt, clen, nprec, nscale, nullable, default, dtp = r
    if dt == 'ARRAY':
        return {'int4': 'integer', 'int8': 'bigint'}.get(udt.lstrip('_'), udt.lstrip('_')) + '[]'
    return {'timestamp with time zone': 'timestamptz', 'character varying': 'varchar', 'USER-DEFINED': udt}.get(dt, dt)


def length(r):
    t, c, pos, dt, udt, clen, nprec, nscale, nullable, default, dtp = r
    if clen:
        return clen
    if dt == 'numeric':
        return f'{nprec},{nscale}' if nprec else 'unconstrained'
    return ''


def clean_default(d):
    d = re.sub(r"::(text|regclass|jsonb|numeric|integer|bigint|date|timestamp with time zone|text\[\]|character varying)", '', d)
    return d


# ---------------------------------------------------------------- styles
HFILL = PatternFill('solid', fgColor='0F4761')
HFONT = Font(bold=True, color='FFFFFF', name='Segoe UI', size=10)
BFONT = Font(name='Segoe UI', size=9)
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')
thin = Side(style='thin', color='D0D7DE')
BORDER = Border(bottom=thin)


def sheet(wb, title, headers, rows, widths, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.fill = HFILL; c.font = HFONT; c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BFONT
            c.alignment = WRAP if (c.column - 1) in wrap_cols else TOP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C2' if len(headers) > 4 else 'B2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.row_dimensions[1].height = 30
    return ws


# ---------------------------------------------------------------- Tables
area_order = {a: i for i, a in enumerate(tm.AREAS)}
tables_sorted = sorted([t for t, _ in D['tables']], key=lambda t: (area_order[tm.AREA_OF[t]], t))
table_rows = []
pii_tables = set()
col_rows = []
pii_count = collections.Counter()
for t in tables_sorted:
    cset = {r[1] for r in COLS[t]}
    for r in COLS[t]:
        c = r[1]
        desc, _ = gen.describe(t, c, r[3], r[9])
        uq = ''
        for u in UQ[t]:
            ucols = [x for x in u if x != 'partial']
            if c in ucols:
                tag = 'Yes' if len(ucols) == 1 else 'Composite (' + ', '.join(ucols) + ')'
                if 'partial' in u:
                    tag += ', partial'
                uq = tag if not uq else uq
        p = pii.classify(t, c, cset)
        if p:
            pii_tables.add(t); pii_count[p] += 1
        col_rows.append([t, c, int(r[2]), dtype(r), length(r), 'Yes' if r[8] == 'YES' else 'No', clean_default(r[9]),
                         'Yes' if c in PK.get(t, []) else '', FK_OF.get((t, c), ''), uq, p,
                         M['colmig'].get(f'{t}.{c}', '') if t not in VIEWS else '', desc])
    area, mod, desc, ret = tm.table_info(t)
    table_rows.append([t, 'View' if t in VIEWS else 'Table', area, mod, desc, ', '.join(PK.get(t, [])) or ('(none, view)' if t in VIEWS else ''),
                       ROWS.get(t, '') if t not in VIEWS else '', len(COLS[t]),
                       M['created'].get(t, 'db/migrate.js' if t == 'schema_migrations' else ('0116_bank_reconciliation.sql' if t in VIEWS else '')),
                       ret, 'Yes' if t in pii_tables else ''])

# views: creating migration
for v in VIEWS:
    for f in sorted(__import__('os').listdir('/home/user/BDOI-OOTB/backend/src/db/migrations')):
        txt = open('/home/user/BDOI-OOTB/backend/src/db/migrations/' + f).read().lower()
        if re.search(r'create (or replace )?view (public\.)?' + v + r'\b', txt):
            for row in table_rows:
                if row[0] == v:
                    row[8] = f
            break

# ---------------------------------------------------------------- Relationships
rel_rows = []
for name, t, cols, ft, fcols, dl, ul in sorted(FKS, key=lambda x: (x[1], x[2])):
    first = cols.split(',')[0]
    rel_rows.append([name, t, cols.replace(',', ', '), ft, fcols.replace(',', ', '), ACT[dl], ACT[ul],
                     tm.AREA_OF[t], tm.AREA_OF[ft], 'Yes' if first in FIRSTCOL[t] else 'No',
                     gen.describe(t, first, [r for r in COLS[t] if r[1] == first][0][3], '')[0] if ',' not in cols else 'Composite key'])

# ---------------------------------------------------------------- Indexes
idx_rows = []
for t, iname, u, p, pred, am, cols, idef in sorted(IDX):
    kind = 'Primary key' if p else ('Unique' if u else 'Non-unique')
    idx_rows.append([t, iname, kind, am, cols, pred, idef])

# ---------------------------------------------------------------- Reference values
ref = []
seen = set()
# 1. code-defined and curated status lists (from the column descriptions written from the code)
for t in tables_sorted:
    for r in COLS[t]:
        c = r[1]
        key = f'{t}.{c}'
        if key in gen.TCOL and 'Values:' in gen.TCOL[key]:
            vals = gen.TCOL[key].split('Values:', 1)[1].strip()
            ref.append(['Status and code values', t, c, '', gen.TCOL[key].split('. Values:')[0], vals, 'Application code (status vocabulary)'])
            seen.add((t, c))
# 2. check constraints
for name, t, cols, df in sorted(CHK, key=lambda x: (x[1], x[0])):
    m = re.match(r"CHECK \(\(\(?(\w+) = ANY \(ARRAY\[(.*)\]\)\)\)?\)$", df)
    if m:
        vals = ', '.join(re.findall(r"'([^']*)'::text", m.group(2)))
        ref.append(['Check constraint (allowed values)', t, m.group(1), '', name, vals, 'Database CHECK constraint'])
        seen.add((t, m.group(1)))
    else:
        rule = re.sub(r'::(text|numeric|integer)', '', df)
        ref.append(['Check constraint (rule)', t, cols.replace(',', ', '), '', name, rule, 'Database CHECK constraint'])
# 3. migration comment value lists
for key, com in M['comments'].items():
    t, c = key.split('.', 1)
    if (t, c) in seen or t not in COLS or key in gen.NO_COMMENT:
        continue
    if '|' in com:
        vals = ', '.join(v.strip() for v in com.split('|') if v.strip())
        ref.append(['Status and code values', t, c, '', 'Documented in the migration', vals, 'Migration comment (' + M['colmig'].get(key, M['created'].get(t, '')) + ')'])
        seen.add((t, c))
# 4. values found in the loaded test database for other code columns
VALS = json.load(open('vals.json'))
for key, vals in sorted(VALS.items()):
    t, c = key.split('.', 1)
    if (t, c) in seen or not vals or t in VIEWS:
        continue
    vs = [a for a, b in vals if a not in ("''",)]
    if not vs or len(vs) > 25 or any(len(v) > 60 for v in vs):
        continue
    ref.append(['Status and code values', t, c, '', f'Values present in the {"reference data" if REFERENCE_ONLY else "loaded test database"} (not an exhaustive list)', ', '.join(vs), 'Reference data' if REFERENCE_ONLY else 'Test data'])
# 5. application settings
for key, grp, label, typ, value, editable in D['settings']:
    ref.append(['Application setting (app_settings)', 'app_settings', key, grp, label, value, f'Type {typ}; ' + ('editable' if editable == 'true' else 'maintained by the system') + '; seeded value'])
# 6. master types and generic master records
for code, label, cat, storage, tname, n in q("select code, label, category, storage, coalesce(table_name,''), (select count(*) from master_records r where r.type_code = t.code) from master_types t order by category, code"):
    ref.append(['Master type (master_types)', 'master_types', code, cat, label,
                f'storage {storage}' + (f', table {tname}' if tname else f', {n} records in master_records'), 'Seed / migration'])
for typ, code, name, status in q("select type_code, coalesce(code,''), coalesce(name,''), status from master_records where status <> 'deleted' order by type_code, code"):
    ref.append(['Master record (master_records)', typ, code, '', name, status, 'Seed (reference or sample data)'])
# 7. document numbering
for code, name, module, prefix, pattern, w, reset, start in q("select code, name, coalesce(module,''), prefix, pattern, seq_width, reset_rule, start_number from document_numbering order by code"):
    ref.append(['Document number series (document_numbering)', 'document_numbering', code, module, name, f'{pattern} (prefix {prefix}, {w} digits, reset {reset}, start {start})', 'Seed / migration'])
# 8. roles, tax codes, bank transaction types, checklist, scheduled jobs, posting events
for code, name in q("select code, name from roles order by id"):
    ref.append(['Role (roles)', 'roles', code, '', name, '', 'Seed'])
for code, desc, ttype, rate, atc, acc in q("select code, coalesce(description,''), tax_type, rate::text, coalesce(atc,''), coalesce(gl_account,'') from tax_codes order by tax_type, code"):
    ref.append(['Tax code (tax_codes)', 'tax_codes', code, ttype, desc, f'rate {rate}%' + (f', ATC {atc}' if atc else '') + (f', GL {acc}' if acc else ''), 'Seed'])
for code, name, direction, action in q("select code, name, direction, action from bank_transaction_types order by sort_order, code"):
    ref.append(['Bank transaction type (bank_transaction_types)', 'bank_transaction_types', code, direction, name, action, 'Seed / migration'])
for code, label, itype, sev in q("select code, label, item_type, severity from period_close_checklist order by sort_order"):
    ref.append(['Month-end checklist item (period_close_checklist)', 'period_close_checklist', code, itype, label, sev, 'Migration'])
for code, name, cron, enabled in q("select code, name, cron, enabled::text from scheduled_jobs order by code"):
    ref.append(['Scheduled job (scheduled_jobs)', 'scheduled_jobs', code, '', name, f'cron {cron}; ' + ('enabled' if enabled == 'true' else 'disabled'), 'Seed / migration'])
for ev, name, src, et in q("select distinct on (event_code) event_code, name, source, coalesce(entry_type,'') from posting_rules order by event_code, version desc"):
    ref.append(['Posting rule event (posting_rules)', 'posting_rules', ev, src, name, et, 'Migration 0131 and later'])

# ---------------------------------------------------------------- Summary
base_tables = [t for t in tables_sorted if t not in VIEWS]
n_cols_tables = sum(len(COLS[t]) for t in base_tables)
n_cols_all = sum(len(COLS[t]) for t in tables_sorted)
types = collections.Counter(dtype(r) for t in base_tables for r in COLS[t])
summary = [
    ['Database', f'PostgreSQL, schema public (database "{DD_DB}" built with the migrations and the reference seed data, no sample data)'
                 if REFERENCE_ONLY else f'PostgreSQL, schema public (loaded test database "{DD_DB}")'],
    ['Migrations applied', f"{len(q('select name from schema_migrations'))} (0001_core.sql to {q('select max(name) from schema_migrations')[0][0]})"],
    ['Tables', len(base_tables)],
    ['Views', len(VIEWS)],
    ['Columns in tables', n_cols_tables],
    ['Columns including views', n_cols_all],
    ['Foreign keys', len(FKS)],
    ['Foreign keys without an index on the first column', sum(1 for r in rel_rows if r[9] == 'No')],
    ['Indexes', len(IDX)],
    ['  of which primary keys', sum(1 for r in idx_rows if r[2] == 'Primary key')],
    ['  of which further unique indexes', sum(1 for r in idx_rows if r[2] == 'Unique')],
    ['  of which partial indexes', sum(1 for r in idx_rows if r[5])],
    ['Check constraints', len(CHK)],
    ['Unique constraints', sum(1 for c in D['cons'] if c[1] == 'u')],
    ['Triggers', len(D['triggers'])],
    ['Application settings (app_settings)', f"{len(D['settings'])} keys in {len({s[1] for s in D['settings']})} groups"],
    [f'Rows in all tables ({DB_WORDS})', sum(ROWS.values())],
    ['Columns flagged as personal data', sum(pii_count.values())],
    ['  ' + pii.P, pii_count[pii.P]],
    ['  ' + pii.S, pii_count[pii.S]],
    ['  ' + pii.C, pii_count[pii.C]],
    ['  ' + pii.J, pii_count[pii.J]],
    ['Tables holding personal data', len(pii_tables)],
    ['', ''],
    ['Data type', 'Columns (tables only)'],
] + [[k, v] for k, v in types.most_common()] + [['', ''], ['Functional area', f'Tables / rows ({DB_WORDS})']]
for a in tm.AREAS:
    ts = [t for t in base_tables if tm.AREA_OF[t] == a]
    vs = [t for t in tables_sorted if tm.AREA_OF[t] == a and t in VIEWS]
    summary.append([a, f"{len(ts)} tables" + (f" and {len(vs)} views" if vs else '') + f", {sum(ROWS.get(t, 0) for t in ts):,} rows"])
summary += [['', ''], ['Sheets', ''],
            ['Tables', 'One row per table and view: functional area, owning module, description, primary key, rows, creating migration, retention class'],
            ['Columns', 'One row per column: type, length, nullability, default, keys, uniqueness, personal data class, migration that added it, description'],
            ['Relationships', 'All foreign keys with delete and update rules and whether an index supports them'],
            ['Indexes', 'All indexes with type, method, columns and partial predicate'],
            ['Reference values', 'Status vocabularies, CHECK constraints, application settings, master types and records, number series, roles, tax codes, jobs, posting events'],
            ['', ''], ['Note', 'Row counts are those of the reference data a new database starts with (no sample data); transaction tables are empty.'
                       if REFERENCE_ONLY else 'Row counts are those of the loaded test database (reference data and sample data). A production database starts with reference data only.']]

wb = Workbook()
ws = wb.active
ws.title = 'Summary'
ws.append(['BrokerVerse OOTB: Data Dictionary', ''])
ws['A1'].font = Font(bold=True, size=14, color='0F4761', name='Segoe UI')
ws.append(['iorta TechNXT, version 1.0, 03 October 2026', ''])
ws['A2'].font = Font(italic=True, size=9, name='Segoe UI')
ws.append(['', ''])
ws.append(['Item', 'Value'])
for i in (1, 2):
    c = ws.cell(row=4, column=i); c.fill = HFILL; c.font = HFONT
for r in summary:
    ws.append(r)
    if r[0] in ('Data type', 'Functional area', 'Sheets'):
        for i in (1, 2):
            c = ws.cell(row=ws.max_row, column=i); c.fill = PatternFill('solid', fgColor='DCE6EE'); c.font = Font(bold=True, name='Segoe UI', size=10)
    else:
        for i in (1, 2):
            ws.cell(row=ws.max_row, column=i).font = BFONT
            ws.cell(row=ws.max_row, column=i).alignment = WRAP
ws.column_dimensions['A'].width = 52
ws.column_dimensions['B'].width = 110
ws.freeze_panes = 'A5'

sheet(wb, 'Tables', ['Table', 'Type', 'Functional area', 'Owning module', 'Description', 'Primary key', 'Rows (reference data)' if REFERENCE_ONLY else 'Rows (loaded test DB)',
                     'Columns', 'Created by migration', 'Retention class', 'Personal data'],
      table_rows, [30, 8, 26, 22, 70, 22, 12, 9, 32, 13, 10], wrap_cols=(4,))
sheet(wb, 'Columns', ['Table', 'Column', 'Position', 'Data type', 'Length / precision', 'Nullable', 'Default', 'Primary key',
                      'Foreign key to', 'Unique', 'Personal data (RA 10173)', 'Added by migration', 'Description'],
      col_rows, [28, 26, 8, 13, 11, 9, 26, 8, 30, 16, 22, 30, 80], wrap_cols=(12,))
sheet(wb, 'Relationships', ['Constraint', 'Child table', 'Child column(s)', 'Parent table', 'Parent column(s)', 'On delete', 'On update',
                            'Child functional area', 'Parent functional area', 'Index on FK column', 'Meaning of the child column'],
      rel_rows, [40, 28, 24, 28, 18, 11, 11, 26, 26, 10, 70], wrap_cols=(10,))
sheet(wb, 'Indexes', ['Table', 'Index', 'Kind', 'Method', 'Columns / expression', 'Partial (WHERE)', 'Definition'],
      idx_rows, [28, 44, 12, 8, 40, 40, 90], wrap_cols=(4, 5, 6))
sheet(wb, 'Reference values', ['Category', 'Table / list', 'Column / key / code', 'Group / qualifier', 'Label / meaning', 'Values / value', 'Source'],
      ref, [30, 26, 34, 18, 60, 60, 30], wrap_cols=(4, 5))
wb.save(OUT)
json.dump({'tables': len(base_tables), 'views': len(VIEWS), 'columns': n_cols_tables, 'columns_all': n_cols_all, 'fks': len(FKS),
           'indexes': len(IDX), 'checks': len(CHK), 'ref_rows': len(ref), 'pii': dict(pii_count), 'pii_tables': sorted(pii_tables),
           'settings': len(D['settings'])}, open('stats.json', 'w'), indent=1)
print(open('stats.json').read()[:600])
