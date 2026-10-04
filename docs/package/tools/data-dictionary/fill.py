"""Refreshes the generated parts of source/data-dictionary.md from the workbook and the database (see ddpaths.py).

Replaced in place (the rest of the document is written by hand):
  - the table under "## Database at a glance"
  - the table under "## Overview" (tables per functional area)
  - the class counts under "## Personal data under the Data Privacy Act" and the table under
    "## Tables and columns holding personal data"
  - the retention class counts under "## Retention"
  - the chapter "# Appendix A: Table list"
It prints the further counts quoted in the text (identifiers, money columns, audit columns, keys) so that the hand-written
sentences can be checked against them.
"""
import collections
import json
import re

from openpyxl import load_workbook

import pii
import tables_meta as tm
from ddpaths import WORKBOOK, SOURCE_MD, psql

wb = load_workbook(WORKBOOK)
T = list(wb['Tables'].iter_rows(min_row=2, values_only=True))
C = list(wb['Columns'].iter_rows(min_row=2, values_only=True))
R = list(wb['Relationships'].iter_rows(min_row=2, values_only=True))
I = list(wb['Indexes'].iter_rows(min_row=2, values_only=True))
st = json.load(open('stats.json'))
D = json.load(open('db.json'))
views = {t for t, ty in D['tables'] if ty == 'VIEW'}
Cb = [c for c in C if c[0] not in views]


def q1(sql):
    return psql(sql)[0][0]


def n(x):
    return f'{x:,}'


nmig = int(q1('select count(*) from schema_migrations'))
lastmig = q1('select max(name) from schema_migrations')
fk_rules = collections.Counter(r[5] for r in R)
pii_cols = collections.Counter(c[10] for c in C if c[10])
pii_tables = {c[0] for c in C if c[10]}
ph = {
    'migrations': nmig, 'last migration': lastmig,
    'id text': sum(1 for c in Cb if c[1] == 'id' and c[7] == 'Yes' and c[3] == 'text'),
    'id serial': sum(1 for c in Cb if c[1] == 'id' and c[7] == 'Yes' and c[3] in ('integer', 'bigint')),
    'numbering series': q1('select count(*) from document_numbering'),
    'numeric(14,2)': sum(1 for c in Cb if c[3] == 'numeric' and c[4] == '14,2'),
    'numeric(16,2)': sum(1 for c in Cb if c[3] == 'numeric' and c[4] == '16,2'),
    'date': sum(1 for c in Cb if c[3] == 'date'), 'timestamptz': sum(1 for c in Cb if c[3] == 'timestamptz'),
    'jsonb': sum(1 for c in Cb if c[3] == 'jsonb'),
    'created_at': sum(1 for c in Cb if c[1] == 'created_at'), 'created_by': sum(1 for c in Cb if c[1] == 'created_by'),
    'updated_at': sum(1 for c in Cb if c[1] == 'updated_at'), 'updated_by': sum(1 for c in Cb if c[1] == 'updated_by'),
    'status': sum(1 for c in Cb if c[1] == 'status'),
    'fk': len(R), 'fk rules': dict(fk_rules), 'fk without index': sum(r[9] == 'No' for r in R), 'fk to users': sum(r[3] == 'users' for r in R),
    'checks': st['checks'], 'pii columns': sum(pii_cols.values()), 'pii tables': len(pii_tables),
}
idx_kinds = collections.Counter(i[2] for i in I)
partial = sum(bool(i[5]) for i in I)
gin = sum(1 for i in I if (i[3] or '').lower() == 'gin')
n_tables = sum(1 for t in T if t[1] == 'Table')
n_cols_t = len(Cb)

# ---------------------------------------------------------------- generated blocks
glance = '\n'.join([
    '| Item | Value (reference data only) |', '|---|---|',
    '| Database engine | PostgreSQL, schema `public` |',
    f'| Migrations applied | {nmig} (0001_core.sql to {lastmig}) |',
    f'| Tables | {n_tables} |',
    f"| Views | {len(views)} ({', '.join(f'`{v}`' for v in sorted(views))}) |",
    f'| Columns | {n(n_cols_t)} in tables, {n(len(C))} including the views |',
    f"| Foreign keys | {len(R)} ({fk_rules.get('No action', 0)} no action, {fk_rules.get('Cascade', 0)} cascade, {fk_rules.get('Set null', 0)} set null) |",
    f"| Indexes | {len(I)} ({idx_kinds.get('Primary key', 0)} primary keys, {idx_kinds.get('Unique', 0)} further unique, {partial} partial, {gin} GIN) |",
    f"| CHECK constraints | {st['checks']} |",
    f"| Application settings | {len(D['settings'])} keys in {len({s[1] for s in D['settings']})} groups |",
    f"| Rows in all tables | {n(sum((t[6] or 0) for t in T))} (reference data a new database starts with) |",
    f"| Columns holding personal data | {sum(pii_cols.values())} in {len(pii_tables)} tables |",
])

KEY = {
    'Party and client': 'clients, insurance_companies, client_beneficial_owners', 'Product and rating': 'products, product_templates, premium_charge_rules',
    'Sales and quotation': 'leads, quotes, payment_links', 'Placement': 'broker_slips, insurer_offers, placements, risk_participants',
    'Policy and endorsement': 'policies, endorsements, cover_notes', 'Billing, receipts and payments': 'receivables, receipts, receipt_applications, disbursements',
    'Remittance and insurer accounting': 'remittances, direct_bill_items, commission_debit_notes', 'Commission': 'commissions, commission_rates, override_agreements',
    'Claims': 'claims, claim_settlement_movements, claim_document_items', 'Renewals': 'renewals, renewal_quotes',
    'General ledger and period end': 'journal_vouchers, journal_lines, gl_accounts, posting_rules, accounting_periods',
    'Bank and insurer reconciliation': 'bank_statements, bank_rec_matches, bank_reconciliations, insurer_statements', 'Credit control': 'collection_items, premium_instalment_plans',
    'Reinsurance': 'reinsurance_treaties, cessions, fac_placements', 'Incentives': 'incentive_programs, incentive_results',
    'Security and audit': 'users, roles, audit_log', 'Configuration and schedules': 'app_settings, master_records, document_numbering, scheduled_jobs, data_load_batches',
    'Distribution and marketing': 'lead_assignment_rules, distribution_channels, motor_programmes, campaigns',
    'Fleet, marine and motor claims': 'fleet_schedules, open_covers, claim_repair_estimates',
    'Payables and fixed assets': 'supplier_invoices, supplier_payments, fixed_assets',
    'BIR returns and invoicing': 'bir_return_filings, sales_invoices, eis_submissions',
    'Compliance': 'aml_risk_assessments, aml_cases, compliance_licences, complaints, personal_data_breaches',
    'Integrations': 'integration_connectors, integration_outbox, ctpl_authentications, bank_payment_batches',
}
ov = ['| Functional area | Tables | Key tables |', '|---|---|---|']
for a in tm.AREAS:
    nt = sum(1 for t in T if t[2] == a and t[1] == 'Table')
    nv = sum(1 for t in T if t[2] == a and t[1] == 'View')
    ov.append(f"| {a} | {nt}{f' + {nv} views' if nv else ''} | " + ', '.join(f'`{k.strip()}`' for k in KEY[a].split(',')) + ' |')
overview = '\n'.join(ov)

by = collections.defaultdict(lambda: collections.defaultdict(list))
for c in C:
    if c[10]:
        by[c[0]][c[10]].append(c[1])
short = {pii.P: 'Personal', pii.S: 'Sensitive', pii.C: 'Credential', pii.J: 'May contain'}
lines = ['| Table | Columns holding personal data |', '|---|---|']
area_order = {a: i for i, a in enumerate(tm.AREAS)}
for t in sorted(by, key=lambda t: (area_order[tm.AREA_OF[t]], t)):
    parts = [f'{short[k]}: ' + ', '.join(by[t][k]) for k in (pii.P, pii.S, pii.C, pii.J) if by[t].get(k)]
    lines.append(f'| `{t}` | ' + '<br>'.join(parts) + ' |')
pii_table = '\n'.join(lines)

ret = collections.Counter(t[9] for t in T if t[1] == 'Table')
app = ['# Appendix A: Table list', '',
       'This appendix lists every table and view with a one-line description and the number of rows of reference data a new '
       'database starts with (transaction tables are empty). The columns of each table are in the workbook, sheet Columns.', '']
for a in tm.AREAS:
    app += [f'## {a}', '', '| Table | Description | Rows |', '|---|---|---|']
    for t in T:
        if t[2] == a:
            rows = 'view' if t[1] == 'View' else f'{t[6]:,}'
            app.append(f'| `{t[0]}` | {t[4]} | {rows} |')
    app.append('')
appendix = '\n'.join(app).rstrip() + '\n'

# ---------------------------------------------------------------- write into the document
s = open(SOURCE_MD).read()


def replace_table_after(s, heading, new):
    """Replace the first pipe table after the heading line."""
    i = s.index(heading)
    m = re.compile(r'(^\|.*\n)+', re.M).search(s, i)
    return s[:m.start()] + new + '\n' + s[m.end():]


s = replace_table_after(s, '## Database at a glance', glance)
s = replace_table_after(s, '## Overview', overview)
s = replace_table_after(s, '## Tables and columns holding personal data', pii_table)
i = s.index('## Personal data under the Data Privacy Act')
m = re.compile(r'(^\|.*\n)+', re.M).search(s, i)
cls = m.group(0)
for k in (pii.P, pii.S, pii.C):
    label = {pii.P: 'Personal information', pii.S: 'Sensitive personal information', pii.C: 'Credential / security secret'}[k]
    cls = re.sub(rf'(^\| {re.escape(label)} \|.*\| )\d+( \|)$', rf'\g<1>{pii_cols.get(k, 0)}\2', cls, flags=re.M)
cls = re.sub(r'(^\| May contain personal information \|.*\| )\d+( \|)$', rf'\g<1>{pii_cols.get(pii.J, 0)}\2', cls, flags=re.M)
s = s[:m.start()] + cls + s[m.end():]
i = s.index('Each table also carries a retention class')
m = re.compile(r'(^\|.*\n)+', re.M).search(s, i)
rt = m.group(0)
for label in ('Financial', 'Transaction', 'Reference', 'Log / audit', 'Temporary'):
    rt = re.sub(rf'(^\| {re.escape(label)} \| )\d+( \|)', rf'\g<1>{ret.get(label, 0)}\2', rt, flags=re.M)
s = s[:m.start()] + rt + s[m.end():]
s = s[:s.index('# Appendix A: Table list')] + appendix
open(SOURCE_MD, 'w').write(s)
print({k: v for k, v in ph.items()})
print('retention', dict(ret), 'pii', dict(pii_cols))
