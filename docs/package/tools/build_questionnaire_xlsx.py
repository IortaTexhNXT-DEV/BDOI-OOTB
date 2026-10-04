"""Builds BrokerVerse_Security_Due_Diligence_Questionnaire.xlsx from the question tables of the source document,
so the Word version and the workbook always carry the same answers.

    python3 build_questionnaire_xlsx.py ../source/security-due-diligence-questionnaire.md \
        ../04_Onboarding_and_Go_Live/BrokerVerse_Security_Due_Diligence_Questionnaire.xlsx
"""
import re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC = sys.argv[1] if len(sys.argv) > 1 else '../source/security-due-diligence-questionnaire.md'
OUT = sys.argv[2] if len(sys.argv) > 2 else '../04_Onboarding_and_Go_Live/BrokerVerse_Security_Due_Diligence_Questionnaire.xlsx'
NAVY = '0B2A4A'
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
HEAD_FONT = Font(name='Segoe UI', bold=True, color='FFFFFF', size=10)
BODY_FONT = Font(name='Segoe UI', size=9)
THIN = Side(style='thin', color='BFC7D1')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
FILLS = {'Yes': 'E2F0D9', 'Partial': 'FFF2CC', 'No': 'F8CBAD', 'Planned': 'DDEBF7', 'Broker': 'EDEDED', 'N/A': 'FFFFFF'}


def cells(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def plain(text):
    return re.sub(r'\*\*([^*]+)\*\*', r'\1', text).replace('`', '')


def parse(path):
    chapter, questions, evidence, open_items, block = '', [], [], [], None
    for line in open(path, encoding='utf-8'):
        if line.startswith('# '):
            chapter = line[2:].strip()
            block = None
        elif line.startswith('|') and not re.match(r'^\|\s*-', line):
            row = cells(line)
            if row[0] in ('Ref', 'Code', '#'):
                block = row[0]
                continue
            if block == 'Ref' and len(row) == 5:
                questions.append([row[0], chapter] + [plain(c) for c in row[1:]])
            elif block == 'Code' and len(row) == 2:
                evidence.append([plain(c) for c in row])
            elif block == '#' and chapter.startswith('Open items'):
                open_items.append([plain(c) for c in row])
    return questions, evidence, open_items


def sheet(wb, title, header, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = BODY_FONT, WRAP, BOX
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    return ws


def main():
    questions, evidence, open_items = parse(SRC)
    wb = Workbook()
    header = ['Ref', 'Area', 'Question', 'Response', 'Answer', 'Evidence document reference', 'Reviewer comment', 'Follow-up owner']
    ws = sheet(wb, 'Questionnaire', header, [q + ['', ''] for q in questions], [9, 24, 38, 11, 70, 30, 30, 18], first=True)
    for row in ws.iter_rows(min_row=2):
        fill = FILLS.get(row[3].value)
        if fill:
            row[3].fill = PatternFill('solid', fgColor=fill)
        if '[to confirm' in (row[4].value or ''):
            row[4].font = Font(name='Segoe UI', size=9, color='9C0006')
    dv = DataValidation(type='list', formula1='"Yes,Partial,No,Planned,Broker,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'D2:D{ws.max_row}')
    sheet(wb, 'Evidence documents', ['Code', 'Document'], evidence, [14, 80])
    sheet(wb, 'Open items', ['#', 'Item', 'Owner'], open_items, [5, 80, 40])
    legend = [['Yes', 'In place today'], ['Partial', 'In place with a stated limit'], ['No', 'Not in place'],
              ['Planned', 'On the product backlog'], ['Broker', "The broker's own control, outside iorta TechNXT's scope"],
              ['N/A', 'Not applicable'], ['[to confirm]', 'Organisational fact to be answered by iorta TechNXT management before release (red text)']]
    lg = sheet(wb, 'Legend', ['Response', 'Meaning'], legend, [14, 80])
    for row in lg.iter_rows(min_row=2):
        fill = FILLS.get(row[0].value)
        if fill:
            row[0].fill = PatternFill('solid', fgColor=fill)
    wb.save(OUT)
    print(f'{OUT}: {len(questions)} questions, {len(evidence)} evidence documents, {len(open_items)} open items')


if __name__ == '__main__':
    main()
