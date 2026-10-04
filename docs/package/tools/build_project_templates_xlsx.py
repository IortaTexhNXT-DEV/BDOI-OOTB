"""Builds the project control workbooks of the implementation:

    python3 build_project_templates_xlsx.py ../05_Delivery

  BrokerVerse_Fit_Gap_Register.xlsx   Instructions, Register, Lists
  BrokerVerse_RAID_Log_Template.xlsx  Instructions, Summary, Risks, Assumptions, Issues, Dependencies, Lists

The classes and fields follow the Implementation Approach and Plan (fit-gap classes Fit, Configure, Procedure, Gap;
RAID log fields; steering committee escalation of high items). Rows marked "Example" show how a line is filled in and
are deleted at mobilisation.
"""
import os
import sys
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '05_Delivery')
NAVY = '0B2A4A'
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
HEAD_FONT = Font(name='Segoe UI', bold=True, color='FFFFFF', size=10)
BODY_FONT = Font(name='Segoe UI', size=10)
TITLE_FONT = Font(name='Segoe UI', bold=True, color=NAVY, size=14)
BOLD = Font(name='Segoe UI', bold=True, size=10, color=NAVY)
THIN = Side(style='thin', color='BFC7D1')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
EXAMPLE_FILL = PatternFill('solid', fgColor='F3F6FA')
ROWS = 200  # rows prepared with validation and formulas


def table(ws, headers, widths, rows=(), start=1):
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=start, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[start].height = 32
    for i, r in enumerate(rows, start + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border, c.alignment = BODY_FONT, BOX, WRAP
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f'A{start}:{get_column_letter(len(headers))}{start + max(len(rows), 1) + ROWS}'


def boxes(ws, first, last, cols):
    for r in range(first, last + 1):
        for j in range(1, cols + 1):
            c = ws.cell(row=r, column=j)
            c.border, c.alignment = BOX, WRAP
            if c.font != BODY_FONT and not c.font.bold:
                c.font = BODY_FONT


def listsheet(wb, lists):
    ws = wb.create_sheet('Lists')
    for j, (name, values) in enumerate(lists.items(), 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        ws.column_dimensions[get_column_letter(j)].width = 26
        for i, v in enumerate(values, 2):
            ws.cell(row=i, column=j, value=v).font = BODY_FONT
    ws.freeze_panes = 'A2'
    return {name: f"Lists!${get_column_letter(j)}$2:${get_column_letter(j)}${len(v) + 1}" for j, (name, v) in enumerate(lists.items(), 1)}


def validate(ws, ref, col, first, last):
    dv = DataValidation(type='list', formula1=f'={ref}', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{col}{first}:{col}{last}')


def instructions(ws, title, intro, rows):
    ws['A1'] = title
    ws['A1'].font = TITLE_FONT
    ws['A2'] = intro
    ws['A2'].font, ws['A2'].alignment = BODY_FONT, WRAP
    ws.merge_cells('A2:B2')
    ws.row_dimensions[2].height = 60
    table(ws, ['Item', 'Guidance'], [28, 110], rows, start=4)
    ws.auto_filter.ref = None


# ------------------------------------------------------------------ fit-gap register
def fit_gap():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Instructions'
    instructions(ws, 'BrokerVerse Fit-Gap Register',
                 'One line per requirement found in the discovery workshops. Each line is classed, owned and decided. '
                 'The register is signed by the process owners at the end of discovery (Small week 2, Medium week 3, '
                 'Large week 5) and is then kept up to date until hypercare exit. Lines marked Example show how a line is '
                 'filled in; delete them before the first workshop.',
                 [
                     ('Requirement ID', 'FG-<area code>-<number>, for example FG-ACC-012. Never reuse an ID; a dropped line keeps its ID with status Withdrawn.'),
                     ('Area', 'The process area of the discovery workshops (list on the Lists sheet).'),
                     ('Requirement', 'What the broker needs, in its own words, with the volume or rule that matters.'),
                     ('Classification', 'Fit: delivered as it is. Configure: met by a setting or master data. Procedure: the broker changes how it works (workaround). Gap: needs a change to the product.'),
                     ('Solution', 'For Fit and Configure: the screen, setting or master (for example Master > Configuration > Billing, Collections & Credit). For Procedure: the new way of working, which goes into training. For Gap: the proposed change.'),
                     ('CR reference', 'For a Gap only: the change request number (Change Request Procedure and Form). A Gap is not in the go-live scope unless the steering committee approves it.'),
                     ('Effort', 'For Configure: hours of configuration. For Gap: the estimate in person-days from the change request.'),
                     ('Priority', 'Must have for go-live, Should have, Could have (after go-live).'),
                     ('Owner', 'One named person of the broker or of iorta TechNXT who closes the line.'),
                     ('Status', 'Open, Under analysis, Decided, Configured, Closed, Withdrawn.'),
                     ('Decision date', 'The date the class and solution were agreed (workshop or steering committee).'),
                     ('Decided by', 'The process owner, or the steering committee for a Gap.'),
                     ('Summary', 'The counts at the top of the Register sheet update themselves; report them in the weekly status report.'),
                 ])
    reg = wb.create_sheet('Register')
    reg['A1'] = 'Lines by classification:'
    reg['A1'].font = BOLD
    for j, cls in enumerate(['Fit', 'Configure', 'Procedure', 'Gap'], 2):
        reg.cell(row=1, column=2 * j - 2, value=cls).font = BOLD
        reg.cell(row=1, column=2 * j - 1, value=f'=COUNTIF($D$4:$D${ROWS + 3},"{cls}")').font = BODY_FONT
    headers = ['Requirement ID', 'Area', 'Requirement', 'Classification', 'Solution', 'CR reference', 'Effort',
               'Priority', 'Owner', 'Status', 'Decision date', 'Decided by', 'Remarks']
    widths = [15, 18, 46, 15, 46, 13, 10, 14, 18, 14, 13, 18, 30]
    examples = [
        ('FG-ACC-001', 'Accounting and tax', 'Example. VAT-registered broker: output VAT on commission debit notes at 12%.', 'Fit',
         'Delivered: direct bill debit notes add VAT at the tax code direct_bill.commission_vat_code (VAT12-OUT).', '', '', 'Must have', '[Accounting Manager]', 'Decided', '', '[name]', 'Example: delete'),
        ('FG-BIL-002', 'Billing and collection', 'Example. Official receipts e-mailed to the client with the PDF when recorded.', 'Configure',
         'Switch on receipts.email_on_record (Master > Configuration > Billing, Collections & Credit > Receipts); SMTP mailbox set.', '', '1 hour', 'Should have', '[System Administrator]', 'Decided', '', '[name]', 'Example: delete'),
        ('FG-REG-003', 'Regulatory reporting', 'Example. Quarterly report to the Insurance Commission in its prescribed format.', 'Gap',
         'Figures from the financial statements and production reports; a report in the IC format needs a change request.', 'CR-[number]', '[days]', 'Could have', '[Compliance officer]', 'Under analysis', '', 'Steering committee', 'Example: delete'),
        ('FG-REM-004', 'Remittance', 'Example. Remittances above PHP 1,000,000.00 approved by the Accounting Manager.', 'Configure',
         'Authority Matrix, transaction type Remittance approval: Accounting PHP 1,000,000.00, Accounting Manager without limit (delivered default).', '', '0.5 hour', 'Must have', '[Accounting Manager]', 'Decided', '', '[name]', 'Example: delete'),
        ('FG-POL-005', 'Policy and servicing', 'Example. Brokers send the policy schedule by courier only.', 'Procedure',
         'Print from Operations > Policy > Policy detail > Policy schedule; procedure added to the Operations training.', '', '', 'Should have', '[Operations head]', 'Decided', '', '[name]', 'Example: delete'),
    ]
    table(reg, headers, widths, examples, start=3)
    boxes(reg, 4, ROWS + 3, len(headers))
    for r in range(4, 4 + len(examples)):
        for j in range(1, len(headers) + 1):
            reg.cell(row=r, column=j).fill = EXAMPLE_FILL
    for r in range(4, ROWS + 4):
        reg.cell(row=r, column=11).number_format = 'DD MMM YYYY'
    refs = listsheet(wb, {
        'Area': ['Organisation and users', 'Sales and quotation', 'Placement', 'Policy and servicing', 'Claims', 'Renewals',
                 'Billing and collection', 'Remittance', 'Commission and incentives', 'Accounting and tax', 'Period end',
                 'Bank and insurer reconciliation', 'Reinsurance', 'Reports and dashboards', 'Regulatory reporting',
                 'Data privacy', 'Integrations', 'Data migration', 'Security and access', 'Hosting and operations'],
        'Classification': ['Fit', 'Configure', 'Procedure', 'Gap'],
        'Priority': ['Must have', 'Should have', 'Could have'],
        'Status': ['Open', 'Under analysis', 'Decided', 'Configured', 'Closed', 'Withdrawn'],
    })
    for name, col in (('Area', 'B'), ('Classification', 'D'), ('Priority', 'H'), ('Status', 'J')):
        validate(reg, refs[name], col, 4, ROWS + 3)
    reg.conditional_formatting.add(f'D4:D{ROWS + 3}', CellIsRule(operator='equal', formula=['"Gap"'], fill=PatternFill('solid', fgColor='FCE4D6')))
    reg.freeze_panes = 'B4'
    wb.active = 1
    wb.save(os.path.join(OUT, 'BrokerVerse_Fit_Gap_Register.xlsx'))


# ------------------------------------------------------------------ RAID log
def raid():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Instructions'
    instructions(ws, 'BrokerVerse RAID Log',
                 'One shared log of risks, assumptions, issues and dependencies, kept by the iorta TechNXT project manager '
                 'and reviewed in every weekly status meeting. Items rated High go to the steering committee. The risks '
                 'and dependencies pre-filled from the Implementation Approach and Plan are typical of a BrokerVerse '
                 'implementation: keep, change or close them at mobilisation.',
                 [
                     ('ID', 'R-nnn, A-nnn, I-nnn or D-nnn with a running number. Never reuse an ID.'),
                     ('Probability (risks)', '1 Rare (under 10%), 2 Unlikely (10 to 30%), 3 Possible (30 to 50%), 4 Likely (50 to 80%), 5 Almost certain (over 80%).'),
                     ('Impact (risks and issues)', '1 Negligible, 2 Minor (a task slips), 3 Moderate (a milestone slips up to one week), 4 Major (go-live or budget at risk), 5 Severe (go-live missed, compliance or data loss).'),
                     ('Score and rating', 'Score = probability x impact. High: 15 to 25. Medium: 8 to 14. Low: 1 to 7. Calculated on the sheet; do not type over the formula.'),
                     ('Issue priority', 'An issue has happened. Priority follows impact: 4 or 5 High, 3 Medium, 1 or 2 Low.'),
                     ('Assumptions', 'Each assumption has a date by which it must be validated; an assumption found false becomes an issue or a risk.'),
                     ('Dependencies', 'What the project needs from someone outside the team (the broker, an insurer, a bank, the hosting provider), by when, and its state.'),
                     ('Status', 'Open, In progress, Closed. Fill the closure date and keep closed items on the sheet.'),
                     ('Summary', 'The Summary sheet counts open items by type and rating for the status report and the steering committee.'),
                 ])
    refs_data = {
        'Status': ['Open', 'In progress', 'Closed'],
        'Score': ['1', '2', '3', '4', '5'],
        'Category': ['Scope', 'Schedule', 'Data', 'People', 'Technical', 'Integration', 'Compliance', 'Commercial', 'Hosting'],
        'Priority': ['High', 'Medium', 'Low'],
        'Dependency state': ['On track', 'At risk', 'Late', 'Delivered'],
    }
    rating = lambda r: f'=IF(OR(G{r}="",H{r}=""),"",IF(G{r}*H{r}>=15,"High",IF(G{r}*H{r}>=8,"Medium","Low")))'
    score = lambda r: f'=IF(OR(G{r}="",H{r}=""),"",G{r}*H{r})'

    # risks
    rk = wb.create_sheet('Risks')
    rh = ['ID', 'Risk', 'Cause', 'Effect', 'Category', 'Owner', 'Probability', 'Impact', 'Score', 'Rating',
          'Mitigation', 'Due date', 'Status', 'Raised on', 'Closed on']
    rw = [8, 38, 28, 30, 13, 18, 11, 9, 8, 10, 46, 12, 12, 12, 12]
    pre = [
        ('Data arrives late or incomplete', 'Extracts from the old system not ready', 'Mock loads slip; go-live moves', 'Data', 4, 4,
         'Data requests in week 1; templates handed out at kick-off; data readiness tracked weekly; steering committee escalation after one week of delay'),
        ('Poor data quality in the old system', 'Duplicates, missing fields, old balances', 'Rejected rows; receivables or balances do not reconcile', 'Data', 3, 4,
         'Two to four mock loads with reconciliation; the broker cleanses at source; rejected-row reports from each upload'),
        ('Requests for customisation', 'Old system habits; new wishes', 'Scope and timeline grow', 'Scope', 3, 4,
         'Configure-first principle; fit-gap classes; change request process with steering committee approval'),
        ('Key users unavailable', 'Business as usual takes priority', 'Decisions late; weak UAT and training', 'People', 3, 3,
         'Named key users with committed time; backups named; workshops scheduled at kick-off'),
        ('Opening balances and open items do not agree', 'Cut-off differences', 'Ledger wrong from day one', 'Data', 2, 5,
         'Premiums Receivable balance equals the open item total; Due to Insurers reconciled to insurer statements; Accounting Manager signs off'),
        ('Tax rates or ATC not confirmed', 'Tax adviser not engaged in time', 'Wrong VAT or withholding', 'Compliance', 2, 4,
         'Tax adviser review in the configuration phase; VAT Summary checked in UAT'),
        ('Official receipt numbering out of line with the Authority to Print', 'Last numbers not supplied', 'Compliance exposure', 'Compliance', 2, 4,
         'Numbering set from the last number used; checked in UAT'),
        ('E-mail or payment gateway credentials late', 'Provider onboarding', 'Approval links, receipts by e-mail, password resets and online payments not working', 'Integration', 3, 3,
         'Requested in mobilisation; the system works without the gateway; sandbox first'),
        ('Users not ready at go-live', 'Training too early or skipped', 'Errors and slow work', 'People', 2, 3,
         'Role-based training, assessment, floor walkers in the first week'),
        ('First month-end close outside hypercare', 'Go-live late in the month', 'Close issues found late', 'Schedule', 2, 3,
         'Hypercare extends to the first close'),
    ]
    rows = [(f'R-{i:03d}', r[0], r[1], r[2], r[3], '[owner]', r[4], r[5], None, None, r[6], None, 'Open', None, None) for i, r in enumerate(pre, 1)]
    table(rk, rh, rw, rows)
    boxes(rk, 2, ROWS + 1, len(rh))
    for r in range(2, ROWS + 2):
        rk.cell(row=r, column=9, value=score(r))
        rk.cell(row=r, column=10, value=rating(r))
        for c in (12, 14, 15):
            rk.cell(row=r, column=c).number_format = 'DD MMM YYYY'
    for name, cols in (('Score', 'GH'), ('Category', 'E'), ('Status', 'M')):
        for col in cols:
            dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True) if name == 'Score' else None
            if dv:
                rk.add_data_validation(dv); dv.add(f'{col}2:{col}{ROWS + 1}')
    rk.conditional_formatting.add(f'J2:J{ROWS + 1}', CellIsRule(operator='equal', formula=['"High"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    rk.conditional_formatting.add(f'J2:J{ROWS + 1}', CellIsRule(operator='equal', formula=['"Medium"'], fill=PatternFill('solid', fgColor='FFE699')))
    rk.conditional_formatting.add(f'J2:J{ROWS + 1}', CellIsRule(operator='equal', formula=['"Low"'], fill=PatternFill('solid', fgColor='C6E0B4')))

    # assumptions
    asn = wb.create_sheet('Assumptions')
    ah = ['ID', 'Assumption', 'Basis', 'Effect if false', 'Owner', 'Validate by', 'Validated', 'Status', 'Remarks']
    aw = [8, 44, 30, 34, 18, 13, 11, 12, 30]
    arows = [
        ('A-001', 'The size of the implementation (Small, Medium or Large) agreed at mobilisation stays valid once the data volumes are known.', 'Implementation Approach and Plan, sizing', 'Plan and price re-baselined', '[PM]', None, '', 'Open', ''),
        ('A-002', 'The broker provides its SMTP mailbox before the integrations phase.', 'Dependencies on the broker', 'E-mails stay queued in the outbox', '[IT head]', None, '', 'Open', ''),
        ('A-003', 'The delivered chart of accounts is used with the accountant\'s mapping, not replaced.', 'Configure-first principle', 'Posting rules and account determination redone', '[Accounting Manager]', None, '', 'Open', ''),
    ]
    table(asn, ah, aw, arows)
    boxes(asn, 2, ROWS + 1, len(ah))
    for r in range(2, ROWS + 2):
        asn.cell(row=r, column=6).number_format = 'DD MMM YYYY'

    # issues
    iss = wb.create_sheet('Issues')
    ih = ['ID', 'Issue', 'Effect', 'Category', 'Owner', 'Impact', 'Priority', 'Action', 'Due date', 'Status', 'Raised on', 'Closed on', 'From risk']
    iw = [8, 40, 32, 13, 18, 8, 10, 44, 12, 12, 12, 12, 10]
    table(iss, ih, iw, [('I-001', '[Example: bank statement export of one bank does not match the delivered formats]', '[Bank reconciliation not testable in SIT]', 'Integration', '[owner]', 3, None, '[Add an insurer or bank statement format in Master > Finance]', None, 'Open', None, None, '')])
    boxes(iss, 2, ROWS + 1, len(ih))
    for r in range(2, ROWS + 2):
        iss.cell(row=r, column=7, value=f'=IF(F{r}="","",IF(F{r}>=4,"High",IF(F{r}=3,"Medium","Low")))')
        for c in (9, 11, 12):
            iss.cell(row=r, column=c).number_format = 'DD MMM YYYY'
    iss.conditional_formatting.add(f'G2:G{ROWS + 1}', CellIsRule(operator='equal', formula=['"High"'], fill=PatternFill('solid', fgColor='F8CBAD')))

    # dependencies
    dep = wb.create_sheet('Dependencies')
    dh = ['ID', 'Dependency', 'Provided by', 'Needed for', 'Needed by', 'Owner', 'State', 'Remarks']
    dw = [8, 50, 22, 30, 14, 18, 12, 30]
    dpre = [
        ('Named project manager, key users and Accounting Manager', 'Broker', 'Mobilisation', 'Week 1'),
        ('Insurer list and agreements, commission and referrer rates, product list', 'Broker', 'Configuration', 'End of discovery'),
        ('Chart of accounts and the accountant\'s mapping decisions', 'Broker', 'Ledger configuration', 'Start of configuration'),
        ('BIR Authority to Print or registered series; last numbers used', 'Broker', 'Document numbering', 'Configuration'),
        ('SMTP mailbox and credentials; payment gateway merchant account', 'Broker', 'Integrations', 'Integrations phase'),
        ('One recent statement export per bank account', 'Broker', 'Bank reconciliation formats', 'Integrations phase'),
        ('In-force policy, open receivable and trial balance extracts in the templates', 'Broker', 'Mock loads and cutover', 'Each mock load'),
        ('User list with roles and e-mail addresses', 'Broker', 'Training and go-live', 'Two weeks before training'),
        ('UAT testers and sign-off authority', 'Broker', 'UAT', 'UAT phase'),
        ('Production environment provisioned and backups tested', 'Hosting provider', 'Go-live', 'Before mock load 2'),
    ]
    table(dep, dh, dw, [(f'D-{i:03d}', d[0], d[1], d[2], d[3], '[owner]', 'On track', '') for i, d in enumerate(dpre, 1)])
    boxes(dep, 2, ROWS + 1, len(dh))

    refs = listsheet(wb, refs_data)
    for ws_, col in ((rk, 'E'), (iss, 'D')):
        validate(ws_, refs['Category'], col, 2, ROWS + 1)
    for ws_, col in ((rk, 'M'), (asn, 'H'), (iss, 'J')):
        validate(ws_, refs['Status'], col, 2, ROWS + 1)
    validate(dep, refs['Dependency state'], 'G', 2, ROWS + 1)
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True)
    iss.add_data_validation(dv); dv.add(f'F2:F{ROWS + 1}')
    dep.conditional_formatting.add(f'G2:G{ROWS + 1}', CellIsRule(operator='equal', formula=['"Late"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    dep.conditional_formatting.add(f'G2:G{ROWS + 1}', CellIsRule(operator='equal', formula=['"At risk"'], fill=PatternFill('solid', fgColor='FFE699')))

    # summary
    sm = wb.create_sheet('Summary', 1)
    sm['A1'] = 'RAID summary (open and in progress items)'
    sm['A1'].font = TITLE_FONT
    N = ROWS + 1
    open_ = lambda sheet, col: f'COUNTIFS({sheet}!$A$2:$A${N},"<>",{sheet}!${col}$2:${col}${N},"<>Closed")'
    rows = [
        ('Risks', f'=COUNTIFS(Risks!$J$2:$J${N},"High",Risks!$M$2:$M${N},"<>Closed")', f'=COUNTIFS(Risks!$J$2:$J${N},"Medium",Risks!$M$2:$M${N},"<>Closed")', f'=COUNTIFS(Risks!$J$2:$J${N},"Low",Risks!$M$2:$M${N},"<>Closed")', '=' + open_('Risks', 'M')),
        ('Issues', f'=COUNTIFS(Issues!$G$2:$G${N},"High",Issues!$J$2:$J${N},"<>Closed")', f'=COUNTIFS(Issues!$G$2:$G${N},"Medium",Issues!$J$2:$J${N},"<>Closed")', f'=COUNTIFS(Issues!$G$2:$G${N},"Low",Issues!$J$2:$J${N},"<>Closed")', '=' + open_('Issues', 'J')),
        ('Assumptions', '', '', '', '=' + open_('Assumptions', 'H')),
        ('Dependencies late or at risk', '', '', '', f'=COUNTIFS(Dependencies!$G$2:$G${N},"Late")+COUNTIFS(Dependencies!$G$2:$G${N},"At risk")'),
    ]
    table(sm, ['Type', 'High', 'Medium', 'Low', 'Open total'], [30, 12, 12, 12, 14], rows, start=3)
    sm.auto_filter.ref = None
    sm['A10'] = 'Risk matrix: open risks by probability (rows) and impact (columns)'
    sm['A10'].font = BOLD
    sm.cell(row=11, column=1, value='Probability \\ Impact').font = HEAD_FONT
    sm.cell(row=11, column=1).fill = HEAD_FILL
    for i in range(1, 6):
        c = sm.cell(row=11, column=i + 1, value=i); c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
    for p in range(5, 0, -1):
        r = 12 + (5 - p)
        c = sm.cell(row=r, column=1, value=p); c.font, c.border = BOLD, BOX
        for i in range(1, 6):
            cell = sm.cell(row=r, column=i + 1, value=f'=COUNTIFS(Risks!$G$2:$G${N},{p},Risks!$H$2:$H${N},{i},Risks!$M$2:$M${N},"<>Closed")')
            cell.border, cell.font = BOX, BODY_FONT
            cell.alignment = Alignment(horizontal='center')
            cell.fill = PatternFill('solid', fgColor='F8CBAD' if p * i >= 15 else ('FFE699' if p * i >= 8 else 'C6E0B4'))
    sm.freeze_panes = 'A4'
    wb.active = 1
    wb.save(os.path.join(OUT, 'BrokerVerse_RAID_Log_Template.xlsx'))


if __name__ == '__main__':
    fit_gap()
    raid()
    print('written to', OUT)
