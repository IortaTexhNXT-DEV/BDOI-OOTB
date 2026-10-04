"""Builds BrokerVerse_Implementation_Plan.xlsx from the plan model (delivery/plan_model.py):

    python3 build_plan_xlsx.py ../05_Delivery/BrokerVerse_Implementation_Plan.xlsx

Sheets: Read me; Summary (phases and milestones by size); one task-level Gantt per size (Small, Medium, Large,
Enterprise) with duration, owner party and role, predecessors, float and critical path; Milestones; Dependencies (every
link of the plan); RACI. The start and finish dates of each Gantt follow the kick-off date typed on the sheet.
"""
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'delivery'))
import plan_model as pm  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else '../05_Delivery/BrokerVerse_Implementation_Plan.xlsx'
NAVY = '0F4761'
HEAD = PatternFill('solid', fgColor='DCE9F7')
THIN = Side(style='thin', color='BFC7D1')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PARTY_FILL = {pm.IORTA: '2E75B6', pm.BROKER: '548235', pm.PARTNER: 'BF8F00'}
CRIT_FILL = 'C00000'
PHASE_BAND = PatternFill('solid', fgColor='EEF3F8')
VERSION = 'version 1.1, 04 October 2026'
KICKOFF = '2026-11-02'   # example kick-off date (a Monday); the project manager types the real one


def head(c, text):
    c.value = text
    c.font = Font(bold=True, color=NAVY, size=10)
    c.fill = HEAD
    c.border = BOX
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def links_text(t, i):
    out = []
    for p, kind, lag in t[7]:
        if pm.task(p)[6][i] is None:
            continue
        lag = pm.lag_of(lag, i)
        out.append(p + ('' if kind == 'FS' else kind) + (f'+{lag}d' if lag else ''))
    return ', '.join(out)


GATE_TEXT = {pm.GOLIVE: 'Go-live', pm.FEATURE: 'Feature only', pm.EXIT: 'Hypercare exit'}


def gantt(ws, size):
    i = pm.SIZES.index(size)
    info = pm.SIZE_INFO[size]
    s = pm.schedule(size)
    weeks = info['weeks']
    go_day = s['X6']['start']
    ws['A1'] = f'BrokerVerse OOTB implementation plan: {size} broker ({weeks} weeks, go-live at the start of week {info["golive_week"]})'
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    ws['A2'] = 'Kick-off date (a Monday):'
    ws['A2'].font = Font(bold=True, size=10)
    from datetime import date
    ws['D2'] = date.fromisoformat(KICKOFF)
    ws['D2'].number_format = 'DD MMM YYYY'
    ws['D2'].fill = PatternFill('solid', fgColor='FFF2CC')
    ws['D2'].border = BOX
    ws['F2'] = ('Type the kick-off date in the yellow cell: start and finish dates follow (working days, weekends skipped; '
                'add holidays as a third WORKDAY argument). Environments: ' + info['envs'] + '. ' + info['mocks'] + '.')
    ws['F2'].font = Font(italic=True, size=9, color='555F6D')
    ws['A3'] = ('Bars: blue iorta TechNXT, green broker, gold partner; red border and red ID = on the critical path to go-live. '
                'Float = working days a task can slip before the go-live date moves (empty: the task does not hold the go-live).')
    ws['A3'].font = Font(italic=True, size=9, color='555F6D')
    cols = ['ID', 'Phase', 'Task', 'Party', 'Owner role', 'Deliverable or evidence', 'Predecessors', 'Duration (working days)',
            'Start day', 'Finish day', 'Start date', 'Finish date', 'Start week', 'End week', 'Float to go-live (days)',
            'Critical', 'Holds back']
    r = 5
    for j, t in enumerate(cols, 1):
        head(ws.cell(r, j), t)
    base = len(cols)
    for w in range(1, weeks + 1):
        head(ws.cell(r, base + w), f'W{w}')
    ws.row_dimensions[r].height = 42
    order = [t for t in pm.TASKS if t[6][i] is not None]
    phase_rank = {p: k for k, (p, _) in enumerate(pm.PHASES)}
    order.sort(key=lambda t: (phase_rank[t[1]], s[t[0]]['start'], s[t[0]]['finish']))
    last_phase = None
    for t in order:
        tid = t[0]
        v = s[tid]
        if t[1] != last_phase:
            r += 1
            ws.cell(r, 1, pm.PHASE_NAME[t[1]]).font = Font(bold=True, color=NAVY, size=10)
            for j in range(1, base + weeks + 1):
                ws.cell(r, j).fill = PHASE_BAND
                ws.cell(r, j).border = BOX
            last_phase = t[1]
        r += 1
        a, b = pm.weeks(tid, v)
        crit = v['float_golive'] == 0
        vals = [tid, pm.PHASE_NAME[t[1]], t[2], t[3], t[4], t[5], links_text(t, i), v['dur'], v['start'], v['finish'],
                f'=WORKDAY($D$2,I{r})', f'=IF(H{r}=0,K{r},WORKDAY($D$2,J{r}-1))', a, b,
                v['float_golive'] if v['float_golive'] is not None else None, 'Yes' if crit else '', GATE_TEXT[t[8]]]
        for j, val in enumerate(vals, 1):
            c = ws.cell(r, j, val)
            c.border = BOX
            c.font = Font(size=9, bold=(j == 1), color=CRIT_FILL if (crit and j == 1) else None)
            c.alignment = Alignment(vertical='center', wrap_text=j in (3, 5, 6, 7), horizontal='center' if j >= 8 else 'left')
        for j in (11, 12):
            ws.cell(r, j).number_format = 'DD MMM YYYY'
        for w in range(1, weeks + 1):
            c = ws.cell(r, base + w)
            c.border = BOX
            if tid in pm.MILESTONES and w == a:
                c.value = 'GO' if tid == 'X6' else '◆'
                c.font = Font(bold=True, color='FFFFFF' if tid == 'X6' else CRIT_FILL, size=9)
                if tid == 'X6':
                    c.fill = PatternFill('solid', fgColor='00B050')
                c.alignment = Alignment(horizontal='center', vertical='center')
            elif tid not in pm.MILESTONES and a <= w <= b:
                c.fill = PatternFill('solid', fgColor=PARTY_FILL[t[3]])
                if crit:
                    red = Side(style='medium', color=CRIT_FILL)
                    c.border = Border(left=THIN, right=THIN, top=red, bottom=red)
            elif w == info['golive_week']:
                c.fill = PatternFill('solid', fgColor='E2EFDA')
        ws.row_dimensions[r].height = max(30, 12 * (max(len(t[2]) // 52, len(t[4]) // 22, len(t[5]) // 28) + 1) + 4)
    widths = [6, 16, 54, 13, 24, 30, 18, 10, 7, 7, 12, 12, 7, 7, 10, 8, 11]
    for j, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    for w in range(1, weeks + 1):
        ws.column_dimensions[get_column_letter(base + w)].width = 4
    ws.freeze_panes = 'D6'
    ws.auto_filter.ref = f'A5:{get_column_letter(base)}{r}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '5:5'
    return go_day


def summary(ws):
    ws['A1'] = 'Summary: phases and milestones by broker size'
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    ws['A2'] = 'Weeks from the kick-off (week 1). A phase runs from the first week of its first task to the last week of its last task.'
    ws['A2'].font = Font(italic=True, size=9, color='555F6D')
    r = 4
    head(ws.cell(r, 1), 'Phase')
    for k, size in enumerate(pm.SIZES):
        head(ws.cell(r, 2 + k), f'{size} ({pm.SIZE_INFO[size]["weeks"]} weeks)')
    scheds = {size: pm.schedule(size) for size in pm.SIZES}
    for code, name in pm.PHASES:
        r += 1
        ws.cell(r, 1, name).font = Font(bold=True, size=10)
        ws.cell(r, 1).border = BOX
        for k, size in enumerate(pm.SIZES):
            s = scheds[size]
            ids = [t[0] for t in pm.TASKS if t[1] == code and t[0] in s]
            if not ids:
                txt = 'Not applicable'
            else:
                a = min(pm.weeks(x, s[x])[0] for x in ids)
                b = max(pm.weeks(x, s[x])[1] for x in ids)
                txt = f'W{a}' if a == b else f'W{a} to W{b}'
            c = ws.cell(r, 2 + k, txt)
            c.border = BOX
            c.alignment = Alignment(horizontal='center')
    r += 2
    head(ws.cell(r, 1), 'Milestone')
    for k, size in enumerate(pm.SIZES):
        head(ws.cell(r, 2 + k), size)
    for tid in ['M1', 'D5', 'C5', 'T2', 'T5', 'X4', 'X6', 'H3']:
        r += 1
        ws.cell(r, 1, pm.task(tid)[2]).border = BOX
        for k, size in enumerate(pm.SIZES):
            v = scheds[size][tid]
            c = ws.cell(r, 2 + k, ('Start of ' if tid in pm.START_OF_WEEK else 'End of ') + f'W{pm.weeks(tid, v)[0]}')
            c.border = BOX
            c.alignment = Alignment(horizontal='center')
    r += 2
    head(ws.cell(r, 1), 'Size')
    for j, t in enumerate(['Named users', 'Environments', 'Mock loads', 'Hypercare'], 2):
        head(ws.cell(r, j), t)
    for size in pm.SIZES:
        r += 1
        info = pm.SIZE_INFO[size]
        for j, val in enumerate([size, info['users'], info['envs'], info['mocks'], info['hypercare'] + ', extended to the first month-end close'], 1):
            c = ws.cell(r, j, val)
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 45
    ws.column_dimensions['A'].width = 44
    for k in range(4):
        ws.column_dimensions[get_column_letter(2 + k)].width = 30
    ws.column_dimensions['F'].width = 30
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=1):
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def milestones(ws):
    ws['A1'] = 'Milestones and their evidence'
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    r = 3
    hdr = ['ID', 'Milestone', 'Evidence', 'Accountable'] + [f'{s} week' for s in pm.SIZES]
    for j, t in enumerate(hdr, 1):
        head(ws.cell(r, j), t)
    scheds = {size: pm.schedule(size) for size in pm.SIZES}
    for tid in ['M1', 'D5', 'C5', 'T2', 'T5', 'X4', 'X6', 'H3']:
        t = pm.task(tid)
        r += 1
        vals = [tid, t[2], t[5], t[4]] + [f'W{pm.weeks(tid, scheds[s][tid])[0]}' for s in pm.SIZES]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for j, wdt in enumerate([6, 50, 40, 30, 12, 12, 12, 12], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt


def dependencies(ws):
    ws['A1'] = 'Dependencies: every link of the plan'
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    ws['A2'] = ('FS = the task starts when the predecessor finishes; SS = the task starts when the predecessor has started (plus the lag). '
                'Lags in working days, by size.')
    ws['A2'].font = Font(italic=True, size=9, color='555F6D')
    r = 4
    hdr = ['Predecessor', 'Predecessor task', 'Successor', 'Successor task', 'Type', 'Lag S / M / L / E (days)', 'Holds back']
    for j, t in enumerate(hdr, 1):
        head(ws.cell(r, j), t)
    for t in pm.TASKS:
        for p, kind, lag in t[7]:
            r += 1
            lags = ' / '.join(str(pm.lag_of(lag, k)) for k in range(4))
            for j, v in enumerate([p, pm.task(p)[2], t[0], t[2], kind, lags, GATE_TEXT[t[8]]], 1):
                c = ws.cell(r, j, v)
                c.border = BOX
                c.alignment = Alignment(wrap_text=True, vertical='top')
    for j, wdt in enumerate([11, 50, 10, 50, 7, 16, 13], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:G{r}'


RACI_ROLES = ['iorta TechNXT PM', 'iorta TechNXT consultants', 'iorta TechNXT technical / DevOps', 'Broker sponsor',
              'Broker PM', 'Broker key users / process owners', 'Broker IT / System Administrator', 'Broker Accounting Manager',
              'Broker compliance officer / DPO', 'Partners (banks, SMS, CTPL provider, insurers, gateway)']
RACI = [
    ('Mobilisation', 'Project charter and baselined plan', 'R', 'C', 'C', 'A', 'R', 'I', 'I', 'I', 'I', 'I'),
    ('Mobilisation', 'Governance set-up (steering committee, status, RAID log)', 'R', 'I', 'I', 'A', 'R', 'I', 'I', 'I', 'I', 'I'),
    ('Mobilisation', 'Data request list and go-live kits issued and tracked', 'R', 'C', 'C', 'I', 'A', 'C', 'C', 'C', 'C', 'I'),
    ('Mobilisation', 'Hosting option, environment set and data location decision', 'C', 'I', 'C', 'A', 'R', 'I', 'R', 'I', 'C', 'I'),
    ('Environments and pipeline', 'Dev, SIT, UAT and Production provisioned; Pre-Prod when needed', 'A', 'I', 'R', 'I', 'I', 'I', 'C', 'I', 'I', 'I'),
    ('Environments and pipeline', 'GitHub Environments, protection rules, reviewers, variables and secrets', 'C', 'I', 'A', 'I', 'I', 'I', 'C', 'I', 'I', 'I'),
    ('Environments and pipeline', 'Encryption keys (DATA_ENCRYPTION_KEY, PII_ENCRYPTION_KEY): generation, custody, escrow copy, rotation', 'I', 'I', 'R', 'I', 'I', 'I', 'A', 'I', 'C', 'I'),
    ('Environments and pipeline', 'Backups, restore test, monitoring', 'A', 'I', 'R', 'I', 'I', 'I', 'C', 'I', 'I', 'I'),
    ('Discovery and fit-gap', 'Discovery workshops', 'A', 'R', 'I', 'I', 'R', 'R', 'C', 'R', 'C', 'I'),
    ('Discovery and fit-gap', 'Configuration decisions and fit-gap register sign-off', 'C', 'R', 'I', 'I', 'A', 'R', 'C', 'R', 'C', 'I'),
    ('Configuration', 'Configuration kit filled (masters, users, chart of accounts, insurers, products, taxes, numbering)', 'C', 'R', 'I', 'I', 'A', 'R', 'R', 'R', 'I', 'I'),
    ('Configuration', 'Configuration kit validated and loaded; errors workbook', 'A', 'R', 'C', 'I', 'I', 'C', 'R', 'I', 'I', 'I'),
    ('Configuration', 'Product Configurator: templates, rating factors, acceptance rules, document templates, market mapping', 'A', 'R', 'I', 'I', 'C', 'R', 'C', 'I', 'I', 'I'),
    ('Configuration', 'Tax codes, premium taxes, LGU rates, posting rules, account determination', 'A', 'R', 'I', 'I', 'C', 'C', 'I', 'R', 'I', 'I'),
    ('Configuration', 'Document numbering and invoice serial range (ATP or CAS)', 'A', 'R', 'I', 'I', 'C', 'C', 'C', 'R', 'I', 'I'),
    ('Configuration', 'Approvals, maker-checker, authority matrix, roles and view:pii grants', 'A', 'R', 'C', 'I', 'C', 'R', 'R', 'R', 'C', 'I'),
    ('Compliance set-up', 'AML settings, risk factors, monitoring rules', 'C', 'R', 'I', 'I', 'I', 'I', 'C', 'I', 'A', 'I'),
    ('Compliance set-up', 'Screening lists obtained, licensed and loaded', 'I', 'C', 'I', 'I', 'C', 'I', 'C', 'I', 'A', 'R'),
    ('Compliance set-up', 'IC licence register, fit and proper records, insurer certificates of authority', 'I', 'R', 'I', 'I', 'C', 'C', 'C', 'I', 'A', 'I'),
    ('Compliance set-up', 'Complaints and breach register settings; DPO and NPC registration confirmed', 'I', 'R', 'I', 'C', 'C', 'C', 'C', 'I', 'A', 'I'),
    ('Branding and brand pack', 'Theme, letterhead, branded documents and e-mails, e-signatures mapped, brand pack', 'A', 'R', 'C', 'C', 'C', 'C', 'R', 'I', 'I', 'I'),
    ('Branding and brand pack', 'Contract reference covering the use of client marks (client brand pack)', 'C', 'I', 'I', 'A', 'R', 'I', 'I', 'I', 'I', 'C'),
    ('Data migration', 'Extraction and cleansing of the old system data', 'I', 'C', 'I', 'I', 'A', 'R', 'R', 'R', 'I', 'I'),
    ('Data migration', 'Mapping to the migration kit', 'A', 'R', 'C', 'I', 'C', 'R', 'C', 'C', 'I', 'I'),
    ('Data migration', 'Mock loads and reconciliation', 'A', 'R', 'R', 'I', 'C', 'R', 'C', 'R', 'I', 'I'),
    ('Data migration', 'Sign-off of opening balances and open items', 'C', 'C', 'I', 'I', 'C', 'C', 'I', 'A', 'I', 'I'),
    ('Integrations', 'SMTP mailbox and credentials', 'I', 'I', 'R', 'I', 'C', 'I', 'A', 'I', 'I', 'I'),
    ('Integrations', 'Bank statement formats and bank payment file layouts', 'A', 'R', 'C', 'I', 'I', 'R', 'C', 'C', 'I', 'R'),
    ('Integrations', 'SMS or Viber, CTPL authentication and LTO, insurer API connectors, payment gateway', 'C', 'R', 'R', 'I', 'A', 'C', 'R', 'I', 'I', 'R'),
    ('Integrations', 'BIR EIS connector and enrolment', 'I', 'R', 'C', 'I', 'C', 'I', 'C', 'A', 'I', 'I'),
    ('Regulatory registrations', 'BIR ATP or CAS, EIS enrolment; AMLC registration; NPC registration; IC licences', 'I', 'C', 'I', 'A', 'R', 'I', 'I', 'R', 'R', 'I'),
    ('Training', 'Train-the-trainer', 'A', 'R', 'I', 'I', 'C', 'R', 'C', 'C', 'C', 'I'),
    ('Training', 'End-user training and assessment', 'C', 'C', 'I', 'I', 'A', 'R', 'C', 'C', 'C', 'I'),
    ('Training', 'Compliance officer and DPO training', 'C', 'R', 'I', 'I', 'C', 'I', 'C', 'I', 'A', 'I'),
    ('SIT and UAT', 'SIT execution and defect fixing', 'A', 'R', 'R', 'I', 'I', 'C', 'I', 'I', 'I', 'C'),
    ('SIT and UAT', 'UAT execution', 'C', 'C', 'I', 'I', 'A', 'R', 'R', 'R', 'R', 'I'),
    ('SIT and UAT', 'UAT sign-off', 'C', 'C', 'I', 'A', 'R', 'R', 'C', 'R', 'R', 'I'),
    ('Rehearsal and cutover', 'Release and frozen configuration to Production; comparison report', 'A', 'R', 'R', 'I', 'C', 'I', 'C', 'I', 'I', 'I'),
    ('Rehearsal and cutover', 'Pre-Prod from a production backup; masking approval', 'C', 'I', 'R', 'I', 'C', 'I', 'C', 'I', 'A', 'I'),
    ('Rehearsal and cutover', 'Cutover rehearsal, final loads and reconciliation', 'A', 'R', 'R', 'I', 'C', 'R', 'C', 'R', 'I', 'I'),
    ('Rehearsal and cutover', 'Go/no-go decision', 'C', 'C', 'C', 'A', 'R', 'C', 'C', 'C', 'C', 'I'),
    ('Go-live and hypercare', 'Go-live communication to users, insurers, banks and partners', 'C', 'I', 'I', 'A', 'R', 'C', 'C', 'I', 'I', 'I'),
    ('Go-live and hypercare', 'Hypercare issue handling', 'A', 'R', 'R', 'I', 'R', 'C', 'R', 'C', 'C', 'C'),
    ('Go-live and hypercare', 'First month-end close and first BIR and AMLC filings', 'C', 'R', 'C', 'I', 'I', 'R', 'I', 'A', 'R', 'I'),
    ('Go-live and hypercare', 'Handover to production support', 'A', 'R', 'R', 'I', 'C', 'I', 'R', 'I', 'I', 'I'),
    ('Change control', 'Change request impact assessment', 'A', 'R', 'R', 'I', 'C', 'C', 'C', 'C', 'C', 'I'),
    ('Change control', 'Change request approval', 'C', 'I', 'I', 'A', 'R', 'C', 'C', 'C', 'C', 'I'),
]


def raci(ws):
    ws['A1'] = 'RACI: iorta TechNXT, the broker and the partners'
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    ws['A2'] = 'R = Responsible, A = Accountable (one per activity), C = Consulted, I = Informed.'
    ws['A2'].font = Font(italic=True, size=9, color='555F6D')
    r = 4
    for j, t in enumerate(['Phase', 'Activity'] + RACI_ROLES):
        head(ws.cell(r, j + 1), t)
    ws.row_dimensions[r].height = 58
    fills = {'R': 'C6E0B4', 'A': 'F4B183', 'C': 'DDEBF7', 'I': 'F2F2F2'}
    for row in RACI:
        assert len(row) == 2 + len(RACI_ROLES) and row[2:].count('A') == 1, row
        r += 1
        for j, v in enumerate(row):
            c = ws.cell(r, j + 1, v)
            c.border = BOX
            if j >= 2:
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.fill = PatternFill('solid', fgColor=fills[v])
                c.font = Font(bold=v in 'RA', size=10)
            else:
                c.alignment = Alignment(vertical='center', wrap_text=True)
                c.font = Font(size=10, bold=j == 0)
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 56
    for j in range(len(RACI_ROLES)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 14
    ws.freeze_panes = 'C5'
    ws.auto_filter.ref = f'A4:{get_column_letter(2 + len(RACI_ROLES))}{r}'
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Read me'
    lines = [
        ('BrokerVerse OOTB Implementation Plan', Font(bold=True, size=14, color=NAVY)),
        (f'Companion workbook of the BrokerVerse Implementation Approach and Plan and of the Dependency Map and Critical Path, {VERSION}, iorta TechNXT.', None),
        ('', None),
        ('Sheets', Font(bold=True, color=NAVY)),
        ('Summary: the weeks of each phase and milestone for the four broker sizes.', None),
        ('Small (8 weeks), Medium (12 weeks), Large (20 weeks), Enterprise (26 weeks): the task-level plan of each size, with duration, party, owner role, deliverable, predecessors, dates from the kick-off date, float to go-live and the critical path, and the Gantt bars.', None),
        ('Milestones: the eight milestones of the plan, their evidence and their week by size.', None),
        ('Dependencies: every link between tasks, with its type and lag by size.', None),
        ('RACI: who is Responsible, Accountable, Consulted and Informed, for iorta TechNXT, the broker roles and the partners.', None),
        ('', None),
        ('Sizes (named users, as in the Service Catalogue and Rate Annex)', Font(bold=True, color=NAVY)),
    ] + [(f'{s}: {pm.SIZE_INFO[s]["users"]} users; {pm.SIZE_INFO[s]["envs"]}; hypercare {pm.SIZE_INFO[s]["hypercare"]}.', None) for s in pm.SIZES] + [
        ('', None),
        ('Notes', Font(bold=True, color=NAVY)),
        ('Durations are working days of elapsed time, not effort. They and the lead times of the regulatory registrations are planning assumptions; the steering committee re-baselines them at mobilisation.', None),
        ('A registration the broker already holds (ATP or CAS, AMLC registration, NPC registration, IC licences) takes no lead time: confirm it at mobilisation and close the task.', None),
        ('Tasks marked "Feature only" do not hold the go-live: the feature goes live later and the documented fallback is used meanwhile (Dependency Map and Critical Path).', None),
        ('Hypercare extends to cover the first month-end close when the close falls after the planned hypercare exit.', None),
        ('The plan assumes an out-of-the-box implementation: configuration only, no customisation.', None),
        ('Source: docs/package/tools/delivery/plan_model.py. Rebuild: python3 build_plan_xlsx.py ../05_Delivery/BrokerVerse_Implementation_Plan.xlsx', None),
    ]
    for k, (t, f) in enumerate(lines, 1):
        c = ws.cell(k, 1, t)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if f:
            c.font = f
    ws.column_dimensions['A'].width = 130
    summary(wb.create_sheet('Summary'))
    for size in pm.SIZES:
        gantt(wb.create_sheet(size), size)
    milestones(wb.create_sheet('Milestones'))
    dependencies(wb.create_sheet('Dependencies'))
    raci(wb.create_sheet('RACI'))
    wb.save(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
