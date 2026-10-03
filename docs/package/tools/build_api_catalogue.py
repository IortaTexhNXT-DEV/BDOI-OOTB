#!/usr/bin/env python3
"""Build BrokerVerse_API_and_Dependency_Catalogue.xlsx from the current code.

    python3 docs/package/tools/build_api_catalogue.py [output.xlsx]

Sources (collected by api_catalogue_sources.mjs, which needs node and the installed backend and brokerverse
packages):
  - the route registry (define() calls in backend/src/modules/*/router.js, loaded as `npm run export:api` does):
    one row per endpoint on the APIs sheet;
  - the side menu (brokerverse/src/components/SideBar/list.js), the routes (src/routes/MainRoute.js) and the
    call-site parser of brokerverse/scripts/check-api-calls.js: one row per menu item on the Screen dependencies
    sheet, with the endpoints its screens call and the backend modules behind them;
  - PREREQUISITES below: the data each backend module needs before its endpoints work (maintained here).

Sheets: APIs, Screen dependencies, Summary (per backend module, with the totals and the source notes).
"""
import json
import os
import re
import subprocess
import sys
from collections import OrderedDict, defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
DEFAULT_OUT = os.path.join(REPO, "docs", "package", "out", "BrokerVerse_API_and_Dependency_Catalogue.xlsx")

# Data that must exist before a module's endpoints are useful (per backend module folder).
PREREQUISITES = {
    "access-control": "Roles and permissions (seeded); users",
    "accounting": "Posted journals; chart of accounts; accounting periods",
    "bank-reconciliation": "Bank accounts linked to GL accounts; bank statement formats and transaction types; "
                           "imported statements; posted journals",
    "claims": "Issued policy; claim settings (claims.*); signatories for letters",
    "clients": "Document numbering (client); address masters",
    "collections": "Open receivables; jobs receivable-ageing and collection-reminders",
    "commission": "Issued policies with commission; Commission Rate Matrix; referrers (agents / sub-agents)",
    "commission-rates": "Insurance Company and Products masters",
    "credit-control": "Policies with receivables; credit.* settings; client credit limits",
    "dashboard": "Transactions in leads, quotations, policies, claims and receivables",
    "disbursements": "Payables (commission, refunds, premium due to insurers); bank accounts and checkbooks; "
                     "posting rules; open period",
    "documents": "Signatories and system settings (letterhead, logo); the record being printed",
    "endorsements": "Issued policy; posting rules for premium changes; open accounting period",
    "incentive": "Incentive Programs master; issued policies; eligible roles setting",
    "insurer-reconciliation": "Insurer statement formats; remittances to reconcile against",
    "journal-vouchers": "Chart of accounts; transaction codes; open accounting period",
    "leads": "Document numbering (lead); address masters",
    "notifications": "SMTP_URL in the server environment for e-mail delivery",
    "packages": "Package Bundles; Insurer Rate Tables in force for the date; Premium Taxes & LGU Rates; "
                "Payment Gateways for payment links",
    "payment-gateway": "Gateway enabled under Master > Finance > Payment Gateways; its credentials in the server "
                       "environment",
    "period-end": "Fiscal years and accounting periods; chart of accounts; tax codes; close checklist; posting rules",
    "placement": "Client or prospect; Insurance Company and Products masters; document numbering (broker slip, "
                 "placement)",
    "policies": "Accepted quotation or bound placement; insurer credit terms; posting rules and accounting.account.* "
                "settings for the booking journal; open accounting period",
    "posting-rules": "Chart of accounts (active GL codes); account roles under Account Determination",
    "premium-charges": "LGU tax rates and charge rules (Master > Finance > Premium Taxes & LGU Rates)",
    "privacy": "Clients and prospects; privacy.* settings (notice version, request due days, retention years)",
    "product-configurator": "Product templates; insurers for market mapping",
    "quotations": "Products and Insurance Company masters; coverages and vehicle masters (motor); motor tariff "
                  "(seeded); premium tax settings (premium.taxes_by_lob); commission rate matrix; a prospect or client",
    "receipts": "Open bills (receivables) from issued policies, endorsements, renewals or go-live opening items; "
                "bank accounts; posting rules; open accounting period",
    "reinsurance": "Reinsurance Treaty master and reinsurers; policies and claims to cede",
    "remittance": "Collected premiums due to insurers; Remittance Master (settlement limits); insurers",
    "renewals": "Policies near expiry; renewals.* settings; jobs renewal-queue, renewal-pipeline and renewal-notices "
                "(Master > Schedules)",
    "reports": "Report definitions (seeded); data in the source module",
    "schedules": "Scheduled jobs (seeded); general.timezone",
    "users": "Roles (seeded)",
}

HEADER_FILL = PatternFill("solid", fgColor="0F4761")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(wrap_text=True, vertical="center")
CELL_ALIGN = Alignment(wrap_text=True, vertical="top")
METHOD_ORDER = {"GET": 1, "POST": 2, "PUT": 3, "PATCH": 4, "DELETE": 5}


def collect():
    script = os.path.join(HERE, "api_catalogue_sources.mjs")
    run = subprocess.run(["node", script], cwd=REPO, capture_output=True, text=True, check=False)
    if run.returncode != 0:
        sys.stderr.write(run.stderr)
        raise SystemExit("api_catalogue_sources.mjs failed")
    data = json.loads(run.stdout)
    if data["skipped"]:
        raise SystemExit(f"backend modules failed to load: {data['skipped']}")
    if data["parseErrors"]:
        sys.stderr.write("front-end files not parsed:\n  " + "\n  ".join(data["parseErrors"]) + "\n")
    return data


def access(r):
    if not r["auth"]:
        return "Public (no sign-in)"
    parts = [f"role:{x}" for x in r["roles"]] + list(r["permissions"])
    return ", ".join(parts) if parts else "Any signed-in user"


def names_menu(label, menu):
    """True when a registry screen label names the menu item (or a screen under it)."""
    for part in label.replace(";", " / ").split(" / "):
        part = part.strip()
        if part == menu or part.startswith(menu + " >") or part.startswith(menu + " ("):
            return True
    return False


def write_sheet(ws, header, rows, widths, bold_rows=(), filter_rows=None):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGN
    for row in rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = CELL_ALIGN
    for i in bold_rows:
        for c in ws[i]:
            c.font = Font(bold=True)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    last_col = ws.cell(1, len(header)).column_letter
    ws.auto_filter.ref = f"A1:{last_col}{filter_rows or ws.max_row}"


def build(out):
    data = collect()
    registry = data["registry"]
    for r in registry:
        r["full"] = "/api" + r["path"]
        r["key"] = f"{r['method']} {r['path']}"
    index = {r["key"]: i for i, r in enumerate(registry)}
    by_key = {r["key"]: r for r in registry}

    # ---------- APIs ----------
    # summaries are quoted from the code; long dashes become " - " (writing rules)
    api_rows = [[r["folder"], r["module"], r["method"], r["full"], re.sub(r"\s*[\u2013\u2014]\s*", " - ", r["summary"]), access(r),
                 r["screen"] or None]
                for r in registry]

    # ---------- Screen dependencies ----------
    screen_rows = []
    modules_used = defaultdict(int)
    for s in data["screens"]:
        keys = set(s["endpoints"])
        keys |= {r["key"] for r in registry if r["screen"] and names_menu(r["screen"], s["menu"])}
        keys = sorted(keys, key=lambda k: index[k])
        folders = sorted({by_key[k]["folder"] for k in keys})
        for f in folders:
            modules_used[f] += 1
        prereq = []
        if s["masters"]:
            prereq.append("Masters read: " + ", ".join(s["masters"]))
        prereq += [f"{f}: {PREREQUISITES[f]}" for f in folders if f in PREREQUISITES]
        screen_rows.append([
            s["menu"], s["name"], s["route"], s["component"], s["file"],
            "\n".join(s["serviceFiles"]) or None, len(keys),
            "\n".join(f"{by_key[k]['method']} {by_key[k]['full']}" for k in keys) or None,
            "\n".join(folders) or None, "\n".join(prereq) or None,
        ])

    # ---------- Summary ----------
    per = OrderedDict()
    for r in registry:
        m = per.setdefault(r["folder"], {"labels": set(), "n": 0, "GET": 0, "POST": 0, "PUTPATCH": 0, "DELETE": 0,
                                         "public": 0})
        m["labels"].add(r["module"])
        m["n"] += 1
        meth = r["method"]
        m["PUTPATCH" if meth in ("PUT", "PATCH") else meth] += 1
        m["public"] += 0 if r["auth"] else 1
    summary_rows = []
    for folder in sorted(per):
        m = per[folder]
        summary_rows.append([folder, "; ".join(sorted(m["labels"])), m["n"], m["GET"], m["POST"], m["PUTPATCH"],
                             m["DELETE"], m["public"], modules_used.get(folder, 0), PREREQUISITES.get(folder)])
    total = ["TOTAL", None] + [sum(row[i] for row in summary_rows) for i in range(2, 8)] + [None, None]
    today = date.today().strftime("%d %B %Y")
    notes = [
        f"Source: route registry (backend/src/lib/registry.js) loaded by the module loader, {today}: every endpoint, "
        "its permission and its screen touchpoint.",
        "Health endpoints GET /api/health and GET /api/health/live are mounted in src/app.js outside the registry and "
        "are not listed.",
        "Screen dependencies: menu items from brokerverse/src/components/SideBar/list.js, routes from "
        "src/routes/MainRoute.js, then the import closure of the screen component (and the routes in the menu item's "
        "includes list). A service file counts only through the functions the screen code calls.",
        "Endpoints called = API call sites found in that closure (brokerverse/scripts/check-api-calls.js parser) plus "
        "the endpoints whose registry screen label names the menu item. Generic helpers (for example the masters and "
        "petty cash services) list every endpoint of the helper.",
        'Prerequisites are per backend module; "Masters read" lists the master types the screen reads through '
        "mastersService.",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "APIs"
    write_sheet(ws, ["Backend module (folder)", "Module label", "Method", "Path", "Summary", "Permission / access",
                     "Screen (touchpoint)"], api_rows,
                {"A": 22, "B": 22, "C": 9, "D": 52, "E": 70, "F": 34, "G": 52})
    ws = wb.create_sheet("Screen dependencies")
    write_sheet(ws, ["Menu path", "Screen", "Route", "Component", "Front-end file", "Service file(s)", "Endpoints",
                     "Endpoints called (method, path)", "Backend module(s)",
                     "Prerequisites (data that must exist first)"], screen_rows,
                {"A": 40, "B": 22, "C": 38, "D": 22, "E": 46, "F": 26, "G": 10, "H": 52, "I": 24, "J": 70})
    ws = wb.create_sheet("Summary")
    rows = summary_rows + [total]
    write_sheet(ws, ["Backend module (folder)", "Module label(s)", "Endpoints", "GET", "POST", "PUT/PATCH", "DELETE",
                     "Public", "Menu screens using it", "Prerequisites"], rows,
                {"A": 24, "B": 36, "C": 11, "D": 8, "E": 8, "F": 10, "G": 9, "H": 8, "I": 12, "J": 80},
                bold_rows=(len(rows) + 1,))
    ws.append([])
    for n in notes:
        ws.append([n])
    ws.auto_filter.ref = f"A1:J{len(rows) + 1}"
    wb.save(out)
    print(f"{out}: {len(api_rows)} endpoints, {len(screen_rows)} menu screens, {len(summary_rows)} backend modules")
    without = [s["menu"] for s in data["screens"] if not s["component"]]
    if without:
        print("menu items without a route component: " + "; ".join(without))


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUT)
