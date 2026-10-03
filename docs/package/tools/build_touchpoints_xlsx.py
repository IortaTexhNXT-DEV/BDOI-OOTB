#!/usr/bin/env python3
"""Build BrokerVerse_Communication_Touchpoints.xlsx from the communication templates document.

    python3 docs/package/tools/build_touchpoints_xlsx.py [--db DATABASE_URL] [output.xlsx]

Source: docs/package/source/communication-templates.md (written from the code: every call that queues an e-mail or
creates a notification). Sheets:
  - E-mails: one row per e-mail (E01, E02, ...) from its "## Enn <name>" section: the Item / Detail table, the
    stage and recipient group of its chapter, and the Subject / Body code blocks. When the document gives only a
    summary of the body, the seeded template (the *_template / *_body setting named under Configured in) is used.
  - Notifications: the tables under "# In-app notifications", one row per event, with the section as journey stage.
  - Documents: the first table under "# Printed documents".
  - Settings: every application setting the document names (and EXTRA_SETTINGS), with its value read from the database given by --db
    (or DATABASE_URL): a database built with the migrations and seeds gives the OOTB seeded values.
"""
import argparse
import json
import os
import re
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE = os.path.normpath(os.path.join(HERE, "..", "source", "communication-templates.md"))
DEFAULT_OUT = os.path.normpath(os.path.join(HERE, "..", "out", "BrokerVerse_Communication_Touchpoints.xlsx"))

HFILL = PatternFill("solid", fgColor="0F4761")
HFONT = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
BFONT = Font(name="Segoe UI", size=9)
WRAP = Alignment(wrap_text=True, vertical="top")

# Chapter of the document -> recipient group of its e-mails; a few e-mails go to more than one kind of party.
GROUP_OF_CHAPTER = {"E-mails to clients": "Client", "E-mails to insurers and agents": "Insurer", "E-mails to staff": "Internal"}
GROUP_OVERRIDE = {"E15": "Agent / insurer", "E17": "Any"}
# Settings of the printed documents and message numbering that the document refers to only as a group
# ("documents.* settings") or not by key, kept in the Settings sheet.
EXTRA_SETTINGS = {"documents.accent_color", "documents.default_logo_path", "documents.default_signatory", "limits.renewal_notice_days",
                  "motor.pricing_template_code", "numbering.product_template.prefix", "numbering.remittance_notice.prefix",
                  "product.template_statuses"}
EMAIL_COLUMNS = ["No.", "Recipient group", "Journey stage", "E-mail", "Trigger", "Screen", "Recipient", "Outbox template code",
                 "Configured in (setting)", "Where to edit", "Subject (seed)", "Body (seed)", "Merge fields", "Attachments",
                 "On / off", "Notes"]
EMAIL_WIDTHS = [6, 12, 16, 28, 40, 40, 34, 22, 30, 34, 40, 70, 40, 30, 30, 40]


def plain(text):
    """Cell text from a Markdown table cell: no code ticks or bold marks, <br> as a line break, no long dashes."""
    text = text.replace("<br>", "\n").replace("`", "").replace("**", "")
    return re.sub(r"\s*[–—]\s*", " - ", text).strip()


def table_rows(lines, start):
    """Rows of the pipe table that starts at or after line `start`: (header, rows, index after the table)."""
    i = start
    while i < len(lines) and not lines[i].startswith("|"):
        i += 1
    header = [c.strip() for c in lines[i].strip().strip("|").split("|")]
    i += 2  # header and |---| line
    rows = []
    while i < len(lines) and lines[i].startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if len(cells) > len(header):
            # a literal "|" inside a cell: keep it in the message (or the second-to-last) column
            at = header.index("Message") if "Message" in header else len(header) - 2
            extra = len(cells) - len(header)
            cells = cells[:at] + [" | ".join(cells[at:at + extra + 1])] + cells[at + extra + 1:]
        rows.append(cells)
        i += 1
    return header, rows, i


def code_block(lines, i):
    """Text of the ``` block that starts at or after line i, and the index after it."""
    while not lines[i].startswith("```"):
        i += 1
    j = i + 1
    out = []
    while not lines[j].startswith("```"):
        out.append(lines[j])
        j += 1
    return out, j + 1


def sections(lines, level):
    """(title, first line, end line) of every heading of the given level ("# " or "## ")."""
    marks = [(i, l[len(level):].strip()) for i, l in enumerate(lines) if l.startswith(level) and not l.startswith(level + "#")]
    out = []
    for n, (i, title) in enumerate(marks):
        end = marks[n + 1][0] if n + 1 < len(marks) else len(lines)
        if level == "## ":
            nxt = [k for k, l in enumerate(lines) if k > i and l.startswith("# ")]
            end = min(end, nxt[0]) if nxt else end
        out.append((title, i + 1, end))
    return out


def read_settings(db_url, keys):
    sql = "select key, value::text from app_settings where key = any(string_to_array('%s', ','))" % ",".join(sorted(keys))
    run = subprocess.run(["psql", db_url, "-AtF", "\x1f", "-c", sql], capture_output=True, text=True, check=False)
    if run.returncode:
        raise SystemExit("cannot read app_settings: " + run.stderr.strip())
    return dict(line.split("\x1f", 1) for line in run.stdout.splitlines() if line)


def all_setting_keys(db_url):
    run = subprocess.run(["psql", db_url, "-Atc", "select key from app_settings"], capture_output=True, text=True, check=False)
    if run.returncode:
        raise SystemExit("cannot read app_settings: " + run.stderr.strip())
    return set(run.stdout.split())


def write_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HFILL, HFONT, WRAP
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = BFONT, WRAP
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(ord('A') + len(header) - 1)}{ws.max_row}"
    return ws


def build(out, db_url):
    lines = open(SOURCE, encoding="utf-8").read().split("\n")
    chapters = sections(lines, "# ")
    db_keys = all_setting_keys(db_url)
    text = "\n".join(lines)
    named = {k for k in db_keys if re.search(r"(?<![\w.])" + re.escape(k) + r"(?![\w])", text)}
    named |= EXTRA_SETTINGS & db_keys
    values = read_settings(db_url, named)

    # ---------- E-mails ----------
    stage, group, overview_recipient = {}, {}, {}
    for title, a, b in chapters:
        if title in GROUP_OF_CHAPTER:
            header, rows, _ = table_rows(lines, a)
            for r in rows:
                no = r[0]
                stage[no] = r[header.index("Stage")]
                group[no] = GROUP_OVERRIDE.get(no, GROUP_OF_CHAPTER[title])
                overview_recipient[no] = r[header.index("Recipient")]
    emails = []
    for title, a, b in sections(lines, "## "):
        m = re.match(r"(E\d\d) (.+)$", title)
        if not m:
            continue
        no, name = m.groups()
        _, rows, after = table_rows(lines, a)
        item = {r[0]: r[1] for r in rows}
        subject = body = None
        k = after
        while k < b:
            line = lines[k].strip()
            if line == "Subject:":
                block, k = code_block(lines, k)
                subject = "\n".join(block)
                continue
            if line == "Body:":
                block, k = code_block(lines, k)
                body = "".join(block)
                continue
            if line.startswith("Body (summary):"):
                body = line.split(":", 1)[1].strip()
                seeded = [s.strip() for s in item.get("Configured in", "").split(";")[0].split(",")
                          if re.search(r"(_template|_body)$|^email\.template\.", s.strip())]
                if seeded and seeded[0] in values:
                    v = json.loads(values[seeded[0]])
                    body = v.get("html", body) if isinstance(v, dict) else v
            k += 1
        configured = item.get("Configured in", "")
        setting, _, where = configured.rpartition("; ") if re.search(r"; (Master|Accounts|Operations) >", configured) else (configured, "", "-")
        emails.append([no, group.get(no), stage.get(no), name, item.get("Trigger"), item.get("Screen"),
                       item.get("Recipient") or overview_recipient.get(no), item.get("Template code (outbox)"), setting, where,
                       subject, body, item.get("Merge fields"), item.get("Attachments"), item.get("On / off"), item.get("Notes")])
    emails.sort(key=lambda r: r[0])
    emails = [[plain(c) if isinstance(c, str) else c for c in r] for r in emails]

    # ---------- Notifications ----------
    notes = []
    chapter = next(c for c in chapters if c[0] == "In-app notifications")
    for title, a, b in sections(lines, "## "):
        if not (chapter[1] <= a < chapter[2]):
            continue
        header, rows, _ = table_rows(lines, a)
        for r in rows:
            notes.append([title] + [plain(c) for c in r])

    # ---------- Documents ----------
    chapter = next(c for c in chapters if c[0] == "Printed documents")
    header, rows, _ = table_rows(lines, chapter[1])
    docs = [[plain(c) for c in r] for r in rows]

    # ---------- Settings ----------
    settings = [[k, re.sub(r"\s*[–—]\s*", " - ", values[k])] for k in sorted(values)]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "E-mails", EMAIL_COLUMNS, emails, EMAIL_WIDTHS)
    write_sheet(wb, "Notifications", ["Journey stage", "Event", "Recipient", "Type", "Title", "Message", "On / off"], notes,
                [22, 40, 40, 10, 40, 60, 26])
    write_sheet(wb, "Documents", ["Document", "Where", "For", "Format", "Configured by"], docs, [40, 50, 14, 10, 40])
    write_sheet(wb, "Settings", ["Setting", "Value (OOTB seed)"], settings, [44, 120])
    wb.save(out)
    print(f"{out}: {len(emails)} e-mails, {len(notes)} notifications, {len(docs)} documents, {len(settings)} settings")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("out", nargs="?", default=DEFAULT_OUT)
    ap.add_argument("--db", default=os.environ.get("DATABASE_URL"), help="database with the seeded settings (default DATABASE_URL)")
    args = ap.parse_args()
    if not args.db:
        sys.exit("give --db or DATABASE_URL: a database built with the migrations and seeds")
    build(args.out, args.db)
