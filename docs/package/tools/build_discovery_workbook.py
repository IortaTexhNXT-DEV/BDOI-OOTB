"""Builds BrokerVerse_Discovery_and_Configuration_Workbook.xlsx: the workbook used at mobilisation and in the
discovery workshops to record every configuration decision of a broker, one sheet per area.

    python3 build_discovery_workbook.py ../out/BrokerVerse_Discovery_and_Configuration_Workbook.xlsx

Screens use the menu names of brokerverse/src/components/SideBar/list.js. OOTB defaults are the values seeded by
backend/src/db/seeds and the migrations (settings keys in the Notes column). The Reports sheet is read from the
Reports Book workbook so both stay aligned.
"""
import os, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '../out/BrokerVerse_Discovery_and_Configuration_Workbook.xlsx')
REPORTS_BOOK = os.path.join(HERE, '../out/BrokerVerse_Reports_Book.xlsx')

NAVY = '0B2A4A'
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
HEAD_FONT = Font(name='Segoe UI', bold=True, color='FFFFFF', size=10)
BODY = Font(name='Segoe UI', size=9)
BOLD = Font(name='Segoe UI', size=9, bold=True)
TITLE = Font(name='Segoe UI', size=14, bold=True, color=NAVY)
ANSWER_FILL = PatternFill('solid', fgColor='FFF9E5')
THIN = Side(style='thin', color='BFC7D1')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
STATUSES = ['Open', 'Answered', 'Confirmed', 'Configured', 'Not applicable', 'Gap (CR)']

HEADER = ['ID', 'Topic', 'Question or configuration item', 'BrokerVerse screen', 'OOTB default', 'Client answer',
          'Owner (broker)', 'Needed by', 'Due date', 'Status', 'Notes and setting key']
WIDTHS = [9, 18, 46, 36, 34, 36, 20, 16, 12, 13, 36]

# Phases of the implementation plan (Implementation Approach and Plan)
MOB, DISC, CONF, MIG, UAT = 'Mobilisation', 'Discovery', 'Configuration', 'Mock load 1', 'UAT'

M = 'Master > '
G = M + 'Generals > '
F = M + 'Finance > '
UM = G + 'User Management > '
IM = G + 'Insurance Management > '
CFG = M + 'Configuration'

# Each row: (topic, question, screen, default, owner, needed by, notes)
SHEETS = {}

SHEETS['Company and Branches'] = ('CO', [
    ('Legal identity', 'Registered company name exactly as on the SEC and BIR registration', G + 'Organization > Company', 'Seeded company ITX "iorta TechNXT Corp.": replace', 'Sponsor', MOB, 'Primary company is the letterhead on documents and BIR forms'),
    ('Legal identity', 'TIN and RDO code', G + 'Organization > Company', 'Blank', 'Accounting Manager', MOB, 'Fields TIN, RDO Code'),
    ('Legal identity', 'Registered address, ZIP code, city, province', G + 'Organization > Company', 'Blank', 'Accounting Manager', MOB, 'Fields Address Line 1 to 3, Pin Code, City, State'),
    ('Legal identity', 'Insurance Commission certificate of authority (licence) number and validity', G + 'Organization > Company', 'Blank', 'Sponsor', MOB, 'Field License Number'),
    ('Legal identity', 'Company e-mail, telephone, website', G + 'Organization > Company', 'Blank', 'System Administrator', CONF, ''),
    ('Legal identity', 'More than one legal entity on the system? Which one is the letterhead company', G + 'Organization > Company', 'One company', 'Sponsor', DISC, 'Field "Letterhead company - used on documents and reports"'),
    ('Branding', 'Logo for printed documents (PNG, at least 300 px wide)', G + 'Organization > Company', 'iorta TechNXT logo (fallback)', 'System Administrator', CONF, 'documents.default_logo_path is the fallback only'),
    ('Branding', 'Application name on the sign-in page, side bar and browser tab', M + 'System Settings', 'BrokerVerse', 'Sponsor', CONF, 'general.system_name'),
    ('Branding', 'Screen logo and colours', M + 'System Settings', 'Primary #0072d8, secondary #004ea8', 'System Administrator', CONF, 'branding.logo_url, branding.primary_color, branding.secondary_color'),
    ('Branding', 'Accent colour of printed documents and report PDFs', CFG + ' (Company & Branding)', '#1f4e79', 'System Administrator', CONF, 'documents.accent_color'),
    ('Branches', 'List of branches: code, name, address, e-mail, telephone', G + 'Organization > Branch', 'HO Head Office, Makati City', 'Operations head', MIG, 'Branch_Upload_Template.xlsx'),
    ('Branches', 'Departments per branch', G + 'Organization > Branch (Department_Upload_Template.xlsx)', 'Seeded departments', 'Operations head', MIG, 'Department upload through the API route written on the template'),
    ('People', 'Employee hierarchy (ranks) and designations', G + 'Employee Management > Hierarchy; Designation', 'Seeded hierarchy and designations', 'HR', MIG, 'Hierarchy_ and Designation_Upload_Template.xlsx'),
    ('People', 'Employees and sales persons with branch, department, reporting line', G + 'Employee Management > Employee', 'None', 'HR', MIG, 'Employee codes EMP-0001 (series employee)'),
    ('Calendar', 'Fiscal year start month', CFG + ' (Accounting & Tax)', '1 (January)', 'Accounting Manager', DISC, 'accounting.fiscal_year_start_month'),
    ('Calendar', 'Time zone and date format', M + 'System Settings', 'Asia/Manila; DD/MM/YYYY', 'System Administrator', CONF, 'general.timezone, general.date_format'),
    ('Calendar', 'Go-live date and first open accounting period', 'Accounts > Period End > Period Management', 'None', 'Sponsor', MOB, 'Drives cutover and opening balance date'),
    ('Calendar', 'Month-end close calendar, BIR filing dates and peak renewal months to avoid for releases', 'Not a system setting', 'Not applicable', 'Accounting Manager', DISC, 'Used for the maintenance window and the BCP'),
    ('Dashboards', 'Management targets: revenue, active policies, new business, claims ratio, retention, satisfaction', CFG + ' (Reports & Dashboards)', 'Revenue PHP 5,000,000.00; 250 policies; new business PHP 2,000,000.00; claims ratio 70; retention 90; satisfaction 95', 'Sponsor', CONF, 'dashboard.targets'),
    ('Dashboards', 'Sum insured above which a policy is flagged as high value', CFG + ' (Reports & Dashboards)', 'PHP 5,000,000.00', 'Sponsor', CONF, 'dashboard.high_sum_insured'),
    ('Hosting', 'Hosting option: AWS Singapore, Azure Singapore, Philippine partner, or own data centre', 'Not a system setting', 'AWS Singapore (implementation assumption)', 'IT head', MOB, 'Architecture document; data residency decision recorded by the DPO'),
    ('Hosting', 'Planned maintenance window', 'Not a system setting', 'Saturday 20:00 to Sunday 06:00 PHT', 'IT head', DISC, 'Production Support Approach and Standards'),
])

SHEETS['Users and Roles'] = ('US', [
    ('Users', 'List of named users: full name, user name, e-mail, branch, department, role', UM + 'User', 'Administrator BrokerVerse only', 'System Administrator', MIG, 'Users_Provisioning_Template.xlsx; user list due two weeks before training'),
    ('Roles', 'Map each job title to one of the seven delivered roles', UM + 'Role', 'System Administrator, Sales & Marketing, Processing Team, Operations, Claims, Accounting, Accounting Manager', 'Sponsor', DISC, 'One role per user is the norm'),
    ('Roles', 'Changes to the permissions of a delivered role', UM + 'Role Permissions', 'Delivered grants', 'Sponsor', DISC, 'Review with the User Access Matrix'),
    ('Administrators', 'Names of the System Administrator and a deputy', UM + 'User', 'None', 'IT head', MOB, 'At least two recommended'),
    ('Approvers', 'Names of a second Accounting Manager and second checkers for each maker-checker flow', UM + 'User', 'None', 'Accounting Manager', DISC, 'Maker and checker must be different users'),
    ('Security', 'Roles that must use two-step verification (authenticator app)', CFG + ' (Security & Access)', 'None', 'IT head', CONF, 'security.require_2fa_roles; recommended system-admin, accounting, accounting-manager'),
    ('Security', 'Password policy: length, character classes, history, maximum age', CFG + ' (Security & Access)', '8 characters; upper, lower, digit, symbol; last 5 refused; 90 days', 'IT head', CONF, 'security.password_*'),
    ('Security', 'Failed sign-ins before the account is locked', CFG + ' (Security & Access)', '5', 'IT head', CONF, 'limits.max_login_attempts'),
    ('Security', 'Sign-in attempts per IP and user name', CFG + ' (Security & Access)', '10 per 300 seconds', 'IT head', CONF, 'security.login_rate_limit'),
    ('Security', 'Idle sign-out', CFG + ' (Security & Access)', '30 minutes', 'IT head', CONF, 'limits.session_idle_minutes'),
    ('Security', 'Days without sign-in before an account is deactivated', CFG + ' (Security & Access)', '90 days', 'IT head', CONF, 'access.dormant_days'),
    ('Scope', 'Roles that see only their own book (records and clients they own)', CFG + ' (Security & Access)', 'None', 'Sponsor', DISC, 'security.scoped_roles'),
    ('Segregation of duties', 'Accept or change the SoD rules', UM + 'Segregation of Duties', 'Block: placement and payment; placement and payment approval; claims and payment. Warn: sales and collection; sales and claims', 'Sponsor', DISC, 'access.sod_enforced = true'),
    ('Access reviews', 'Frequency and reviewer of access recertification', UM + 'Access Reviews', 'Not scheduled', 'IT head', DISC, 'Twice a year recommended'),
    ('Delegations', 'Delegations of approval authority needed at go-live (leave, travel)', UM + 'Delegations', 'None', 'Accounting Manager', UAT, ''),
    ('Commission', 'Roles that earn commission on the policies they produce', CFG + ' (Commission & Incentives)', 'Sales & Marketing', 'Sales head', DISC, 'commission.eligible_roles'),
    ('Notifications', 'Roles told when a client approves a quotation', CFG + ' (Sales, Quotations & Placement)', 'Processing Team', 'Sales head', DISC, 'quotations.approval_notify_roles'),
    ('Notifications', 'Roles that approve renewals', CFG + ' (Policies, Endorsements & Renewals)', 'Processing Team', 'Operations head', DISC, 'renewals.approver_roles'),
])

SHEETS['Insurers'] = ('IN', [
    ('Insurer master', 'Insurers the broker places with: legal name, short name, code, TIN, address', IM + 'Insurance Company', 'Six Philippine non-life insurers seeded (MAPFRE, Malayan, Pioneer, FPG, Standard, Mercantile) with example contact e-mails', 'Processing Team lead', MIG, 'Insurance_Company_Upload_Template.xlsx; replace example e-mails'),
    ('Insurer master', 'Default commission rate per insurer', IM + 'Insurance Company', '14% to 17.5% on the seeded insurers', 'Accounting Manager', MIG, 'Rate matrix rows override it'),
    ('Credit terms', 'Premium payment warranty days per insurer', IM + 'Insurance Company', 'Blank', 'Accounting Manager', MIG, 'Field Premium Payment Warranty (days)'),
    ('Credit terms', 'Remittance terms: days after collection per insurer', IM + 'Insurance Company', 'Blank; fallback 30 days', 'Accounting Manager', MIG, 'Field Remittance Terms; remittance.default_due_days'),
    ('Billing', 'Default billing mode per insurer: broker billed or direct bill', IM + 'Insurance Company', 'Broker billed', 'Accounting Manager', DISC, 'Field Default Billing Mode; direct_bill.default_billing_mode'),
    ('Contacts', 'Underwriting, claims and accounting contacts per insurer (for slips, loss advice, remittance statements, debit notes)', IM + 'Insurance Company', 'Example addresses only', 'Processing Team lead', CONF, ''),
    ('Claims', 'Default recipient of the preliminary loss advice', CFG + ' (Claims)', 'claims@brokerverse.local (placeholder)', 'Claims head', CONF, 'claims.pla_default_recipient; replace before go-live'),
    ('Statements', 'Format of each insurer statement of account (columns, file type) and a sample file', F + 'Insurer Statement Formats', 'None', 'Accounting', CONF, 'Used by Accounts > Insurer Reconciliation'),
    ('Co-insurance', 'Lines usually co-insured; lead insurer practice; who confirms shares', 'Operations > Sales & Marketing > Placement Slips', 'Shares must total 100%; rounding to the lead', 'Processing Team lead', DISC, ''),
    ('Placement', 'Days an insurer offer stays valid', CFG + ' (Sales, Quotations & Placement)', '30 days', 'Processing Team lead', DISC, 'placement.offer_validity_days'),
    ('Signatories', 'Authorised signatories printed on documents', IM + 'Signatories', 'None in production', 'Sponsor', CONF, 'Signatories_Upload_Template.xlsx'),
    ('Reinsurance', 'Does the broker arrange reinsurance? Reinsurers and treaties', F + 'Reinsurance Treaty', 'Minimum security rating A-; treaty approval required; retention target 65%', 'Processing Team lead', DISC, 'reinsurance.*; skip if not used'),
    ('Direct bill', 'EWT the insurer withholds on direct-bill commission', F + 'Taxation', 'WC139 10%', 'Accounting Manager', CONF, 'direct_bill.insurer_ewt_code'),
    ('Direct bill', 'Commission debit note: title, remarks, due days, e-mail text', CFG + ' (Billing, Collections & Credit)', 'Commission Debit Note; 30 days', 'Accounting Manager', CONF, 'direct_bill.debit_note_*'),
    ('Reconciliation', 'Tolerance for remittance reconciliation differences', CFG + ' (Remittance & Reconciliation)', 'PHP 0.50', 'Accounting Manager', CONF, 'remittance.reconciliation_tolerance'),
])

SHEETS['Products and LOB'] = ('PR', [
    ('Lines', 'Lines of business written; the three with the most premium', IM + 'Line of Business', 'Philippine non-life lines delivered', 'Sales head', DISC, 'Number of lines drives the implementation fee'),
    ('Products', 'Products per line, mapped to the delivered products', IM + 'Product', 'Motor, CTPL, Fire, Marine Cargo, PA, CGL, Group EB, Surety Bond; travel, householder, micro-insurance, group PA, CAR, EAR, machinery breakdown, marine hull, money and securities', 'Processing Team lead', DISC, 'Product_Upload_Template.xlsx for additions'),
    ('Products', 'Policy types per product', IM + 'Product', 'For example Comprehensive, TPL, Own Damage / Theft; Residential, Commercial; Import, Domestic', 'Processing Team lead', CONF, 'Policy_Type_Upload_Template.xlsx'),
    ('Products', 'Package or non-package classification per product', IM + 'Product', 'Set per delivered product', 'Processing Team lead', DISC, 'Drives the placement journey'),
    ('Products', 'Tax regime per product: VAT or premium tax', IM + 'Product', 'As delivered per product', 'Accounting Manager', DISC, 'Confirm with each insurer'),
    ('Covers', 'Covers and limits offered; excess bodily injury and property damage limits on motor', IM + 'Cover', 'BI and PD limits PHP 100,000.00 to 500,000.00', 'Processing Team lead', CONF, 'quote.bodily_injury_limits, quote.property_damage_limits'),
    ('Journey', 'Placement journey per line: Request for Quotation, Quotation Slip, Placement Slip, Record Issued Policy', CFG + ' (Sales, Quotations & Placement)', 'Motor: Quotation Slip required. Fire, IAR, Marine, Casualty, Engineering: Placement Slip required. Others optional', 'Processing Team lead', DISC, 'placement.journey'),
    ('Motor', 'Confirm the motor tariff, CTPL 1-year and 3-year amounts per vehicle class', 'Product Configurator > Product Templates (MOT-003-2025)', '2025 tariff, confirmed 29 September 2026', 'Processing Team lead', CONF, 'motor.pricing_template_code'),
    ('Motor', 'Auto Passenger PA rates and limits', 'Product Configurator > Product Templates', 'Per seat from the motor template', 'Processing Team lead', CONF, ''),
    ('Motor', 'Vehicle makes, models and variants to add', IM + 'Vehicle', 'Toyota, Mitsubishi, Honda, Ford, Nissan, Hyundai, Suzuki, Isuzu with models', 'Processing Team lead', MIG, 'Vehicle_* templates'),
    ('KYC', 'Identifiers required before issue, per line', CFG + ' (Policies, Endorsements & Renewals)', 'Motor: ID type, number, image, chassis, motor and plate number', 'Compliance officer', DISC, 'policy.kyc_required_fields'),
    ('KYC', 'Accepted government ID documents', CFG + ' (Policies, Endorsements & Renewals)', "PhilSys ID, UMID, Passport, Driver's License, PRC ID, SSS ID, GSIS and others", 'Compliance officer', DISC, 'policy.kyc_id_types'),
    ('Quick Quote', 'Package bundles and insurer rate tables for Quick Quote and Compare Insurers', F + 'Package Bundles; Insurer Rate Tables', 'Delivered bundles', 'Sales head', CONF, 'packages.*'),
    ('Rules', 'Acceptance rules and product approval workflows', 'Product Configurator > Acceptance Rules; Approval Workflows', 'Delivered rules', 'Processing Team lead', CONF, ''),
    ('Documents', 'Policy schedule, quotation and endorsement layouts; wording to print', 'Product Configurator > Document Manager', 'Delivered layouts', 'Processing Team lead', CONF, 'Layout changes outside OOTB are change requests'),
    ('Quotation', 'Quotation validity; maker-checker on quotations; approval link validity', CFG + ' (Sales, Quotations & Placement)', '30 days; on; 168 hours', 'Sales head', DISC, 'limits.quote_validity_days, workflow.quote_maker_checker, quotations.approval_link_ttl_hours'),
    ('Policy', 'Default policy term', CFG + ' (Policies, Endorsements & Renewals)', '12 months', 'Processing Team lead', DISC, 'policies.default_term_months'),
    ('Renewals', 'Renewal timetable: pipeline, notices, grace, reinstatement', CFG + ' (Policies, Endorsements & Renewals)', 'Pipeline 90 days; notices 60, 30, 15 days; grace 30 days; reinstatement 90 days', 'Operations head', DISC, 'renewals.pipeline_days, limits.renewal_notice_days, renewals.grace_period_days, renewals.reinstatement_days'),
    ('Renewals', 'Renewal pricing: claims loading, loyalty discount, maker-checker', CFG + ' (Policies, Endorsements & Renewals)', 'Loading 10% (cap 30%); loyalty 2% (cap 10%); maker-checker on', 'Operations head', DISC, 'renewals.claims_loading_*, renewals.loyalty_discount_*'),
    ('Claims', 'Claims controls: refuse claims on unpaid premium, check loss date, settlement maker-checker, SLA days, ageing', CFG + ' (Claims)', 'On; on; on; 20 days; 7, 15, 30 days', 'Claims head', DISC, 'claims.block_unpaid_premium, claims.sla_days, claims.aging_thresholds'),
    ('Packages', 'Pro-rata pricing of section endorsements on package policies', CFG + ' (Sales, Quotations & Placement)', 'On', 'Processing Team lead', DISC, 'packages.endorsement_prorata'),
    ('Prospects', 'Age limits for a prospect or client date of birth', CFG + ' (Sales, Quotations & Placement)', '18 to 100 years', 'Sales head', CONF, 'leads.min_age_years, leads.max_age_years'),
])

SHEETS['Commission Rates'] = ('CM', [
    ('Insurer commission', 'Commission rate by insurer, product or line, new or renewal, with effective dates', F + 'Commission Rate Matrix', 'Fallback 15%', 'Accounting Manager', CONF, 'commission.default_rate; no upload template, entered on screen'),
    ('Referrers', 'Referrers and sub-agents: name, payee type, TIN, bank account', 'Commission > Agents/Referrer Accounts', 'None in production', 'Accounting', MIG, 'A referrer needs a bank account before payout (commission.require_bank_account)'),
    ('Sharing', 'Sub-agent (comsub) share by level', CFG + ' (Commission & Incentives)', 'Level 1 8%; level 2 5%', 'Sales head', DISC, 'commission.comsub_rate_by_level'),
    ('Withholding', 'Withholding tax code per payee type', CFG + ' (Commission & Incentives)', 'Agent WI515 5%; Sub-agent WI515 5%; External WC515 10%', 'Accounting Manager', DISC, 'commission.wht_code_by_type; confirm with the tax adviser'),
    ('Payout rule', 'Commission payable only when the premium is fully paid', CFG + ' (Commission & Incentives)', 'On; eligible automatically on full payment', 'Accounting Manager', DISC, 'commission.require_full_payment, commission.auto_eligible_on_full_payment'),
    ('Commission master', 'Commission master records (codes) used by sales', G + 'Commission', 'COM- series', 'Sales head', CONF, ''),
    ('VAT', 'Is the broker VAT-registered; is commission quoted VAT-inclusive', CFG + ' (Billing, Collections & Credit)', 'VAT-registered: yes; VAT-inclusive: no', 'Accounting Manager', DISC, 'direct_bill.broker_vat_registered, direct_bill.commission_vat_inclusive'),
    ('Incentives', 'Incentive programmes for sales staff: basis, period, targets, payout', F + 'Incentive Programs', 'None in production', 'Sales head', CONF, 'Skip if not used'),
    ('Open balances', 'Commission due to referrers earned before go-live', 'Accounts > Disbursement', 'Part of the commission payable opening balance', 'Accounting', MIG, 'Paid with payment vouchers'),
])

GL_ACCOUNTS = [
    ('Cash on hand', '1101001'), ('Cash in bank (default)', '1102001'), ('Petty cash fund', '1103001'),
    ('Premium receivable', '1202001'), ('Commission receivable', '1203001'), ('Agent receivable', '1204001'),
    ('Insurer refund receivable', '1203002'), ('Employee advances', '1205001'), ('Input VAT', '1301001'),
    ('Creditable withholding tax', '1302001'), ('Output VAT', '2204003'), ('Due to insurers', '2201001'),
    ('Commission payable', '2203001'), ('Withholding tax payable', '2204001'), ('Client refund payable', '2205001'),
    ('Supplier payable', '2206001'), ('Commission income', '3201001'), ('Commission expense', '4401010'), ('Write-off', '4401009'),
]
SHEETS['Chart of Accounts'] = ('GL', [
    ('Chart', 'Adopt the delivered Philippine broker chart of accounts, or map the broker\'s chart to it', F + 'Account Category; Main Account; Sub Account', 'Delivered chart of about 110 accounts', 'Accounting Manager', DISC, 'Chart_of_Accounts_Upload_Template.xlsx'),
    ('Chart', 'Premium trust account and other bank GL accounts needed', F + 'Main Account; Sub Account', 'Cash in Bank: Operating, E-wallet Clearing, Premium Trust, Payroll', 'Accounting Manager', DISC, ''),
] + [('Account determination', f'GL account for {name.lower()}', F + 'Account Determination', code, 'Accounting Manager', CONF, 'accounting.account.*') for name, code in GL_ACCOUNTS] + [
    ('Account determination', 'Payable account per payee type', F + 'Account Determination', 'Insurer 2201001; Client 2205001; Agent/Referrer 2203001; Supplier 2206001', 'Accounting Manager', CONF, 'accounting.payable_account_by_payee'),
    ('Account determination', 'Cash account per payment mode', F + 'Account Determination', 'Cash 1101001; cheque, bank transfer, card 1102001; GCash and online 1102002', 'Accounting Manager', CONF, 'accounting.cash_account_by_payment_mode'),
    ('Posting rules', 'Review the posting rule of each business event; agree changes', F + 'Posting Rules; Accounting Flow', 'Delivered rules', 'Accounting Manager', CONF, 'Changes approved by a second user in Configuration Approvals'),
    ('Controls', 'Maker-checker on finance documents; journal approval; period close approval', CFG + ' (Accounting & Tax)', 'All on', 'Accounting Manager', DISC, 'finance.maker_checker_enabled, journal.require_approval, accounting.period_close_requires_approval'),
    ('Transaction codes', 'Transaction codes with GL accounts and user limits', F + 'Transaction Code', 'Delivered codes', 'Accounting Manager', CONF, 'Transaction_Code_Upload_Template.xlsx'),
    ('Write-off', 'Write-off reasons, GL account and maximum amount', F + 'Account Determination', 'Delivered reasons', 'Accounting Manager', CONF, 'Write_off_Reason_Upload_Template.xlsx'),
    ('Close', 'Month-end close checklist items', F + 'Close Checklist', 'Delivered checklist incl. sub-ledger tie-out and bank reconciliation', 'Accounting Manager', CONF, ''),
    ('Close', 'Recurring journals (rent, depreciation, accruals)', 'Accounts > Period End > Recurring Journals', 'None', 'Accounting', UAT, ''),
    ('Opening balances', 'Trial balance at the day before go-live and its date', 'Accounts > Period End > Period Management (Import opening balances)', 'None', 'Accounting Manager', MIG, 'Opening_Balances_Upload_Template.xlsx'),
    ('Currency', 'Currencies used and source of exchange rates', F + 'Currency; Exchange Rate', 'PHP base; USD, EUR, SGD, JPY', 'Accounting Manager', CONF, 'Exchange_Rate_Upload_Template.xlsx'),
    ('Receivables', 'Ageing buckets; receivable due days', CFG + ' (Billing, Collections & Credit)', '30, 60, 90, 120 days; due in 30 days', 'Accounting Manager', DISC, 'limits.receivable_ageing_buckets, receivables.due_days'),
    ('Collections', 'Credit days, overdue levels, reminder timing', CFG + ' (Billing, Collections & Credit)', '30 days; 30 and 60 days; 7 days before, every 7 days', 'Accounting Manager', DISC, 'collections.*'),
    ('Petty cash', 'Petty cash funds: custodian, amount, replenishment limit', F + 'Petty Cash', 'None', 'Accounting Manager', CONF, 'PCF- codes'),
    ('External ledger', 'Will a corporate ledger be kept in another package? Export needs', 'Reports > Financial Reports > Trial Balance; Journal', 'Ledger kept in BrokerVerse', 'Accounting Manager', DISC, 'Excel or CSV export; a direct interface is a change request'),
])

SHEETS['Banks and Gateways'] = ('BK', [
    ('Bank accounts', "Broker's bank accounts: bank, branch, account number, purpose (operating, premium trust, payroll), GL cash account", F + 'Bank', 'Banks seeded: BDO, BPI, Metrobank, Land Bank, Security Bank, China Bank; no accounts', 'Accounting Manager', MIG, 'Bank_Account_Upload_Template.xlsx'),
    ('Statements', 'Statement export format per account and one sample statement', F + 'Bank Statement Formats', 'BDO-SAMPLE, BPI-SAMPLE, MBT-SAMPLE, GENERIC', 'Accounting', CONF, ''),
    ('Statements', 'Reconcile-from date per account', F + 'Bank', 'Go-live date', 'Accounting Manager', MIG, 'Older items stay on the old reconciliation'),
    ('Statements', 'Bank transaction types: charges, interest, final tax, others', F + 'Bank Transaction Types', 'Delivered types', 'Accounting', CONF, ''),
    ('Payments in', 'How clients may pay', CFG + ' (Billing, Collections & Credit)', 'Bank transfer, cheque, online, cash', 'Accounting Manager', DISC, 'policy.payment_capture_modes'),
    ('Payments in', 'Default payment mode on receipts; on disbursements', CFG + ' (Billing, Collections & Credit)', 'Bank transfer; cheque', 'Accounting', CONF, 'receipts.default_payment_mode, disbursements.default_payment_mode'),
    ('Gateway', 'Online payment gateway to use: PayMongo, Dragonpay or none; merchant account owner', F + 'Payment Gateways', 'PayMongo and Dragonpay off (sandbox mode); Sandbox gateway on (switch off in production)', 'Accounting Manager', DISC, 'Keys only in the secret store: PAYMONGO_*, DRAGONPAY_*'),
    ('Gateway', 'Methods offered and fee handling (absorbed or charged)', F + 'Payment Gateways', 'Absorb; PayMongo 2.5%; Dragonpay PHP 20.00', 'Accounting Manager', CONF, 'Confirm with the merchant agreement'),
    ('Gateway', 'Payment link validity; automatic official receipt on confirmed payment', CFG + ' (Billing, Collections & Credit)', '72 hours; on', 'Accounting Manager', CONF, 'payments.link_validity_hours, payments.auto_receipt'),
    ('Payments out', 'Transfer methods to insurers (InstaPay, PESONet, RTGS, cheque) and who executes them in the bank portal', 'Accounts > Remittance > Electronic Transfer', 'Recorded and approved in BrokerVerse; executed in the bank portal', 'Accounting Manager', DISC, 'Bank payment files are a change request'),
    ('Payments out', 'Cheque series in use per account', 'Accounts > Disbursement', 'None', 'Accounting', MIG, ''),
    ('Remittance', 'Bank account used for remittance reconciliation', CFG + ' (Remittance & Reconciliation)', 'Blank', 'Accounting', CONF, 'remittance.reconciliation_bank_account'),
])

SERIES = [
    ('lead', 'Lead', 'LD'), ('quote', 'Quotation', 'QT'), ('client', 'Client Code', 'CL'), ('policy', 'Policy (when the insurer number is not given)', 'POL'),
    ('endorsement', 'Endorsement', 'END'), ('renewal', 'Renewal', 'RN'), ('renewal_quote', 'Renewal Quotation', 'RQ'), ('renewal_batch', 'Renewal Batch', 'RB'),
    ('campaign', 'Win-back Campaign', 'WB'), ('claim', 'Claim', 'CLM'), ('broker_slip', 'Broker Slip', 'BS'), ('placement', 'Placement Slip', 'PS'),
    ('insurer_offer', 'Insurer Offer', 'OFR'), ('package_quote', 'Package Quotation', 'PQ'), ('package_policy', 'Package Policy', 'PKG'),
    ('receipt', 'Official Receipt', 'OR'), ('receipt_txn', 'Receipt Transaction', 'RT'), ('acknowledgement_receipt', 'Acknowledgement Receipt', 'AR'),
    ('invoice', 'Invoice / Bill', 'INV'), ('payment_link', 'Payment Link', 'PL'), ('voucher', 'Payment Voucher', 'PV'), ('disbursement_txn', 'Disbursement Transaction', 'DT'),
    ('invoice_list', 'Payable (Invoice List)', 'IL'), ('petty_cash', 'Petty Cash Transaction', 'PC'), ('petty_cash_request', 'Petty Cash Request', 'PCR'),
    ('petty_cash_receipt', 'Petty Cash Receipt', 'PCRC'), ('commission_debit_note', 'Commission Debit Note', 'DN'), ('dn_collection', 'Debit Note Collection', 'DNC'),
    ('journal', 'Journal Voucher', 'JV'), ('period_close', 'Month-end Close', 'MEC'), ('year_end_close', 'Year-end Close', 'YEC'), ('recurring_journal', 'Recurring Journal', 'RJV'),
    ('bir_2307', 'BIR Form 2307', 'CWT'), ('bank_statement', 'Bank Statement', 'BST'), ('bank_reconciliation', 'Bank Reconciliation', 'BRC'),
    ('insurer_statement', 'Insurer Statement Reconciliation', 'ISR'), ('remittance', 'Remittance', 'REM'), ('remittance_bill', 'Direct / Agency Bill', 'BIL'),
    ('remittance_batch', 'Remittance Batch', 'BLK'), ('settlement', 'Settlement', 'SET'), ('adjustment', 'Remittance Adjustment', 'ADJ'), ('transfer', 'Electronic Transfer', 'TRF'),
    ('statement', 'Remittance Statement', 'STMT'), ('remittance_exception', 'Remittance Exception', 'EXC'), ('remittance_notice', 'Remittance Notification', 'NTF'),
    ('remittance_schedule', 'Remittance Schedule', 'SCH'), ('remittance_report', 'Remittance Report', 'RPT'), ('bank_txn', 'Bank Transaction', 'BNK'),
    ('reinsurer', 'Reinsurer', 'RE'), ('treaty', 'Treaty', 'TRT'), ('cession', 'Cession', 'CES'), ('ri_recovery', 'RI Recovery Claim', 'RCL'), ('bordereau', 'Bordereau', 'BDX'),
    ('ri_reconciliation', 'RI Reconciliation', 'REC'), ('incentive_program', 'Incentive Program', 'INC'), ('incentive_calc', 'Incentive Calculation', 'CALC'),
    ('product_template', 'Product Template', 'TPL'), ('data_subject_request', 'Data Subject Request', 'DSR'),
]
MASTER_SERIES = [('commission_master', 'Commission Code', 'COM', '{PREFIX}-{SEQ}, width 5, never reset'),
                 ('employee', 'Employee Code', 'EMP', '{PREFIX}-{SEQ}, width 4, never reset'),
                 ('petty_cash_fund', 'Petty Cash Code', 'PCF', '{PREFIX}-{SEQ}, width 3, never reset')]
BIR_SERIES = {'receipt', 'invoice', 'commission_debit_note', 'bir_2307', 'acknowledgement_receipt'}
num_rows = [('Rules', 'Continue the old system\'s numbers at go-live? Next number per series', M + 'Document Numbering', 'Counters start at 1; can be set to continue', 'Accounting Manager', CONF, 'Counters only move forward')]
for code, name, prefix in SERIES:
    bir = code in BIR_SERIES
    num_rows.append(('BIR document' if bir else 'Series', f'{name}: prefix, pattern and reset rule' + ('; registered series (ATP or system-generated equivalent)' if bir else ''),
                     M + 'Document Numbering', f'{prefix}, {{PREFIX}}-{{YYYY}}-{{SEQ}}, width 5, yearly reset',
                     'Accounting Manager' if bir else 'System Administrator', CONF, f'Series code {code}' + ('; confirm with the tax adviser' if bir else '')))
for code, name, prefix, pat in MASTER_SERIES:
    num_rows.append(('Master codes', f'{name}: prefix and pattern', M + 'Document Numbering', f'{prefix}, {pat}', 'System Administrator', CONF, f'Series code {code}'))
SHEETS['Numbering Series'] = ('NU', num_rows)

AUTH = [
    ('quotation_discount', 'Quotation discount (percent)', 'Sales 10; Operations 10; Processing 15; Accounting Manager 25'),
    ('policy_issue', 'Policy issue (sum insured, PHP)', 'Sales 1,000,000.00; Operations 1,000,000.00; Processing 5,000,000.00; Accounting Manager no limit'),
    ('return_premium', 'Return premium (PHP)', 'Operations 100,000.00; Processing 250,000.00; Accounting 250,000.00; Accounting Manager no limit'),
    ('claim_settlement', 'Claim settlement (PHP)', 'Claims 1,000,000.00; Accounting Manager no limit'),
    ('payment_voucher', 'Payment voucher (PHP)', 'Accounting 2,000,000.00; Accounting Manager no limit'),
    ('journal_voucher', 'Journal voucher (PHP)', 'Accounting 1,000,000.00; Accounting Manager no limit'),
    ('write_off', 'Write-off (PHP)', 'Accounting 1,000.00; Accounting Manager 50,000.00'),
    ('commission_payout', 'Commission payout (PHP)', 'Accounting 500,000.00; Accounting Manager no limit'),
    ('petty_cash', 'Petty cash (PHP)', 'Accounting 10,000.00; Accounting Manager 50,000.00'),
    ('remittance', 'Remittance (PHP)', 'Accounting 1,000,000.00; Accounting Manager no limit'),
    ('remittance_settlement', 'Remittance settlement (PHP)', 'Accounting 1,000,000.00; Accounting Manager no limit'),
]
SHEETS['Approvals and Limits'] = ('AP', [
    ('Authority matrix', f'Approval limit: {label}, per role or named user', UM + 'Authority Matrix', default, 'Sponsor', CONF, f'Transaction type {t}; replace with the board-approved signing authority')
    for t, label, default in AUTH] + [
    ('Authority matrix', 'Who approves a new or changed limit (second administrator)', UM + 'Authority Matrix', 'Another administrator must approve', 'Sponsor', CONF, ''),
    ('Maker-checker', 'Quotation approval before sending to the client', CFG + ' (Sales, Quotations & Placement)', 'On', 'Sales head', DISC, 'workflow.quote_maker_checker'),
    ('Maker-checker', 'Renewal approval', CFG + ' (Policies, Endorsements & Renewals)', 'On', 'Operations head', DISC, 'renewals.maker_checker'),
    ('Maker-checker', 'Claim settlement approval; settle automatically on approval', CFG + ' (Claims)', 'On; on', 'Claims head', DISC, 'claims.settlement_maker_checker, claims.auto_settle_on_approval'),
    ('Maker-checker', 'Finance documents (payment vouchers, cheques, payouts, remittances, debit notes, petty cash)', CFG + ' (Accounting & Tax)', 'On', 'Accounting Manager', DISC, 'finance.maker_checker_enabled'),
    ('Maker-checker', 'Journal voucher approval', CFG + ' (Accounting & Tax)', 'On', 'Accounting Manager', DISC, 'journal.require_approval'),
    ('Maker-checker', 'Period close approval', CFG + ' (Accounting & Tax)', 'On', 'Accounting Manager', DISC, 'accounting.period_close_requires_approval'),
    ('Maker-checker', 'Reinsurance treaty approval', CFG + ' (Remittance & Reconciliation)', 'On', 'Processing Team lead', DISC, 'reinsurance.treaty_requires_approval'),
    ('Maker-checker', 'Posting rule and account determination changes', F + 'Configuration Approvals', 'Second approver required', 'Accounting Manager', DISC, ''),
    ('Remittance', 'Remittance approval levels by amount', 'Accounts > Remittance > Approval Workflow', 'Delivered levels', 'Accounting Manager', DISC, 'Confirm levels at the workshop'),
    ('Credit control', 'Client credit limits; instalment plan rules', 'Accounts > Credit Control > Client Credit Limits; Instalment Plans', 'None', 'Accounting Manager', CONF, ''),
    ('Notifications', 'Approval request notifications to approvers and decision notices to makers', CFG + ' (Notifications & E-mail)', 'On', 'System Administrator', CONF, 'notification.approval_requests'),
])

EMAIL = [
    ('Sending', 'Switch e-mail sending on (after the SMTP test)', CFG + ' (Notifications & E-mail)', 'Off', 'System Administrator', UAT, 'notification.email_enabled'),
    ('Sending', 'Sender name and address', CFG + ' (Notifications & E-mail)', 'BrokerVerse <connect@iortatechnxt.com>: replace', 'System Administrator', CONF, 'notification.from_address'),
    ('Sending', 'Largest total attachment size per e-mail', CFG + ' (Notifications & E-mail)', '10 MB', 'System Administrator', CONF, 'email.max_attachment_mb'),
    ('Clients', 'E-mail the official receipt PDF automatically when recorded; subject and text', CFG + ' (Billing, Collections & Credit)', 'Off; delivered text', 'Accounting Manager', CONF, 'receipts.email_on_record, receipts.email_subject, receipts.email_template'),
    ('Clients', 'E-mail the premium invoice PDF automatically when issued; subject, text, payment instructions', CFG + ' (Billing, Collections & Credit)', 'Off; delivered text', 'Accounting Manager', CONF, 'billing.email_on_issue, billing.email_*, billing.payment_instructions'),
    ('Clients', 'Premium payment reminder: subject and text', CFG + ' (Billing, Collections & Credit)', 'Delivered text', 'Accounting Manager', CONF, 'collections.email_subject, collections.email_template'),
    ('Clients', 'Renewal notices: labels, subject and text; order enforced', CFG + ' (Policies, Endorsements & Renewals)', 'First, Second, Final Notice; delivered text; order on', 'Operations head', CONF, 'renewals.notice_*'),
    ('Clients', 'Claim status updates and renewal reminders on or off', CFG + ' (Notifications & E-mail)', 'On; on', 'Operations head', CONF, 'notification.claim_status, notification.renewal_reminder'),
    ('Insurers', 'Preliminary loss advice: subject and text', CFG + ' (Claims)', 'Delivered text', 'Claims head', CONF, 'claims.pla_subject, claims.pla_template'),
    ('Insurers', 'Commission debit note e-mail', CFG + ' (Billing, Collections & Credit)', 'Delivered text', 'Accounting Manager', CONF, 'direct_bill.email_subject, direct_bill.email_body'),
    ('Insurers', 'Remittance statement and agency bill e-mails', CFG + ' (Remittance & Reconciliation)', 'Delivered text', 'Accounting Manager', CONF, 'remittance.statement_email_*, remittance.bill_email_*'),
    ('Users', 'Password reset e-mail', CFG + ' (Security & Access)', 'Delivered text; code valid 15 minutes', 'IT head', CONF, 'security.reset_email_*, security.reset_code_minutes'),
    ('Documents', 'Claim letters: acknowledgement, discharge voucher, data sheet', CFG + ' (Claims)', 'Delivered wording', 'Claims head', CONF, 'claims.documents'),
    ('Documents', 'Footer on official receipts (BIR permit or ATP details)', CFG + ' (Billing, Collections & Credit)', 'Empty', 'Accounting Manager', CONF, 'documents.receipt_footer'),
    ('Documents', 'Payment instructions on billing statements', CFG + ' (Billing, Collections & Credit)', 'Delivered text', 'Accounting Manager', CONF, 'documents.payment_instructions'),
    ('Documents', 'Disclaimer on the insurer comparison given to the client', CFG + ' (Sales, Quotations & Placement)', 'Delivered disclaimer', 'Compliance officer', CONF, 'packages.comparison_disclaimer'),
    ('Documents', 'Note on acknowledgement receipts', CFG + ' (Billing, Collections & Credit)', '"This acknowledgement receipt is not an official receipt..."', 'Accounting Manager', CONF, 'documents.acknowledgement_receipt_note'),
    ('Privacy', 'Privacy notice version recorded with consents', CFG + ' (Data Retention, Privacy & Uploads)', '1.0', 'DPO', CONF, 'privacy.notice_version'),
]
SHEETS['E-mail and Templates'] = ('EM', EMAIL)

SHEETS['Taxes and LGU'] = ('TX', [
    ('Premium taxes', 'DST on premium: method (12.5% or PHP 0.50 per PHP 4.00 unit, fraction rounded up) as each insurer applies it', F + 'Premium Taxes & LGU Rates', 'Per unit, PHP 0.50 per PHP 4.00 (12.5%)', 'Accounting Manager', DISC, 'Rule DST'),
    ('Premium taxes', 'VAT on premium for products under the VAT regime', F + 'Premium Taxes & LGU Rates', '12%', 'Accounting Manager', DISC, 'Rule VAT'),
    ('Premium taxes', 'Premium tax for products under the premium tax regime', F + 'Premium Taxes & LGU Rates', '2%', 'Accounting Manager', DISC, 'Rule PT'),
    ('Premium taxes', 'Fire service tax: lines to which it applies', F + 'Premium Taxes & LGU Rates', '2% on fire lines (fire, IAR, householder)', 'Accounting Manager', DISC, 'Rule FST'),
    ('LGU', 'LGT rate per city or municipality where clients are located', F + 'Premium Taxes & LGU Rates', '0.75% when the location has no rate', 'Accounting Manager', CONF, 'Rule LGT and LGU tax rates'),
    ('Other charges', 'Other charges on premium (for example notarial fee)', F + 'Premium Taxes & LGU Rates', 'Notarial fee rule delivered switched off', 'Accounting Manager', DISC, 'Rule NOTARIAL'),
    ('Tax codes', 'Confirm output and input VAT codes', F + 'Taxation', 'VAT12-OUT, VAT12-IN, VAT0, VATEX', 'Accounting Manager', CONF, 'Confirm with the tax adviser'),
    ('Tax codes', 'Confirm EWT codes and ATC used by the broker', F + 'Taxation', 'WC139, WC140, WI139, WI140, WI515, WI516, WC515, WC158, WI158, WC160, WI160, WC100, WI100, WC120, WI010, WI011, WC010, WC011', 'Accounting Manager', CONF, 'bir.atc_by_payee'),
    ('Tax codes', 'Final withholding on bank interest', F + 'Taxation', 'FWT-INT 20%', 'Accounting Manager', CONF, ''),
    ('VAT status', 'VAT-registered or non-VAT (percentage tax)', CFG + ' (Billing, Collections & Credit)', 'VAT-registered', 'Accounting Manager', DISC, 'direct_bill.broker_vat_registered; no percentage tax working paper'),
    ('BIR forms', 'Registered name, address, TIN of the withholding agent for BIR forms', G + 'Organization > Company', 'From the letterhead company', 'Accounting Manager', CONF, 'bir.* settings are a fallback only'),
    ('BIR forms', 'Which BrokerVerse documents serve as registered invoices and receipts under the EOPT Act', M + 'Document Numbering', 'INV, OR, DN series available', 'Accounting Manager', DISC, 'Decision with the tax adviser; registration is the broker\'s'),
    ('BIR forms', 'Computerized accounting system registration status and plan', 'Not a system setting', 'Not applicable', 'Accounting Manager', DISC, 'Philippine Regulatory Compliance Matrix'),
    ('BIR forms', 'SLSP: is the broker required to submit', 'Accounts > Tax > SLSP Sales; SLSP Purchases', 'Delivered reports', 'Accounting Manager', DISC, ''),
    ('Fallbacks', 'Fallback tax rates used only when no rule exists', CFG + ' (Accounting & Tax)', 'VAT 12%; DST 12.5%; LGT 0.75%; FST 2%', 'Accounting Manager', CONF, 'tax.vat_rate, tax.dst_rate, tax.lgt_rate, tax.fst_rate'),
])

MIGRATION = [
    ('Organisation', 'Company, branches, departments', 'Company_, Branch_, Department_Upload_Template.xlsx'),
    ('Organisation', 'Employee hierarchy, designations, employees', 'Hierarchy_, Designation_Upload_Template.xlsx'),
    ('Access', 'Users with one role each', 'Users_Provisioning_Template.xlsx'),
    ('Insurance', 'Insurers with commission rate, warranty days, remittance terms, billing mode', 'Insurance_Company_Upload_Template.xlsx'),
    ('Insurance', 'Lines of business, products, policy types, covers (only what is missing)', 'Line_of_Business_, Product_, Policy_Type_, Cover_Upload_Template.xlsx'),
    ('Insurance', 'Vehicle brands, models, variants', 'Vehicle_* templates'),
    ('Insurance', 'Signatories', 'Signatories_Upload_Template.xlsx'),
    ('Finance', 'Chart of accounts', 'Chart_of_Accounts_Upload_Template.xlsx'),
    ('Finance', 'Banks and bank accounts', 'Bank_, Bank_Account_Upload_Template.xlsx'),
    ('Finance', 'Currencies and exchange rates', 'Currency_, Exchange_Rate_Upload_Template.xlsx'),
    ('Finance', 'Transaction codes, write-off reasons', 'Transaction_Code_, Write_off_Reason_Upload_Template.xlsx'),
    ('Open business', 'In-force policies and their clients at go-live (insurer at 100%; co-insured entered on screen)', 'Policies_Upload_Template.xlsx (go-live mode)'),
    ('Open business', 'Open premium receivables at the day before go-live', 'Open_Items_Upload_Template.xlsx'),
    ('Open business', 'Trial balance at the day before go-live', 'Opening_Balances_Upload_Template.xlsx'),
    ('Open business', 'Amounts due to insurers on old bills (paid by payment voucher)', 'Disbursements_Upload_Template.xlsx'),
    ('Open business', 'Commission due to referrers earned before go-live', 'Disbursements_Upload_Template.xlsx'),
    ('Open business', 'Open claims on in-force policies (registered on screen after go-live)', 'No template'),
    ('Optional', 'Prospects and quotations in progress', 'Leads_, Quotations_Upload_Template.xlsx'),
]
mig_rows = []
for group, obj, template in MIGRATION:
    mig_rows.append((group, f'{obj}: source system and table or file, record count, extract owner, cleansing needed', 'See the template Instructions sheet', template, 'Data owner', MIG, 'Data Migration and Cutover Plan'))
mig_rows += [
    ('Scope', 'Data the broker wants to migrate beyond the OOTB scope (expired policies, closed claims, history, documents)', 'Not applicable', 'Out of OOTB scope', 'Sponsor', DISC, 'Change request: PHP 240,000.00 per legacy source (Rate Card)'),
    ('Volumes', 'In-force policies, clients, open bills, active referrers (counts)', 'Not applicable', 'Up to 20,000 rows per import file', 'Data owner', MOB, 'Larger books loaded in several files'),
    ('Schedule', 'Dates of mock load 1, mock load 2 (and 3), final extract', 'Not applicable', 'Two mock loads (small), three (medium, large)', 'Project manager', MOB, 'Implementation Approach and Plan'),
    ('Cutover', 'Freeze time of the old system and parallel run, if any', 'Not applicable', 'Freeze at close of the day before go-live', 'Sponsor', DISC, ''),
]
SHEETS['Data Migration Sources'] = ('DM', mig_rows)

SHEETS['Integrations'] = ('IG', [
    ('E-mail', 'SMTP mailbox: host, port (587, STARTTLS), user, who owns it', 'Environment variable SMTP_URL (secret store)', 'None', 'IT head', CONF, 'Test password reset and quotation approval link'),
    ('E-mail', 'Mail domain records (SPF, DKIM) so e-mails are not marked as spam', 'Broker DNS', 'Not applicable', 'IT head', CONF, ''),
    ('Bank', 'Bank statement file per account (format, frequency, who downloads)', 'Accounts > Bank Reconciliation (Import statement)', 'BDO, BPI, Metrobank, generic', 'Accounting', CONF, ''),
    ('Bank', 'Bank transactions file for remittance reconciliation', 'Accounts > Remittance > Reconciliation (Import)', 'CSV template', 'Accounting', CONF, 'Remittance_Bank_Transactions_Template'),
    ('Payment gateway', 'Merchant account, credentials and notification (webhook) address registration', F + 'Payment Gateways', 'PayMongo, Dragonpay (sandbox)', 'Accounting Manager', UAT, 'Sandbox test, then one live payment'),
    ('Insurers', 'Files exchanged with each insurer: slips, placement orders, statements, debit notes', 'Placement Slips; Insurer Reconciliation; Direct Bill Processing', 'PDF, CSV, Excel and e-mail', 'Processing Team lead', DISC, 'Insurer APIs are change requests'),
    ('Accounting', 'Corporate ledger or consolidation package to feed', 'Reports > Financial Reports', 'Excel or CSV export', 'Accounting Manager', DISC, 'Direct interface is a change request'),
    ('Single sign-on', 'Is sign-in through the corporate directory required', 'Not delivered', 'BrokerVerse user names and passwords with two-step verification', 'IT head', DISC, 'Gap (CR) if required'),
    ('SMS', 'SMS notifications to clients', 'Not delivered', 'E-mail only', 'Operations head', DISC, 'Gap (CR) if required'),
    ('API', 'Other systems that will call the BrokerVerse API (owner, purpose)', 'API (OpenAPI and Postman collection)', 'Every call checked against the user permissions', 'IT head', DISC, 'API and Dependency Catalogue'),
    ('Network', 'Office IP ranges, VPN or allow-listing needs; supported browsers', 'Hosting (WAF)', 'Current Chrome, Edge or Firefox', 'IT head', MOB, ''),
    ('Monitoring', 'Who receives monitoring alarms and the monthly service report', 'Hosting', 'iorta TechNXT on-call', 'IT head', CONF, 'Production Support Approach and Standards'),
])


def report_rows():
    rows = [('Schedules', 'Scheduled jobs: confirm times and which are switched on', M + 'Schedules', 'Renewal notices 06:00; Policy expiry 00:15; Quotation expiry 00:30; Receivable ageing 07:00; Daily reports 05:00; E-mail outbox every 5 minutes; Collection reminders 08:00; Housekeeping 02:45; Dormant accounts 01:45; Remittance schedules off', 'System Administrator', CONF, 'Asia/Manila time'),
            ('Retention', 'Generated report retention', CFG + ' (Data Retention, Privacy & Uploads)', '90 days', 'System Administrator', CONF, 'reports.retention_days'),
            ('Gaps', 'Reports the broker uses today that are not in the catalogue (attach samples)', 'Reports > All Reports', 'Catalogue below', 'Sponsor', DISC, 'New report: change request'),
            ('Regulators', 'Reports the broker submits to the IC and how they are prepared today', 'Reports > Financial Reports', 'No IC-format report', 'Compliance officer', DISC, '')]
    try:
        ws = load_workbook(REPORTS_BOOK, read_only=True)['Reports']
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not r or not r[1]:
                continue
            group, code, name, menu = r[0], r[1], r[2], r[3]
            rows.append((group, f'{name}: used? by whom, how often, any column or filter needs', menu, 'Delivered', 'Report owner', UAT, f'Report code {code}'))
    except Exception as exc:  # the Reports Book is built by another script; keep the sheet usable without it
        print('Reports Book not read:', exc)
    return rows


SHEETS['Reports'] = ('RP', report_rows())


def style_header(ws, row=1):
    for c in ws[row]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 30


def area_sheet(wb, title, prefix, rows):
    ws = wb.create_sheet(title[:31])
    ws.append(HEADER)
    style_header(ws)
    for i, (topic, question, screen, default, owner, needed, notes) in enumerate(rows, 1):
        ws.append([f'{prefix}-{i:02d}', topic, question, screen, default, '', owner, needed, '', 'Open', notes])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = BODY, WRAP, BOX
        row[5].fill = ANSWER_FILL
        row[8].number_format = 'DD MMM YYYY'
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    last = ws.max_row
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'J2:J{last}')
    dd = DataValidation(type='date', operator='greaterThan', formula1='DATE(2026,1,1)', allow_blank=True)
    ws.add_data_validation(dd)
    dd.add(f'I2:I{last}')
    ws.conditional_formatting.add(f'A2:K{last}', FormulaRule(formula=['$J2="Configured"'], fill=PatternFill('solid', fgColor='E2F0D9')))
    ws.conditional_formatting.add(f'A2:K{last}', FormulaRule(formula=['$J2="Gap (CR)"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    ws.print_title_rows = '1:1'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def main():
    wb = Workbook()
    ins = wb.active
    ins.title = 'Instructions'
    ins['A1'] = 'BrokerVerse OOTB: Discovery and Configuration Workbook'
    ins['A1'].font = TITLE
    lines = [
        ('Purpose', 'Records every configuration decision of the broker during mobilisation and the discovery workshops. It is the configuration workbook named in the Implementation Approach and Plan and is signed off at the end of discovery.'),
        ('How to use', 'One sheet per area. Each row is a question or a configuration item, the BrokerVerse screen where it is set, and the value delivered out of the box. The broker writes its decision in Client answer (yellow). Leave Client answer empty to keep the OOTB default and set Status to Confirmed.'),
        ('Columns', 'ID; Topic; Question or configuration item; BrokerVerse screen (menu path as on screen); OOTB default (seeded value); Client answer; Owner (broker role that answers); Needed by (project phase); Due date (set at kick-off); Status; Notes and setting key (technical key on Master > Configuration, upload template or reference document).'),
        ('Status values', 'Open: not yet answered. Answered: broker answer recorded. Confirmed: reviewed in the workshop by the owner and iorta TechNXT. Configured: entered in the test environment and checked. Not applicable: area not used. Gap (CR): cannot be met by configuration; recorded in the fit-gap register as a change request.'),
        ('Needed by', 'Mobilisation (week 1), Discovery (workshops), Configuration, Mock load 1, UAT. Due dates follow the plan of the implementation size.'),
        ('Defaults', 'OOTB defaults are the values seeded in the delivered system on 03 October 2026. Money in PHP. Tax rates, ATC and GL accounts are confirmed by the broker\'s accountant and tax adviser before configuration.'),
        ('Rules', 'Do not change the layout of the sheets. Add rows at the end of a sheet with the next ID if a new item comes up. Personal data (staff lists, client extracts) is not pasted into this workbook; it goes into the upload templates.'),
        ('Companion', 'Discovery Workbook Guide (how to run the workshops); Implementation Approach and Plan; Data Migration and Cutover Plan; upload templates in docs/package/05_Delivery/Upload_Templates.'),
    ]
    for i, (k, v) in enumerate(lines, 3):
        ins.cell(row=i, column=1, value=k).font = BOLD
        c = ins.cell(row=i, column=2, value=v)
        c.font, c.alignment = BODY, WRAP
        ins.cell(row=i, column=1).alignment = WRAP
    ins.column_dimensions['A'].width = 18
    ins.column_dimensions['B'].width = 120

    summary = wb.create_sheet('Summary')
    summary.append(['Sheet', 'Items'] + STATUSES + ['Percent configured'])
    style_header(summary)
    for title, (prefix, rows) in SHEETS.items():
        area_sheet(wb, title, prefix, rows)
        r = summary.max_row + 1
        ref = f"'{title}'!$J$2:$J$1000"
        summary.cell(row=r, column=1, value=title)
        summary.cell(row=r, column=2, value=f"=COUNTA('{title}'!$A$2:$A$1000)")
        for j, st in enumerate(STATUSES):
            summary.cell(row=r, column=3 + j, value=f'=COUNTIF({ref},"{st}")')
        pc = summary.cell(row=r, column=3 + len(STATUSES), value=f'=IF(B{r}=0,0,(F{r}+G{r})/B{r})')
        pc.number_format = '0%'
    r = summary.max_row + 1
    summary.cell(row=r, column=1, value='Total').font = BOLD
    for col in range(2, 3 + len(STATUSES)):
        letter = chr(64 + col)
        summary.cell(row=r, column=col, value=f'=SUM({letter}2:{letter}{r - 1})').font = BOLD
    t = summary.cell(row=r, column=3 + len(STATUSES), value=f'=IF(B{r}=0,0,(F{r}+G{r})/B{r})')
    t.number_format, t.font = '0%', BOLD
    for row in summary.iter_rows(min_row=2):
        for c in row:
            c.border = BOX
            if c.font != BOLD:
                c.font = BODY if not c.font.bold else c.font
    summary.column_dimensions['A'].width = 26
    for col in 'BCDEFGHI':
        summary.column_dimensions[col].width = 13
    summary.freeze_panes = 'A2'
    wb.move_sheet('Summary', offset=1 - wb.sheetnames.index('Summary'))
    for ws in (ins, summary):
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(OUT)
    print(OUT, {k: len(v[1]) for k, v in SHEETS.items()})


if __name__ == '__main__':
    main()
