"""Builds the iNXT BrokerVerse one-page brochure.

Outputs into docs/package/sales:
  iNXT_BrokerVerse_One_Page_Brochure.pdf   print and e-mail attachment (A4)
  iNXT_BrokerVerse_One_Page_Brochure.png   for WhatsApp, Viber and LinkedIn posts
  iNXT_BrokerVerse_One_Page_Brochure.docx  Word copy (the page as one picture)

Needs: playwright (Chromium), qrcode, python-docx.
Run:   python3 docs/package/tools/onepager/build_onepager.py
"""
import os
import qrcode
import docx
from docx.shared import Mm
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, '..', '..', 'sales'))
NAME = 'iNXT_BrokerVerse_One_Page_Brochure'
WEBSITE = 'https://www.iortatechnxt.com'


def make_qr():
    qr = qrcode.QRCode(border=1, box_size=12, error_correction=qrcode.constants.ERROR_CORRECT_M)
    qr.add_data(WEBSITE)
    qr.make(fit=True)
    qr.make_image(fill_color='#0A1F3D', back_color='white').save(os.path.join(HERE, 'img', 'qr.png'))


def render():
    pdf = os.path.join(OUT, NAME + '.pdf')
    png = os.path.join(OUT, NAME + '.png')
    exe = '/opt/pw-browsers/chromium' if os.path.exists('/opt/pw-browsers/chromium') else None
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=exe) if exe else p.chromium.launch()
        page = browser.new_page(viewport={'width': 794, 'height': 1123}, device_scale_factor=2.5)
        page.goto('file://' + os.path.join(HERE, 'onepager.html'), wait_until='networkidle')
        page.evaluate('document.fonts.ready')
        overflow = page.evaluate(
            "(() => { const p = document.querySelector('.page'); return p.scrollHeight - p.clientHeight; })()")
        if overflow > 0:
            raise SystemExit(f'Content overflows the A4 page by {overflow}px')
        page.pdf(path=pdf, format='A4', print_background=True, margin={'top': '0', 'right': '0', 'bottom': '0', 'left': '0'})
        page.locator('.page').screenshot(path=png)
        browser.close()
    return png


def word(png):
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(0)
    s.header_distance = s.footer_distance = Mm(0)
    par = d.paragraphs[0] if d.paragraphs else d.add_paragraph()
    par.paragraph_format.space_before = par.paragraph_format.space_after = 0
    par.add_run().add_picture(png, width=Mm(209.5))
    d.core_properties.title = 'iNXT BrokerVerse - One Page Brochure'
    d.core_properties.author = 'iorta TechNXT'
    d.save(os.path.join(OUT, NAME + '.docx'))


if __name__ == '__main__':
    make_qr()
    word(render())
    print('Written to', OUT)
